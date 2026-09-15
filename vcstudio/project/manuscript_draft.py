"""论文草稿骨架生成(骨架模式,用户已裁决)。

设计宪法在**写作**环节的可执行化——**草稿只组织真实数据,科学论断一律留 [作者补充] 占位,
绝不由工具代写机理/意义/评价**:

· **Methods 全自动**:复用 draftpack/methods_text 从**真实 INCAR/KPOINTS/POTCAR** 生成双语方法学
  + BibTeX(缺件降级 [待确认],绝不编造参数)。
· **Results 半自动**:逐图小节,每句由模板锚定**真实数值**(吸附能来自 adsorption.delta_e_rows,
  对照统计来自确定性 comparison),每个自动句子带溯源锚 ``<!--data: job_dir/figure_id-->``;每小节
  尾留 [作者补充:机理讨论]。
· **标题/摘要机理句/引言/结论**:纯占位([作者撰写/补充]),工具**不代写科学论断**。
· **参考文献**:Methods 的 BibTeX(真实、恒引,非编造)+ [作者补充]。

红线(落到代码与测试):自动句子里的数字**只可能**来自输入(项目真实计算值 / comparison);机理
**形容词黑名单**(优异/卓越…)绝不进自动句子(评价留给作者);占位符必在。fmt='docx' 为可选依赖,
python-docx 缺失则**降级**只出 .md(不抛)。中文注释,英文标识符。
"""
from __future__ import annotations

import os
import re
from pathlib import Path

from vcstudio.project import adsorption, draftpack

#: 机理/评价形容词黑名单——**绝不**出现在工具自动生成的句子里(科学评价留 [作者补充])。
BANNED_MECHANISM_WORDS = (
    '优异', '卓越', '优良', '出色', '杰出', '优越', '极佳', '最佳', '理想', '完美',
    '显著优于', '远超', '突破', '革命性', '前所未有', '有望成为', '证明了', '揭示了机理',
    'excellent', 'superior', 'outstanding', 'remarkable', 'unprecedented',
)

_MINUS = '−'

#: 图 kind → 归入吸附族 / 对照族(其余走占位)。
_ADS_KINDS = frozenset({'', 'adsorption', 'adsorption_bar', 'adsorption_energy',
                        'volcano', 'bar', 'free_energy', 'other'})
_CMP_KINDS = frozenset({'comparison', 'validation', 'mae', 'parity'})

_LABELS = {
    'zh': {
        'title': '[作者补充:论文标题]',
        'abstract_h': '## 摘要',
        'abstract_ph': '[作者补充:研究动机、机理解释与科学意义(工具不代写论断)]',
        'intro_h': '## 引言',
        'intro_ph': '[作者撰写:研究背景与文献综述,可参考 AI 分析报告]',
        'methods_fallback': '[待确认:项目暂无 DONE 成员,Methods 待计算完成后依真实 INCAR 生成]',
        'results_h': '## 结果与讨论',
        'discuss_ph': '[作者补充:机理讨论]',
        'concl_h': '## 结论',
        'concl_ph': '[作者撰写:主要结论与展望]',
        'ref_h': '## 参考文献',
        'ref_ph': '[作者补充:与本工作直接相关的文献引用]',
        'sections': ['标题', '摘要', '引言', '方法', '结果与讨论', '结论', '参考文献'],
    },
    'en': {
        'title': '[TO BE WRITTEN: manuscript title]',
        'abstract_h': '## Abstract',
        'abstract_ph': '[TO BE WRITTEN: motivation, mechanism and significance (not auto-written)]',
        'intro_h': '## Introduction',
        'intro_ph': '[TO BE WRITTEN: background and literature review; see AI analysis report]',
        'methods_fallback': '[TO CONFIRM: no DONE member; Methods pending real INCAR]',
        'results_h': '## Results and Discussion',
        'discuss_ph': '[TO BE WRITTEN: mechanistic discussion]',
        'concl_h': '## Conclusions',
        'concl_ph': '[TO BE WRITTEN: main conclusions and outlook]',
        'ref_h': '## References',
        'ref_ph': '[TO BE WRITTEN: references directly relevant to this work]',
        'sections': ['Title', 'Abstract', 'Introduction', 'Methods',
                     'Results and Discussion', 'Conclusions', 'References'],
    },
}


# ── 数值格式化(所有自动句子统一 2 位小数;负号用排印减号) ──────────────────────
def _ev(v, nd=2):
    f = float(v)
    s = f'{abs(f):.{nd}f}'
    return (_MINUS + s) if f < 0 else s


def _anchor(text, job_dir, figure_id):
    """给一句自动文本挂溯源锚 <!--data: job_dir/figure_id-->。"""
    return f'{text} <!--data: {job_dir}/{figure_id}-->'


def _fig_kind(f):
    return str((f or {}).get('kind') or '').strip().lower()


def _figure_for(figs, kinds, default_id, default_title):
    """从 manifest 找首个匹配 kind 的图;没有则用默认图号/标题。"""
    for f in (figs or []):
        if _fig_kind(f) in kinds:
            return dict(f)
    return {'figure_id': default_id, 'title': default_title}


def _fig_ref(fig):
    fid = str(fig.get('figure_id') or '图')
    panel = str(fig.get('panel') or '')
    return fid + panel


# ── Methods 全自动(复用 draftpack.methods_bundle:真实 INCAR → 双语 + BibTeX) ─────
def _methods_and_bib(project, methods_dir, lang):
    """返回 (methods_body, bibtex)。缺 DONE 成员/异常 → 降级占位,绝不编造参数。"""
    try:
        res = draftpack.methods_bundle(project, methods_dir)
        files = {os.path.basename(p): p for p in res.get('files', [])}
        mkey = 'methods_zh.md' if lang == 'zh' else 'methods_en.md'
        body = ''
        if mkey in files and os.path.isfile(files[mkey]):
            body = Path(files[mkey]).read_text(encoding='utf-8').strip()
        bib = ''
        if 'references.bib' in files and os.path.isfile(files['references.bib']):
            bib = Path(files['references.bib']).read_text(encoding='utf-8').strip()
        return body, bib
    except Exception:                                        # noqa: BLE001 方法学装订任何异常都降级
        return '', ''


# ── Results 数据句(逐行锚定真实吸附能;最稳/极值;对照统计) ──────────────────────
def _adsorption_lines(system, rows, fig):
    """逐构型吸附能句 + 最稳 + 极值区间;每句带溯源锚。数字**只来自** rows 的真实 ΔE。"""
    fref = _fig_ref(fig)
    lines = []
    valued = [r for r in rows if r.get('delta_e') is not None]
    for r in valued:
        species = r.get('species') or r.get('name') or '吸附质'
        job_dir = fig.get('job_dir') or r.get('name') or 'job'
        lines.append('- ' + _anchor(
            f'{system} 对 {species} 的吸附能为 {_ev(r["delta_e"])} eV({fref})',
            job_dir, fig.get('figure_id', '图')))
    if valued:
        job_dir = fig.get('job_dir') or 'jobs'
        stable = min(valued, key=lambda r: r['delta_e'])
        sp = stable.get('species') or stable.get('name') or '吸附质'
        lines.append('- ' + _anchor(
            f'在所纳入构型中,{system} 对 {sp} 的吸附能最低(ΔE = {_ev(stable["delta_e"])} eV),'
            f'为最稳吸附构型({fref})', job_dir, fig.get('figure_id', '图')))
        dmin = min(r['delta_e'] for r in valued)
        dmax = max(r['delta_e'] for r in valued)
        lines.append('- ' + _anchor(
            f'各构型吸附能介于 {_ev(dmin)} 至 {_ev(dmax)} eV({fref})',
            job_dir, fig.get('figure_id', '图')))
    return lines


def _comparison_lines(comparison, fig):
    """文献对照统计句(MAE/RMSE/偏差最大项);数字**只来自** comparison(确定性算出,非 LLM)。"""
    fref = _fig_ref(fig)
    job_dir = fig.get('job_dir') or 'validation'
    fid = fig.get('figure_id', '图')
    n = comparison.get('n') or 0
    lines = []
    if not n:
        return ['- ' + _anchor(f'文献对照:暂无可对齐项({fref})', job_dir, fid)]
    lines.append('- ' + _anchor(
        f'与文献对照共 {n} 项:平均绝对误差 MAE = {_ev(comparison.get("mae"))} eV,'
        f'均方根误差 RMSE = {_ev(comparison.get("rmse"))} eV({fref})', job_dir, fid))
    worst = comparison.get('worst') or []
    if worst:
        w = worst[0]
        lines.append('- ' + _anchor(
            f'偏差最大项为 {w.get("system")}/{w.get("species")}:计算 {_ev(w.get("ours"))} eV 对'
            f'文献 {_ev(w.get("ref"))} eV(Δ = {_ev(w.get("delta"))} eV)({fref})', job_dir, fid))
    return lines


def _results_section(project, rows, comparison, figures_manifest, labels):
    """组装 Results:吸附数据小节 + 对照小节 + manifest 里其余图的占位小节。"""
    system = str((project or {}).get('name') or '本体系')
    figs = list(figures_manifest or [])
    out = [labels['results_h'], '']

    ads_fig = _figure_for(figs, _ADS_KINDS, '图1', '各体系吸附能对比')
    cmp_default = '图3' if ads_fig.get('figure_id') == '图2' else '图2'
    cmp_fig = _figure_for(figs, _CMP_KINDS, cmp_default, '计算值与文献值对照')

    if [r for r in rows if r.get('delta_e') is not None]:
        out += [f'### {ads_fig.get("figure_id", "图1")} {ads_fig.get("title", "吸附能")}', '']
        out += _adsorption_lines(system, rows, ads_fig)
        out += ['', labels['discuss_ph'], '']

    if comparison:
        out += [f'### {cmp_fig.get("figure_id", "图2")} {cmp_fig.get("title", "文献对照")}', '']
        out += _comparison_lines(comparison, cmp_fig)
        out += ['', labels['discuss_ph'], '']

    # manifest 里未被数据小节消费的图 → 占位小节(结构在,论断留作者)
    for f in figs:
        if _fig_kind(f) in _ADS_KINDS and ads_fig.get('figure_id') == f.get('figure_id'):
            continue
        if _fig_kind(f) in _CMP_KINDS and cmp_fig.get('figure_id') == f.get('figure_id'):
            continue
        out += [f'### {f.get("figure_id", "图")} {f.get("title", "")}'.rstrip(), '',
                labels['discuss_ph'], '']
    return out


# ═══════════════════════════════════════════════════════════════════════════════
# 主入口
# ═══════════════════════════════════════════════════════════════════════════════
def _abstract_lines(project, rows, comparison, labels):
    """摘要:数据事实句(带锚) + 机理/意义占位。事实数字只来自 rows/comparison。"""
    system = str((project or {}).get('name') or '本体系')
    valued = [r for r in rows if r.get('delta_e') is not None]
    out = [labels['abstract_h'], '']
    if valued:
        stable = min(valued, key=lambda r: r['delta_e'])
        out.append(_anchor(
            f'本工作对 {system} 计算了 {len(valued)} 个吸附构型,最稳构型吸附能为 '
            f'{_ev(stable["delta_e"])} eV。', 'jobs', '图1'))
    if comparison and (comparison.get('n') or 0):
        out.append(_anchor(
            f'与文献数据对照共 {comparison["n"]} 项,MAE = {_ev(comparison.get("mae"))} eV。',
            'validation', '图2'))
    out += ['', labels['abstract_ph'], '']
    return out


def build_manuscript(project, comparison=None, figures_manifest=None, *,
                     lang='zh', fmt='markdown'):
    """项目 → 论文草稿骨架 {'ok','path','sections','placeholders_count', ...}。

    组装:标题(占位)/摘要(数据事实 + [作者补充])/引言(全占位)/**Methods 全自动**(真实 INCAR)/
    Results(逐图数据句 + 每小节尾 [作者补充:机理讨论])/结论(全占位)/参考文献(BibTeX + [作者补充])。
    每个自动句子带 ``<!--data: job_dir/figure_id-->`` 溯源锚。fmt='markdown' 落 .md;'docx' 时若
    python-docx 可用则**同时**出 .docx(缺依赖则降级只出 .md,不抛)。输出目录取 project['root']/manuscript。
    """
    labels = _LABELS.get(lang, _LABELS['zh'])
    project = project or {}
    out_dir = Path(str(project.get('root') or '.')) / 'manuscript'
    out_dir.mkdir(parents=True, exist_ok=True)

    try:
        rows = adsorption.delta_e_rows(project).get('rows', [])
    except Exception:                                        # noqa: BLE001 取数异常 → 空(不编造)
        rows = []

    methods_body, bibtex = _methods_and_bib(project, str(out_dir / 'methods'), lang)

    lines = [f'# {labels["title"]}', '']
    lines += _abstract_lines(project, rows, comparison, labels)
    lines += [labels['intro_h'], '', labels['intro_ph'], '']
    lines += [methods_body if methods_body else labels['methods_fallback'], '']
    lines += _results_section(project, rows, comparison, figures_manifest, labels)
    lines += [labels['concl_h'], '', labels['concl_ph'], '']
    lines += [labels['ref_h'], '']
    if bibtex:
        lines += ['```bibtex', bibtex, '```', '']
    lines += [labels['ref_ph'], '']

    text = '\n'.join(lines) + '\n'
    md_path = out_dir / ('manuscript_zh.md' if lang == 'zh' else 'manuscript_en.md')
    md_path.write_text(text, encoding='utf-8')

    placeholders_count = text.count('[作者') + text.count('[TO BE WRITTEN')

    result = {'ok': True, 'path': str(md_path), 'md_path': str(md_path),
              'sections': list(labels['sections']),
              'placeholders_count': placeholders_count,
              'docx_path': None, 'docx_available': False}

    if fmt == 'docx':
        docx_path = out_dir / (md_path.stem + '.docx')
        if _write_docx(text, docx_path):
            result['docx_path'] = str(docx_path)
            result['docx_available'] = True
            result['path'] = str(docx_path)
        else:
            result['note'] = '未安装 python-docx,已降级只出 Markdown(.md)'
    return result


def _write_docx(text, docx_path):
    """Markdown 文本 → 简版 .docx(标题层级 + 段落)。python-docx 缺失 → 返回 False(降级)。"""
    try:
        import docx
    except Exception:                                        # noqa: BLE001 可选依赖缺失即降级
        return False
    try:
        doc = docx.Document()
        for line in text.split('\n'):
            if line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.strip():
                doc.add_paragraph(line)
        doc.save(str(docx_path))
        return True
    except Exception:                                        # noqa: BLE001 写盘异常也降级(不拖垮 md)
        return False


def draft_stats(path):
    """草稿溯源统计 {'auto','placeholder','total','auto_ratio'}——诚实展示「AI 组织了多少、留给你多少」。

    auto = 带溯源锚 ``<!--data:`` 的自动句子数;placeholder = [作者…]/[TO BE WRITTEN…] 占位数。
    """
    try:
        text = Path(path).read_text(encoding='utf-8')
    except OSError:
        return {'auto': 0, 'placeholder': 0, 'total': 0, 'auto_ratio': 0.0}
    auto = len(re.findall(r'<!--data:', text))
    placeholder = text.count('[作者') + text.count('[TO BE WRITTEN')
    total = auto + placeholder
    return {'auto': auto, 'placeholder': placeholder, 'total': total,
            'auto_ratio': (round(auto / total, 3) if total else 0.0)}
