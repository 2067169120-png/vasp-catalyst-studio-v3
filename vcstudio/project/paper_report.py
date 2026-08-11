"""Render one portable, paper-style report bundle from a shared data model.

The renderer deliberately has no dependency on project calculation code.  Callers
prepare a JSON-like ``model`` once, and the HTML, DOCX, and PDF renderers consume
that same normalized snapshot.  Figure files are copied by content hash into the
destination bundle before rendering, so reports never need a relative path between
different Windows drives.
"""
from __future__ import annotations

import base64
import copy
import errno
import hashlib
import html
import importlib
import json
import math
import mimetypes
import os
import re
import shutil
import sys
import tempfile
import threading
import time
from contextlib import contextmanager
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Mapping, Sequence
from xml.sax.saxutils import escape as xml_escape


BUNDLE_SCHEMA = "vcstudio.paper-report.bundle/v2"
MODEL_SCHEMA = "vcstudio.paper-report.model/v2"
PREVIEW_SCHEMA = "vcstudio.paper-report.html-preview/v1"
_LEGACY_MODEL_SCHEMA = "vcstudio.paper-report.model/v1"
_REPORT_KINDS = ("diagnostic", "final", "draft")
_QUALIFICATION_LEVELS = (
    "diagnostic",
    "adsorption_result_verified",
    "thermodynamic_path_verified",
    "kinetic_evidence_verified",
    "publication_package_verified",
    "human_scientific_reviewed",
)
_ALLOWED_FORMATS = ("html", "docx", "pdf")
_RASTER_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}
_FIGURE_PATH_KEYS = ("path", "src", "image", "file")
_A4_CONTENT_WIDTH_DXA = 9298  # 210 mm page - 23 mm left/right margins.
_PDF_FONT_LOCK = threading.Lock()
_REPORT_PUBLISH_LOCK = threading.RLock()
_PDF_CJK_REGULAR = "PaperCJK"
_PDF_CJK_BOLD = "PaperCJKBold"
_CONTRACT_ADMIN_KEYS = {
    "created_at", "created_at_utc", "generated_at", "generated_at_utc",
    "updated_at", "updated_at_utc", "validated_at", "validated_at_utc",
}
_CONTRACT_SCHEMAS = {
    "spec": "vcstudio.report-spec/v1",
    "snapshot": "vcstudio.report-snapshot/v1",
    "validation": "vcstudio.report-validation/v1",
}
_VALIDATION_STATUSES = {"passed", "passed_with_warnings", "blocked", "unknown"}
_SHA256_HEX_RE = re.compile(r"[0-9a-f]{64}")
_CONTENT_FINGERPRINT_KEYS = (
    "report_kind", "scientific_qualification", "claim_ceiling",
    "template_ref",
    "locale", "outline", "title", "subtitle", "kicker", "metadata",
    "executive_summary", "key_findings", "candidate_evaluations",
    "adsorption_table", "comparison_table", "figures", "methods",
    "limitations", "recommendations", "comparison_context",
)
_REPORT_THEMES = {
    "academic-a4": {
        "ink": "#17212B",
        "muted": "#66717C",
        "accent": "#1F4F64",
        "highlight": "#8A6A2E",
        "rule": "#23313A",
        "title": "#203846",
        "subtitle": "#435866",
        "caption": "#39464E",
        "screen": "#E9ECEF",
        "page_margin_mm": 23,
        "cover_spacer_mm": 43,
    },
    "compact-brief": {
        "ink": "#132A2E",
        "muted": "#526A6D",
        "accent": "#0F766E",
        "highlight": "#0F766E",
        "rule": "#134E4A",
        "title": "#123F3C",
        "subtitle": "#355E5A",
        "caption": "#355E5A",
        "screen": "#E7EFEE",
        "page_margin_mm": 18,
        "cover_spacer_mm": 26,
    },
    "diagnostic-a4": {
        "ink": "#2B2118",
        "muted": "#756353",
        "accent": "#9A3412",
        "highlight": "#B45309",
        "rule": "#7C2D12",
        "title": "#7C2D12",
        "subtitle": "#795548",
        "caption": "#6B4F3F",
        "screen": "#F3ECE6",
        "page_margin_mm": 23,
        "cover_spacer_mm": 38,
    },
}
_FORMAT_ACCESSIBILITY = {
    "html": {
        "status": "conditional",
        "reason": (
            "生成语义 HTML（含 lang、标题层级和图片 alt）；是否可访问仍取决于每幅"
            "非装饰图片具有有意义的替代文本并完成人工检查。"
        ),
        "reason_zh": (
            "生成语义 HTML（含 lang、标题层级和图片 alt）；是否可访问仍取决于每幅"
            "非装饰图片具有有意义的替代文本并完成人工检查。"
        ),
        "reason_en": (
            "Semantic HTML emits lang, heading structure, and image alt text; "
            "accessibility still depends on meaningful text for every "
            "non-decorative figure and manual review."
        ),
        "visual": True,
        "searchable": True,
        "semantic_structure": True,
        "document_language": True,
        "metadata": True,
        "image_alt": True,
        "tagged": None,
        "pdf_ua": None,
        "manual_review_required": True,
    },
    "docx": {
        "status": "conditional",
        "reason": (
            "Word 输出包含标题样式、文档语言、重复表头、元数据和图片替代文本；仍须"
            "保证替代文本有意义，并通过 Word 可访问性检查器和人工阅读顺序检查。"
        ),
        "reason_zh": (
            "Word 输出包含标题样式、文档语言、重复表头、元数据和图片替代文本；仍须"
            "保证替代文本有意义，并通过 Word 可访问性检查器和人工阅读顺序检查。"
        ),
        "reason_en": (
            "Word output includes heading styles, document language, repeating "
            "table headers, metadata, and image alt text; meaningful alt text plus "
            "Word Accessibility Checker and manual reading-order review remain required."
        ),
        "visual": True,
        "searchable": True,
        "semantic_structure": True,
        "document_language": True,
        "metadata": True,
        "image_alt": True,
        "tagged": None,
        "pdf_ua": None,
        "manual_review_required": True,
    },
    "pdf": {
        "status": "partial",
        "reason": (
            "ReportLab 仅生成可视、可搜索 PDF；当前输出未标记，不包含结构树或图片"
            "替代文本，pdf_ua=false。"
        ),
        "reason_zh": (
            "ReportLab 仅生成可视、可搜索 PDF；当前输出未标记，不包含结构树或图片"
            "替代文本，pdf_ua=false。"
        ),
        "reason_en": (
            "ReportLab produces a visual, searchable PDF only. The output is "
            "untagged, has no structure tree or image alt text, and pdf_ua=false."
        ),
        "visual": True,
        "searchable": True,
        "semantic_structure": False,
        "document_language": True,
        "metadata": True,
        "image_alt": False,
        "tagged": False,
        "pdf_ua": False,
        "manual_review_required": True,
    },
}
_GENERIC_FIGURE_ALT_RE = re.compile(
    r"(?i)^(?:(?:figure|fig\.?|image|plot|chart)\s*(?:#|no\.?\s*)?\d+"
    r"|(?:图|图片|图表)\s*[一二三四五六七八九十百\d]+)[.。]?$"
)
_SECTION_LABELS = {
    "en-US": {
        "executive_summary": "Executive Summary",
        "key_findings": "Key Findings",
        "candidate_evaluations": "Candidate Evaluation",
        "adsorption_table": "Adsorption-Energy Results",
        "comparison_table": "Cross-Project Comparison",
        "figures": "Figures",
        "methods": "Methods",
        "limitations": "Limitations",
        "recommendations": "Recommendations",
    },
    "zh-CN": {
        "executive_summary": "执行摘要",
        "key_findings": "核心结论",
        "candidate_evaluations": "候选评价",
        "adsorption_table": "吸附能结果",
        "comparison_table": "多项目比较",
        "figures": "图表",
        "methods": "计算方法",
        "limitations": "局限性",
        "recommendations": "后续建议",
    },
}


class ReportDependencyError(RuntimeError):
    """Raised when an explicitly requested report format is unavailable."""


class ReportRecoveryError(OSError):
    """Raised when publication failed and automatic rollback was incomplete.

    ``recovery_dir`` is deliberately retained on disk.  It contains the
    surviving pre-publication backups and a machine-readable recovery record;
    callers must not treat the destination as a complete report generation.
    """

    def __init__(self, message: str, *, recovery_dir: Path):
        super().__init__(message)
        self.recovery_dir = Path(recovery_dir)


@contextmanager
def _report_publish_guard(destination: Path, stem: str):
    """Serialize one report generation across threads and OS processes.

    The persistent hidden lock file must not be unlinked after release: removing
    it can split waiters across two inodes and re-introduce the publication race.
    The operating system releases the advisory lock automatically if a writer
    exits or crashes.
    """

    lock_path = Path(destination) / f".{stem}.publish.lock"
    with _REPORT_PUBLISH_LOCK:
        lock_path.parent.mkdir(parents=True, exist_ok=True)
        with lock_path.open("a+b") as handle:
            handle.seek(0, os.SEEK_END)
            if handle.tell() == 0:
                handle.write(b"\0")
                handle.flush()

            if os.name == "nt":
                import msvcrt

                while True:
                    handle.seek(0)
                    try:
                        msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
                        break
                    except OSError as exc:
                        if exc.errno not in {errno.EACCES, errno.EAGAIN, errno.EDEADLK}:
                            raise
                        time.sleep(0.05)
                try:
                    yield lock_path
                finally:
                    handle.seek(0)
                    msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl

                fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
                try:
                    yield lock_path
                finally:
                    fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


def report_capabilities() -> dict:
    """Return render availability separately from accessibility support.

    ``formats`` remains the generation dependency gate used by existing callers.
    ``accessibility`` is deliberately independent: a renderer can be available
    while its output is only conditionally accessible (HTML/DOCX) or explicitly
    partial and non-PDF/UA (the current ReportLab PDF renderer).
    """
    formats = {}
    for fmt in _ALLOWED_FORMATS:
        try:
            _check_dependencies((fmt,))
        except Exception as exc:  # noqa: BLE001 - capability probe is per-format
            formats[fmt] = {"available": False, "reason": str(exc)}
        else:
            formats[fmt] = {"available": True, "reason": ""}
    return {
        "schema": "vcstudio.paper-report.capabilities/v1",
        "formats": formats,
        "accessibility": copy.deepcopy(_FORMAT_ACCESSIBILITY),
    }


def render_report_bundle(
    model: Mapping[str, Any],
    out_dir,
    stem: str = "report",
    formats: Sequence[str] = ("html", "docx", "pdf"),
) -> dict:
    """Render a portable HTML/DOCX/PDF bundle from one semantic report model.

    Parameters
    ----------
    model:
        JSON-like mapping.  Supported top-level fields include ``title``,
        ``subtitle``, ``kicker``, ``metadata``, ``executive_summary``,
        ``key_findings``, ``candidate_evaluations``, ``adsorption_table``,
        ``comparison_table``, ``figures``, ``methods``, ``limitations``, and
        ``recommendations``.
    out_dir:
        Destination directory.  All figures are copied into ``out_dir/assets``.
    stem:
        Safe basename used for report files and the manifest.
    formats:
        Any non-empty subset of ``html``, ``docx``, and ``pdf``.

    Returns
    -------
    dict
        ``files`` maps each requested format plus ``manifest`` to a :class:`Path`;
        ``contract_files`` maps every published full-contract sidecar to a
        :class:`Path`, ``model_file`` points to the canonical normalized model
        projection, and ``assets`` is a list of staged :class:`Path` objects. The
        model hash and schema are also returned for downstream completion markers.
    """
    if not isinstance(model, Mapping):
        raise TypeError("model must be a mapping")
    model = dict(model)
    safe_stem = _validate_stem(stem)
    requested = _normalize_formats(formats)
    _check_dependencies(requested)

    destination = Path(out_dir).expanduser()
    destination.mkdir(parents=True, exist_ok=True)
    if not destination.is_dir():
        raise NotADirectoryError(f"report output is not a directory: {destination}")

    normalized = _normalize_model(model, requested_formats=requested)
    temp_root = Path(tempfile.mkdtemp(prefix=f".{safe_stem}.tmp-", dir=destination))
    preserve_temp_root = False
    try:
        accessibility_records = _report_accessibility(normalized, requested)
        staged_model, asset_records = _stage_figures(
            normalized,
            temp_root / "assets",
            requested,
        )
        report_model_sha256 = _sha256_json(_content_fingerprint(staged_model))
        declared_report_model_sha256 = _text(
            (normalized.get("validation") or {}).get("report_model_sha256")
        ).lower()
        if (declared_report_model_sha256
                and declared_report_model_sha256 != report_model_sha256):
            raise ValueError(
                "rendered report content conflicts with "
                "ValidationResult.report_model_sha256"
            )
        fingerprint_model = _fingerprint_model(staged_model)
        model_sha256 = _sha256_json(fingerprint_model)
        temp_model_file = temp_root / f"{safe_stem}.model.json"
        temp_model_file.write_bytes(_canonical_json_bytes(fingerprint_model))
        model_file_record = _file_record(
            temp_model_file, relative_path=temp_model_file.name)
        if model_file_record["sha256"] != model_sha256:
            raise AssertionError(
                "canonical model sidecar bytes do not match model_sha256"
            )

        temp_outputs: dict[str, Path] = {}
        if "html" in requested:
            path = temp_root / f"{safe_stem}.html"
            _render_html(staged_model, path)
            temp_outputs["html"] = path
        if "docx" in requested:
            path = temp_root / f"{safe_stem}.docx"
            _render_docx(staged_model, path)
            temp_outputs["docx"] = path
        if "pdf" in requested:
            path = temp_root / f"{safe_stem}.pdf"
            _render_pdf(staged_model, path)
            temp_outputs["pdf"] = path

        contract_outputs, contract_records = _stage_contract_sidecars(
            normalized, temp_root, safe_stem)
        obsolete_contract_paths = tuple(
            destination / f"{safe_stem}.{key}.json"
            for key in _CONTRACT_SCHEMAS
            if f"contract_{key}" not in contract_outputs
        )
        obsolete_format_paths = tuple(
            destination / f"{safe_stem}.{fmt}"
            for fmt in _ALLOWED_FORMATS
            if fmt not in requested
        )

        file_records = {
            fmt: _file_record(path, relative_path=path.name)
            for fmt, path in temp_outputs.items()
        }
        manifest = {
            "schema": BUNDLE_SCHEMA,
            "artifact_status": "complete",
            "model_schema": normalized["schema"],
            "source_model_schema": normalized["source_model_schema"],
            "model_sha256": model_sha256,
            "report_model_sha256": report_model_sha256,
            "report_id": normalized["report_id"],
            "report_kind": normalized["report_kind"],
            "scientific_status": normalized["scientific_status"],
            "scientific_qualification": normalized["scientific_qualification"],
            "claim_ceiling": normalized["claim_ceiling"],
            "input_fingerprint": normalized["input_fingerprint"],
            "preset_id": normalized["preset_id"],
            "revision": normalized["revision"],
            "contract_status": normalized["contract_status"],
            "contracts": contract_records,
            "model_file": model_file_record,
            "generated_at_utc": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
            "formats": list(requested),
            "accessibility": accessibility_records,
            "files": file_records,
            "assets": asset_records,
        }
        temp_manifest = temp_root / f"{safe_stem}.manifest.json"
        temp_manifest.write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )

        # Publishing is deliberately serialized across threads and cooperating
        # OS processes.  Manual export, background workers, or two application
        # instances can otherwise interleave replacements for the same stem and
        # leave a manifest referring to another writer's files.  Assets are
        # content-addressed and can safely be committed first.
        with _report_publish_guard(destination, safe_stem):
            asset_paths = _commit_assets(
                temp_root / "assets",
                destination / "assets",
                asset_records,
            )
            published_paths, final_manifest = _commit_report_files(
                {**temp_outputs, "model": temp_model_file, **contract_outputs},
                temp_manifest,
                destination,
                temp_root / "rollback",
                obsolete_paths=(*obsolete_contract_paths, *obsolete_format_paths),
            )
            contract_paths = {
                key.removeprefix("contract_"): published_paths.pop(key)
                for key in list(published_paths)
                if key.startswith("contract_")
            }
            model_file = published_paths.pop("model")
            output_paths = published_paths
    except ReportRecoveryError:
        # The rollback directory is the only remaining copy of one or more
        # pre-publication files.  Keep the entire generation directory intact
        # so a human or recovery tool can restore it.
        preserve_temp_root = True
        raise
    finally:
        if not preserve_temp_root:
            shutil.rmtree(temp_root, ignore_errors=True)

    return {
        "schema": BUNDLE_SCHEMA,
        "artifact_status": "complete",
        "model_sha256": model_sha256,
        "report_model_sha256": report_model_sha256,
        "report_kind": normalized["report_kind"],
        "scientific_status": normalized["scientific_status"],
        "scientific_qualification": normalized["scientific_qualification"],
        "input_fingerprint": normalized["input_fingerprint"],
        "contract_status": normalized["contract_status"],
        "contracts": contract_records,
        "contract_files": contract_paths,
        "model_file": model_file,
        "revision": normalized["revision"],
        "files": output_paths,
        "assets": asset_paths,
        "accessibility": accessibility_records,
        "manifest": final_manifest,
    }


def render_report_html_preview(model: Mapping[str, Any]) -> dict:
    """Render a deterministic, self-contained HTML preview without publishing.

    Preview is deliberately a separate seam from :func:`render_report_bundle`:
    it performs the same model normalization, full-contract validation, section
    planning and HTML rendering, but never creates an output directory, model or
    contract sidecars, a manifest, a marker, history, or a report revision.  Any
    figures are read once and embedded as ``data:`` URIs so the returned payload
    contains no source or temporary filesystem path.

    ``ReportSpec.formats`` describes the eventual publication request, not the
    transport used to inspect this preview.  Consequently the normalizer is
    invoked without a requested-format override; a later formal publication will
    still enforce the exact requested format set and its dependencies.
    """
    if not isinstance(model, Mapping):
        raise TypeError("model must be a mapping")

    normalized = _normalize_model(dict(model))
    preview_model = _inline_preview_figures(normalized)
    report_model_sha256 = _sha256_json(_content_fingerprint(preview_model))
    declared_report_model_sha256 = _text(
        (normalized.get("validation") or {}).get("report_model_sha256")
    ).lower()
    if (declared_report_model_sha256
            and declared_report_model_sha256 != report_model_sha256):
        raise ValueError(
            "rendered report content conflicts with "
            "ValidationResult.report_model_sha256"
        )

    model_sha256 = _sha256_json(_fingerprint_model(preview_model))
    return {
        "schema": PREVIEW_SCHEMA,
        "artifact_status": "preview",
        "html": _html_document(preview_model),
        "model_sha256": model_sha256,
        "report_model_sha256": report_model_sha256,
        "report_id": preview_model["report_id"],
        "report_kind": preview_model["report_kind"],
        "scientific_status": preview_model["scientific_status"],
        "scientific_qualification": preview_model["scientific_qualification"],
        "claim_ceiling": preview_model["claim_ceiling"],
        "input_fingerprint": preview_model["input_fingerprint"],
        "preset_id": preview_model["preset_id"],
        "contract_status": preview_model["contract_status"],
        "contract_refs": _json_safe(preview_model["contract_refs"]),
        "accessibility": _report_accessibility(normalized, ("html",))["html"],
    }


def _validate_stem(stem: str) -> str:
    value = str(stem).strip()
    if (
        not value
        or value in {".", ".."}
        or "/" in value
        or "\\" in value
        or ":" in value
        or Path(value).name != value
    ):
        raise ValueError("stem must be a safe filename without directory components")
    return value


def _normalize_formats(formats: Sequence[str]) -> tuple[str, ...]:
    if isinstance(formats, str):
        formats = (formats,)
    result: list[str] = []
    for item in formats:
        fmt = str(item).strip().lower()
        if fmt not in _ALLOWED_FORMATS:
            raise ValueError(
                f"unsupported report format {item!r}; expected one of {_ALLOWED_FORMATS}"
            )
        if fmt not in result:
            result.append(fmt)
    if not result:
        raise ValueError("at least one report format is required")
    return tuple(result)


def _check_dependencies(formats: Sequence[str]) -> None:
    if "docx" in formats:
        try:
            for module in (
                "docx", "docx.enum.table", "docx.enum.text", "docx.image.image",
                "docx.oxml", "docx.oxml.ns", "docx.shared",
            ):
                importlib.import_module(module)
        except Exception as exc:  # noqa: BLE001 - normalize broken/partial installs
            raise ReportDependencyError(
                "DOCX generation requires python-docx; install vcstudio[docs]"
            ) from exc
    if "pdf" in formats:
        try:
            reportlab_package = importlib.import_module("reportlab")
            ttfonts_module = None
            for module in (
                "reportlab.lib.colors", "reportlab.lib.enums", "reportlab.lib.pagesizes",
                "reportlab.lib.styles", "reportlab.lib.units", "reportlab.pdfbase.pdfmetrics",
                "reportlab.pdfbase.ttfonts", "reportlab.platypus",
            ):
                imported = importlib.import_module(module)
                if module == "reportlab.pdfbase.ttfonts":
                    ttfonts_module = imported
            latin_fonts = _latin_pdf_font_files(reportlab_package)
            cjk_root = _pdf_font_root()
            font_files = {
                **latin_fonts,
                _PDF_CJK_REGULAR: cjk_root / "noto-sans-sc-400.ttf",
                _PDF_CJK_BOLD: cjk_root / "noto-sans-sc-700.ttf",
            }
            if ttfonts_module is None or not hasattr(ttfonts_module, "TTFont"):
                raise ImportError("reportlab.pdfbase.ttfonts.TTFont is unavailable")
            _probe_pdf_font_files(ttfonts_module.TTFont, font_files)
        except Exception as exc:  # noqa: BLE001 - normalize fonts/partial installs
            raise ReportDependencyError(
                f"PDF generation dependencies are unavailable: {exc}"
            ) from exc


def _probe_pdf_font_files(ttfont_class, font_files: Mapping[str, Path]) -> None:
    """Parse every font required by the PDF renderer without registering it.

    Merely checking that a ``.ttf`` path exists produces a false-positive for a
    truncated or partially installed bundle.  ``TTFont`` parses the sfnt tables
    during construction, which is the same boundary the real renderer crosses.
    Probe names are local to these unregistered objects and therefore cannot
    pollute ReportLab's process-global font registry.
    """

    for index, (logical_name, path) in enumerate(font_files.items(), 1):
        candidate = Path(path)
        if not candidate.is_file():
            raise FileNotFoundError(f"PDF font is missing: {candidate}")
        ttfont_class(f"VCStudioCapabilityProbe{index}_{logical_name}", str(candidate))


def _normalize_model(model: Mapping[str, Any], *, requested_formats=()) -> dict:
    source_schema = _text(model.get("schema")) or _LEGACY_MODEL_SCHEMA
    report_spec = _normalize_full_contract(
        model.get("report_spec"), "report_spec")
    report_snapshot = _normalize_full_contract(
        model.get("report_snapshot"), "report_snapshot")
    validation = _normalize_full_contract(
        model.get("validation"), "validation")
    (contracts_validated, report_spec, report_snapshot,
     validation) = _validate_full_contract_chain(
        report_spec, report_snapshot, validation)
    contract_refs = _normalize_contract_refs(
        model.get("contract_refs"),
        report_spec=report_spec,
        report_snapshot=report_snapshot,
        validation=validation,
    )
    report_kind, qualification = _resolve_scientific_state(
        model, validation, contract_refs,
        contracts_validated=contracts_validated)
    locale = _normalize_locale(model.get("locale"))
    context = _validate_model_contract_context(
        model,
        report_spec=report_spec,
        report_snapshot=report_snapshot,
        validation=validation,
        contract_refs=contract_refs,
        locale=locale,
        requested_formats=requested_formats,
        contracts_validated=contracts_validated,
    )
    content = _normalize_content_fields(
        model, locale=locale, outline=context["outline"])

    return {
        "schema": MODEL_SCHEMA,
        "source_model_schema": source_schema,
        "report_id": _text(model.get("report_id")),
        "report_kind": report_kind,
        "scientific_status": report_kind,
        "scientific_qualification": qualification,
        "claim_ceiling": context["claim_ceiling"],
        "input_fingerprint": context["input_fingerprint"],
        "preset_id": context["preset_id"],
        "template_ref": context["template_ref"],
        "policy_refs": context["policy_refs"],
        "revision": _json_safe(model.get("revision") or {}),
        "contract_status": (
            "bound" if contracts_validated else
            "references_only" if contract_refs else "absent"
        ),
        "contract_refs": contract_refs,
        "report_spec": report_spec,
        "report_snapshot": report_snapshot,
        "validation": validation,
        "claims": context["claims"],
        "claim_graph": _json_safe(model.get("claim_graph") or {}),
        "comparison_context": _json_safe(model.get("comparison_context") or {}),
        "extensions": _json_safe(model.get("extensions") or {}),
        **content,
    }


def _normalize_content_fields(model: Mapping[str, Any], *, locale: str,
                              outline: Sequence[str]) -> dict:
    figures = model.get("figures") or []
    if isinstance(figures, (str, os.PathLike, Mapping)):
        figures = [figures]
    if not isinstance(figures, Sequence):
        raise TypeError("model['figures'] must be a sequence")
    return {
        "locale": locale,
        "outline": [str(item) for item in outline],
        "title": _text(model.get("title")) or "Scientific Report",
        "subtitle": _text(model.get("subtitle")),
        "kicker": _text(model.get("kicker")) or "Research Report",
        "metadata": _normalize_metadata(model.get("metadata")),
        "executive_summary": _normalize_blocks(model.get("executive_summary")),
        "key_findings": _normalize_blocks(model.get("key_findings")),
        "candidate_evaluations": _normalize_table(
            model.get("candidate_evaluations"), "Candidate evaluation"),
        "adsorption_table": _normalize_table(
            model.get("adsorption_table"), "Adsorption-energy results"),
        "comparison_table": _normalize_table(
            model.get("comparison_table"), "Cross-project comparison"),
        "figures": [
            _normalize_figure(item, index) for index, item in enumerate(figures, 1)
        ],
        "methods": _normalize_blocks(model.get("methods")),
        "limitations": _normalize_blocks(model.get("limitations")),
        "recommendations": _normalize_blocks(model.get("recommendations")),
        "comparison_context": _json_safe(model.get("comparison_context") or {}),
    }


def report_content_sha256(model: Mapping[str, Any], *, outline=None) -> str:
    """Hash the exact normalized report content independently of its contracts.

    Callers compute this before constructing ``ValidationResult``.  The renderer
    recomputes the same digest after content-addressing figures and refuses a
    bound report whose visible content changed after validation.
    """
    if not isinstance(model, Mapping):
        raise TypeError("model must be a mapping")
    if outline is None:
        from vcstudio.project.report_contracts import DEFAULT_OUTLINE

        outline = model.get("outline") or DEFAULT_OUTLINE
    locale = _normalize_locale(model.get("locale"))
    normalized = _normalize_content_fields(model, locale=locale, outline=outline)
    normalized.update({
        "report_kind": _normalize_report_kind(model.get("report_kind")),
        # Content hashing is also used before ValidationResult exists.  Validate
        # the vocabulary here, while leaving the final/qualification gate to the
        # bound contract constructor and renderer.
        "scientific_qualification": _normalize_qualification(
            model.get("scientific_qualification"), "diagnostic"),
        "claim_ceiling": _text(model.get("claim_ceiling")),
        "template_ref": _json_safe(model.get("template_ref")),
    })
    prepared_figures = []
    for figure in normalized["figures"]:
        source = Path(figure["_source_path"]).expanduser()
        if not source.is_file():
            raise FileNotFoundError(f"report figure does not exist: {source}")
        prepared = {
            key: value for key, value in figure.items() if not key.startswith("_")
        }
        prepared["asset_sha256"] = _sha256_file(source)
        prepared_figures.append(prepared)
    normalized["figures"] = prepared_figures
    return _sha256_json(_content_fingerprint(normalized))


def frozen_report_content_sha256(model: Mapping[str, Any]) -> str:
    """Recompute visible-content SHA-256 from a published model sidecar.

    Canonical ``MODEL_SCHEMA`` sidecars already contain normalized content and
    content-addressed figure digests, so they must not reopen original figure
    paths.  The legacy/source-model branch is retained for lightweight adapters
    and test doubles that contain no staged figures; it normalizes those fields
    through the same pre-render public helper.
    """
    if not isinstance(model, Mapping):
        raise TypeError("frozen model must be a mapping")
    if _text(model.get("schema")) != MODEL_SCHEMA:
        return report_content_sha256(model, outline=model.get("outline"))
    return _sha256_json(_content_fingerprint(_json_safe(dict(model))))


def _normalize_report_kind(value: Any) -> str:
    """Normalize the scientific label without ever inferring ``final`` from prose."""
    kind = _text(value).lower() or "diagnostic"
    if kind not in _REPORT_KINDS:
        raise ValueError(f"model['report_kind'] must be one of {_REPORT_KINDS}")
    return kind


def _normalize_full_contract(value: Any, field_name: str) -> dict:
    """Normalize one optional full contract without hiding wrong container types."""
    if value is None:
        return {}
    if not isinstance(value, Mapping):
        raise TypeError(f"model[{field_name!r}] must be a mapping")
    return _json_safe(value)


def _normalize_qualification(value: Any, report_kind: str) -> str:
    qualification = _text(value).lower() or "diagnostic"
    if qualification not in _QUALIFICATION_LEVELS:
        raise ValueError(
            "model['scientific_qualification'] must be one of "
            f"{_QUALIFICATION_LEVELS}"
        )
    if report_kind == "final" and qualification == "diagnostic":
        raise ValueError("a final report requires an explicit verified qualification")
    return qualification


def _validate_full_contract_chain(report_spec: Any, report_snapshot: Any,
                                  validation: Any) -> tuple[bool, dict, dict, dict]:
    """Validate and canonicalize one exact full contract chain."""
    contracts = (report_spec, report_snapshot, validation)
    present = tuple(isinstance(item, Mapping) and bool(item) for item in contracts)
    if not any(present):
        return False, {}, {}, {}
    if not all(present):
        raise ValueError(
            "full report contracts must include ReportSpec, ReportSnapshot, "
            "and ValidationResult together"
        )
    schemas = (
        report_spec.get("schema"),
        report_snapshot.get("schema"),
        validation.get("schema"),
    )
    expected = (
        "vcstudio.report-spec/v1",
        "vcstudio.report-snapshot/v1",
        "vcstudio.report-validation/v1",
    )
    if schemas != expected:
        raise ValueError(
            "full report contract schemas do not match the supported v1 chain"
        )
    from vcstudio.project.report_contracts import (
        ReportSnapshot,
        ReportSpec,
        ValidationResult,
        validate_bindings,
    )

    spec_obj = ReportSpec.from_mapping(report_spec)
    snapshot_obj = ReportSnapshot.from_mapping(report_snapshot, spec=spec_obj)
    validation_obj = ValidationResult.from_mapping(
        validation, spec=spec_obj, snapshot=snapshot_obj)
    validate_bindings(spec_obj, snapshot_obj, validation_obj)
    return (
        True,
        spec_obj.to_dict(),
        snapshot_obj.to_dict(),
        validation_obj.to_dict(),
    )


def _validate_model_contract_context(model: Mapping[str, Any], *, report_spec: Any,
                                     report_snapshot: Any, validation: Any,
                                     contract_refs: Any, locale: str,
                                     requested_formats=(),
                                     contracts_validated: bool = False) -> dict:
    """Cross-check render context against the validated report contracts."""
    if isinstance(report_spec, Mapping) and report_spec:
        spec_locale = _text(report_spec.get("locale"))
        if spec_locale and _normalize_locale(spec_locale) != locale:
            raise ValueError("model locale conflicts with ReportSpec.locale")
        spec_formats = report_spec.get("formats")
        if spec_formats is not None and requested_formats:
            if set(_normalize_formats(spec_formats)) != set(requested_formats):
                raise ValueError("requested render formats conflict with ReportSpec.formats")

    snapshot_fingerprint = (
        _text(report_snapshot.get("input_fingerprint"))
        if isinstance(report_snapshot, Mapping) else ""
    )
    snapshot_ref = (contract_refs.get("snapshot")
                    if isinstance(contract_refs, Mapping) else {})
    ref_fingerprint = (_text(snapshot_ref.get("input_fingerprint"))
                       if isinstance(snapshot_ref, Mapping) else "")
    if snapshot_fingerprint and ref_fingerprint and snapshot_fingerprint != ref_fingerprint:
        raise ValueError("snapshot contract reference input_fingerprint mismatch")
    contract_fingerprint = snapshot_fingerprint or ref_fingerprint
    model_fingerprint = _text(model.get("input_fingerprint"))
    if model_fingerprint and contract_fingerprint and model_fingerprint != contract_fingerprint:
        raise ValueError("model input_fingerprint conflicts with ReportSnapshot")

    model_preset = _text(model.get("preset_id"))
    model_ceiling = _text(model.get("claim_ceiling"))
    model_claims = _json_safe(model.get("claims") or [])
    model_template = _json_safe(model.get("template_ref"))
    model_policies = _json_safe(model.get("policy_refs") or [])
    model_outline = list(_SECTION_LABELS[locale])
    if contracts_validated:
        spec_preset = _text(report_spec.get("preset_id"))
        spec_template = _json_safe(report_spec.get("template_ref"))
        spec_policies = _json_safe(report_spec.get("policy_refs") or [])
        spec_outline = [str(item) for item in report_spec.get("outline") or []]
        unsupported_sections = sorted(
            set(spec_outline) - set(_SECTION_LABELS[locale])
        )
        if unsupported_sections:
            raise ValueError(
                "ReportSpec.outline contains unsupported sections: "
                + ", ".join(unsupported_sections)
            )
        validation_ceiling = _text(validation.get("claim_ceiling"))
        validation_claims = _json_safe(validation.get("claims") or [])
        if model_preset and model_preset != spec_preset:
            raise ValueError("model preset_id conflicts with ReportSpec.preset_id")
        if model_ceiling and model_ceiling != validation_ceiling:
            raise ValueError(
                "model claim_ceiling conflicts with ValidationResult.claim_ceiling"
            )
        if "claims" in model and model_claims != validation_claims:
            raise ValueError("model claims conflict with ValidationResult.claims")
        if "template_ref" in model and model_template != spec_template:
            raise ValueError("model template_ref conflicts with ReportSpec.template_ref")
        if "policy_refs" in model and model_policies != spec_policies:
            raise ValueError("model policy_refs conflict with ReportSpec.policy_refs")
        model_preset = spec_preset
        model_template = spec_template
        model_policies = spec_policies
        model_outline = spec_outline
        model_ceiling = validation_ceiling
        model_claims = validation_claims
    return {
        "input_fingerprint": model_fingerprint or contract_fingerprint,
        "preset_id": model_preset,
        "template_ref": model_template,
        "policy_refs": model_policies,
        "outline": model_outline,
        "claim_ceiling": model_ceiling,
        "claims": model_claims,
    }


def _resolve_scientific_state(model: Mapping[str, Any], validation: Any,
                              contract_refs: Mapping[str, Any], *,
                              contracts_validated: bool) -> tuple[str, str]:
    """Resolve one fail-closed science label from validation, never from prose."""
    explicit_kind = _text(model.get("report_kind")).lower()
    explicit_qualification = _text(model.get("scientific_qualification")).lower()
    validation_map = validation if isinstance(validation, Mapping) else {}
    validation_kind = _text(
        validation_map.get("effective_kind") or validation_map.get("report_kind")
    ).lower()
    validation_qualification = _text(
        validation_map.get("scientific_qualification")
        or validation_map.get("qualification")
    ).lower()
    validation_ref = contract_refs.get("validation")
    validation_ref = validation_ref if isinstance(validation_ref, Mapping) else {}
    status = _text(validation_map.get("status") or validation_ref.get("status")).lower()
    final_allowed = (validation_map.get("final_allowed")
                     if validation_map else validation_ref.get("final_allowed"))

    if explicit_kind and validation_kind and explicit_kind != validation_kind:
        raise ValueError(
            "model['report_kind'] conflicts with validation.effective_kind"
        )
    report_kind = _normalize_report_kind(validation_kind or explicit_kind)
    if report_kind == "final":
        permitted = (
            contracts_validated
            and final_allowed is True
            and status in {"passed", "passed_with_warnings"}
        )
        if not permitted:
            if validation_map:
                raise ValueError(
                    "a final report requires a complete bound validation chain "
                    "with final_allowed=true"
                )
            # Legacy content without validation remains renderable, but cannot
            # acquire a final badge or verified qualification.
            report_kind = "diagnostic"
            explicit_qualification = "diagnostic"

    if (explicit_qualification and validation_qualification
            and explicit_qualification != validation_qualification):
        raise ValueError(
            "model['scientific_qualification'] conflicts with validation"
        )
    qualification = _normalize_qualification(
        validation_qualification or explicit_qualification, report_kind)
    if not contracts_validated:
        # A caller-provided label or an opaque hash reference is not scientific
        # evidence.  Unbound reports may retain their artifact kind (draft or
        # diagnostic), but their qualification always fails closed.
        qualification = "diagnostic"
    return report_kind, qualification


def _normalize_contract_refs(value: Any, *, report_spec: Any,
                             report_snapshot: Any, validation: Any) -> dict:
    refs = _json_safe({} if value is None else value)
    if not isinstance(refs, Mapping):
        raise TypeError("model['contract_refs'] must be a mapping")
    unknown = sorted(set(refs) - set(_CONTRACT_SCHEMAS))
    if unknown:
        raise ValueError(f"unsupported contract reference keys: {', '.join(unknown)}")
    result = {}
    full_contracts = {
        "spec": report_spec,
        "snapshot": report_snapshot,
        "validation": validation,
    }
    for key, expected_schema in _CONTRACT_SCHEMAS.items():
        declared = refs.get(key)
        contract = full_contracts[key]
        has_contract = isinstance(contract, Mapping) and bool(contract)
        if declared is None and not has_contract:
            continue
        if declared is not None and not isinstance(declared, Mapping):
            raise TypeError(f"model contract_refs[{key!r}] must be a mapping")
        declared = dict(declared or {})
        if has_contract:
            contract_schema = _text(contract.get("schema"))
            if contract_schema != expected_schema:
                raise ValueError(
                    f"full {key} contract schema must be {expected_schema!r}"
                )
            semantic_sha256 = _contract_semantic_sha256(key, contract)
            canonical = {
                "schema": expected_schema,
                "sha256": semantic_sha256,
            }
            if key == "validation":
                canonical.update({
                    "status": _text(contract.get("status")).lower(),
                    "final_allowed": contract.get("final_allowed"),
                    "report_model_sha256": _text(
                        contract.get("report_model_sha256")
                    ).lower(),
                })
            if key == "snapshot":
                canonical["input_fingerprint"] = _text(
                    contract.get("input_fingerprint"))
            for field, expected_value in canonical.items():
                if field in declared and declared[field] != expected_value:
                    raise ValueError(
                        f"model contract_refs[{key!r}].{field} conflicts with the full contract"
                    )
            result[key] = canonical
            continue

        schema = _text(declared.get("schema"))
        if schema != expected_schema:
            raise ValueError(
                f"contract_refs[{key!r}].schema must be {expected_schema!r}"
            )
        digest = _text(declared.get("sha256")).lower()
        if not _SHA256_HEX_RE.fullmatch(digest):
            raise ValueError(
                f"contract_refs[{key!r}].sha256 must be a 64-character SHA-256"
            )
        canonical = {"schema": expected_schema, "sha256": digest}
        if key == "validation":
            status = _text(declared.get("status")).lower()
            if status not in _VALIDATION_STATUSES:
                raise ValueError("validation contract reference has an invalid status")
            final_allowed = declared.get("final_allowed")
            if not isinstance(final_allowed, bool):
                raise TypeError("validation contract reference final_allowed must be bool")
            canonical.update({"status": status, "final_allowed": final_allowed})
            report_model_sha256 = _text(
                declared.get("report_model_sha256")
            ).lower()
            if report_model_sha256:
                if not _SHA256_HEX_RE.fullmatch(report_model_sha256):
                    raise ValueError(
                        "validation contract reference report_model_sha256 "
                        "must be a 64-character SHA-256"
                    )
                canonical["report_model_sha256"] = report_model_sha256
        if key == "snapshot":
            fingerprint = _text(declared.get("input_fingerprint"))
            if fingerprint:
                canonical["input_fingerprint"] = fingerprint
        result[key] = canonical
    return _json_safe(result)


def _contract_semantic_sha256(key: str, contract: Mapping[str, Any]) -> str:
    """Use the owning contract type's precise semantic projection."""
    from vcstudio.project.report_contracts import (
        ReportSnapshot,
        ReportSpec,
        ValidationResult,
    )

    classes = {
        "spec": ReportSpec,
        "snapshot": ReportSnapshot,
        "validation": ValidationResult,
    }
    return classes[key].from_mapping(contract).semantic_sha256


def _reference_semantic_value(value: Any, *, remove_times: bool = False) -> Any:
    """Strip top-level navigation/admin fields from known reference records only."""
    if value is None:
        return None
    if isinstance(value, Mapping):
        result = _json_safe(value)
        for key in tuple(result):
            lowered = str(key).lower()
            if lowered == "locator" or lowered.endswith("_locator"):
                result.pop(key, None)
            elif remove_times and lowered in _CONTRACT_ADMIN_KEYS:
                result.pop(key, None)
        return result
    if isinstance(value, Sequence) and not isinstance(value, (str, bytes)):
        return [
            _reference_semantic_value(item, remove_times=remove_times)
            for item in value
        ]
    return _json_safe(value)


def _stage_contract_sidecars(model: Mapping[str, Any], temp_root: Path,
                             stem: str) -> tuple[dict[str, Path], dict]:
    """Write portable full-contract sidecars before the transactional commit."""
    outputs: dict[str, Path] = {}
    records = copy.deepcopy(model.get("contract_refs") or {})
    for key, model_key in (
        ("spec", "report_spec"),
        ("snapshot", "report_snapshot"),
        ("validation", "validation"),
    ):
        contract = model.get(model_key)
        if not isinstance(contract, Mapping) or not contract:
            continue
        path = temp_root / f"{stem}.{key}.json"
        path.write_text(
            json.dumps(contract, ensure_ascii=False, indent=2, sort_keys=True,
                       allow_nan=False) + "\n",
            encoding="utf-8",
        )
        file_record = _file_record(path, relative_path=path.name)
        record = dict(records.get(key) or {})
        record.update({
            "path": file_record["path"],
            "file_sha256": file_record["sha256"],
            "size": file_record["size"],
        })
        records[key] = record
        outputs[f"contract_{key}"] = path
    return outputs, records


def _normalize_locale(value: Any) -> str:
    locale = _text(value) or "en-US"
    normalized = locale.replace("_", "-").lower()
    if normalized in {"en", "en-us"}:
        return "en-US"
    if normalized in {"zh", "zh-cn", "zh-hans"}:
        return "zh-CN"
    raise ValueError("model['locale'] must be 'en-US' or 'zh-CN'")


def _report_kind_label(model: Mapping[str, Any]) -> str:
    labels = {
        "zh-CN": {
            "final": "最终报告 · 科学资格已通过",
            "diagnostic": "诊断报告 · 不构成最终科学结论",
            "draft": "报告草稿 · 尚未完成科学审核",
        },
        "en-US": {
            "final": "FINAL REPORT · SCIENTIFIC GATE PASSED",
            "diagnostic": "DIAGNOSTIC REPORT · NOT A FINAL SCIENTIFIC CONCLUSION",
            "draft": "DRAFT REPORT · SCIENTIFIC REVIEW INCOMPLETE",
        },
    }
    return labels[model["locale"]][model["report_kind"]]


def _normalize_metadata(value: Any) -> list[dict]:
    if value is None:
        return []
    if isinstance(value, Mapping):
        return [{"label": _text(key), "value": _display(val)} for key, val in value.items()]
    if isinstance(value, (str, bytes)) or not isinstance(value, Sequence):
        return [{"label": "Metadata", "value": _display(value)}]

    rows = []
    for item in value:
        if isinstance(item, Mapping):
            label = item.get("label", item.get("name", item.get("key", "")))
            val = item.get("value", item.get("text", item.get("detail", "")))
            if not label and len(item) == 1:
                label, val = next(iter(item.items()))
            rows.append({"label": _text(label), "value": _display(val)})
        elif isinstance(item, Sequence) and not isinstance(item, (str, bytes)) and len(item) >= 2:
            rows.append({"label": _text(item[0]), "value": _display(item[1])})
        else:
            rows.append({"label": "Metadata", "value": _display(item)})
    return rows


def _normalize_blocks(value: Any) -> list[dict]:
    if value is None or value == "":
        return []
    if isinstance(value, Mapping):
        if any(key in value for key in ("title", "label", "name", "text", "summary", "detail")):
            return [_block_from_mapping(value)]
        return [
            {"title": _text(key), "text": _display(val)}
            for key, val in value.items()
        ]
    if isinstance(value, (str, bytes)) or not isinstance(value, Sequence):
        return [{"title": "", "text": _display(value)}]

    blocks = []
    for item in value:
        if isinstance(item, Mapping):
            blocks.append(_block_from_mapping(item))
        else:
            blocks.append({"title": "", "text": _display(item)})
    return [block for block in blocks if block["title"] or block["text"]]


def _block_from_mapping(item: Mapping[str, Any]) -> dict:
    title = item.get("title", item.get("label", item.get("name", "")))
    text = item.get(
        "text",
        item.get("summary", item.get("detail", item.get("value", item.get("description", "")))),
    )
    if not text:
        remaining = {
            str(key): value
            for key, value in item.items()
            if key not in {"title", "label", "name"}
        }
        text = remaining
    return {"title": _text(title), "text": _display(text)}


def _normalize_table(value: Any, default_title: str) -> dict | None:
    if value is None or value == [] or value == {}:
        return None
    title = default_title
    caption = ""
    columns = None
    rows = value
    if isinstance(value, Mapping):
        if "rows" in value:
            title = _text(value.get("title")) or default_title
            caption = _text(value.get("caption"))
            columns = value.get("columns")
            rows = value.get("rows") or []
        else:
            rows = [{"field": key, "value": val} for key, val in value.items()]
            columns = [
                {"key": "field", "label": "Field"},
                {"key": "value", "label": "Value"},
            ]
    if isinstance(rows, (str, bytes)) or not isinstance(rows, Sequence):
        rows = [{"value": rows}]
    rows = list(rows)
    if not rows:
        return None

    definitions = _column_definitions(columns, rows)
    normalized_rows = []
    for row in rows:
        if isinstance(row, Mapping):
            normalized_rows.append([_display(row.get(col["key"], "")) for col in definitions])
        elif isinstance(row, Sequence) and not isinstance(row, (str, bytes)):
            normalized_rows.append(
                [_display(row[index]) if index < len(row) else "" for index in range(len(definitions))]
            )
        else:
            normalized_rows.append([_display(row)] + [""] * (len(definitions) - 1))
    return {
        "title": title,
        "caption": caption,
        "columns": definitions,
        "rows": normalized_rows,
    }


def _column_definitions(columns: Any, rows: Sequence[Any]) -> list[dict]:
    if columns:
        result = []
        for index, column in enumerate(columns):
            if isinstance(column, Mapping):
                key = _text(column.get("key", column.get("field", index)))
                label = _text(column.get("label", column.get("title", key))) or f"Column {index + 1}"
                align = _text(column.get("align")).lower()
            else:
                key = _text(column)
                label = key or f"Column {index + 1}"
                align = ""
            result.append({"key": key, "label": label, "align": _valid_align(align)})
        return result

    first_mapping = next((row for row in rows if isinstance(row, Mapping)), None)
    if first_mapping is not None:
        keys: list[str] = []
        for row in rows:
            if isinstance(row, Mapping):
                for key in row:
                    text_key = _text(key)
                    if text_key not in keys:
                        keys.append(text_key)
        return [{"key": key, "label": _human_label(key), "align": "auto"} for key in keys]
    first_sequence = next(
        (
            row
            for row in rows
            if isinstance(row, Sequence) and not isinstance(row, (str, bytes))
        ),
        None,
    )
    count = len(first_sequence) if first_sequence is not None else 1
    return [
        {"key": str(index), "label": f"Column {index + 1}", "align": "auto"}
        for index in range(count)
    ]


def _normalize_figure(item: Any, index: int) -> dict:
    if isinstance(item, (str, os.PathLike)):
        source = Path(item).expanduser()
        return {
            "title": f"Figure {index}",
            "caption": "",
            "alt": "",
            "_source_path": source,
        }
    if not isinstance(item, Mapping):
        raise TypeError(f"figure {index} must be a path or mapping")
    raw_path = next((item.get(key) for key in _FIGURE_PATH_KEYS if item.get(key)), None)
    if raw_path is None:
        raise ValueError(f"figure {index} has no path")
    title = _text(item.get("title", item.get("name"))) or f"Figure {index}"
    return {
        "title": title,
        "caption": _text(item.get("caption")),
        # Do not manufacture alternative text from a display title.  Renderers
        # still emit their native description attribute, but an absent source
        # description remains empty and the accessibility record fails closed.
        "alt": _text(item.get("alt")),
        "figure_id": _text(item.get("figure_id")),
        "evidence_refs": _json_safe(item.get("evidence_refs") or []),
        "claim_refs": _json_safe(item.get("claim_refs") or []),
        "source_ids": _json_safe(item.get("source_ids") or []),
        "data_sha256": _text(item.get("data_sha256")),
        "quantity": _text(item.get("quantity")),
        "unit": _text(item.get("unit")),
        "denominator": _json_safe(item.get("denominator")),
        "extensions": _json_safe(item.get("extensions") or {}),
        "_source_path": Path(raw_path).expanduser(),
    }


def _has_meaningful_figure_alt(figure: Mapping[str, Any]) -> bool:
    """Return whether a figure description clears the fail-closed baseline.

    This is intentionally a narrow mechanical gate, not an editorial quality
    claim.  Empty text and generated labels such as ``Figure 1`` are never
    accepted as meaningful alternative text.  A human still needs to confirm
    that non-generic descriptions convey the figure's scientific purpose.
    """

    value = " ".join(_text(figure.get("alt")).split())
    return bool(value) and _GENERIC_FIGURE_ALT_RE.fullmatch(value) is None


def _report_accessibility(model: Mapping[str, Any], formats: Sequence[str]) -> dict:
    """Build format-specific, fail-closed accessibility evidence.

    Records describe only what each renderer encodes.  They do not promote a
    generated artifact to WCAG, Word-checker, or PDF/UA conformance.  Missing or
    generic figure descriptions downgrade HTML/DOCX to ``partial``; ReportLab
    PDF remains partial regardless because it does not encode a structure tree
    or image alternative text.
    """

    figures = list(model.get("figures") or [])
    missing_alt = sum(
        1 for figure in figures
        if not isinstance(figure, Mapping) or not _has_meaningful_figure_alt(figure)
    )
    records = {}
    for fmt in formats:
        record = copy.deepcopy(_FORMAT_ACCESSIBILITY[fmt])
        record["figures_total"] = len(figures)
        record["figures_missing_meaningful_alt"] = missing_alt
        record["source_figure_alt_complete"] = missing_alt == 0
        if fmt in {"html", "docx"} and missing_alt:
            record["status"] = "partial"
            prefix_zh = f"{missing_alt} 幅图片缺少有意义的替代文本；"
            prefix_en = (
                f"{missing_alt} figure(s) lack meaningful alternative text; "
            )
            record["reason_zh"] = prefix_zh + record["reason_zh"]
            record["reason_en"] = prefix_en + record["reason_en"]
            record["reason"] = record["reason_zh"]
        records[fmt] = record
    return records


def _stage_figures(model: dict, assets_dir: Path, formats: Sequence[str]) -> tuple[dict, list[dict]]:
    result = dict(model)
    result["figures"] = []
    records_by_digest: dict[str, dict] = {}
    require_raster = "docx" in formats or "pdf" in formats

    for index, figure in enumerate(model["figures"], 1):
        source = figure["_source_path"]
        if not source.is_file():
            raise FileNotFoundError(f"figure {index} does not exist or is not a file: {source}")
        suffix = source.suffix.lower()
        if require_raster and suffix not in _RASTER_SUFFIXES:
            raise ValueError(
                f"figure {index} uses {suffix or 'an unknown format'}; "
                "DOCX/PDF reports require a raster image"
            )
        digest = _sha256_file(source)
        record = records_by_digest.get(digest)
        if record is None:
            assets_dir.mkdir(parents=True, exist_ok=True)
            safe_name = _safe_asset_stem(source.stem)
            asset_name = f"{safe_name}-{digest[:20]}{suffix or '.img'}"
            target = assets_dir / asset_name
            shutil.copyfile(source, target)
            if _sha256_file(target) != digest:
                raise OSError(f"figure copy verification failed: {source}")
            record = {
                "path": f"assets/{asset_name}",
                "sha256": digest,
                "size": target.stat().st_size,
                "media_type": mimetypes.guess_type(asset_name)[0] or "application/octet-stream",
            }
            records_by_digest[digest] = record

        staged = {key: value for key, value in figure.items() if not key.startswith("_")}
        staged.update(
            {
                "asset_path": record["path"],
                "asset_name": Path(record["path"]).name,
                "asset_sha256": digest,
                "_render_path": assets_dir / Path(record["path"]).name,
            }
        )
        result["figures"].append(staged)
    return result, list(records_by_digest.values())


def _inline_preview_figures(model: dict) -> dict:
    """Return an HTML-ready model whose figures are content-addressed data URIs."""
    result = dict(model)
    result["figures"] = []
    encoded_by_digest: dict[str, tuple[str, str]] = {}

    for index, figure in enumerate(model["figures"], 1):
        source = figure["_source_path"]
        if not source.is_file():
            # Do not disclose a local source locator through a preview error.
            raise FileNotFoundError(
                f"figure {index} does not exist or is not a regular file"
            )
        payload = source.read_bytes()
        digest = hashlib.sha256(payload).hexdigest()
        encoded = encoded_by_digest.get(digest)
        if encoded is None:
            media_type = mimetypes.guess_type(source.name)[0]
            if not media_type or not media_type.startswith("image/"):
                raise ValueError(
                    f"figure {index} has no browser-safe image media type"
                )
            data_uri = (
                f"data:{media_type};base64,"
                + base64.b64encode(payload).decode("ascii")
            )
            encoded = (data_uri, media_type)
            encoded_by_digest[digest] = encoded

        staged = {
            key: value for key, value in figure.items() if not key.startswith("_")
        }
        staged.update({
            "asset_path": encoded[0],
            "asset_sha256": digest,
        })
        result["figures"].append(staged)
    return result


def _commit_assets(source_dir: Path, destination_dir: Path, records: Sequence[dict]) -> list[Path]:
    if not records:
        return []
    destination_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for record in records:
        name = Path(record["path"]).name
        source = source_dir / name
        destination = destination_dir / name
        if destination.exists():
            if not destination.is_file() or _sha256_file(destination) != record["sha256"]:
                raise OSError(f"asset hash collision at destination: {destination}")
            source.unlink()
        else:
            os.replace(source, destination)
        paths.append(destination)
    return paths


def _commit_report_files(
    temp_outputs: Mapping[str, Path],
    temp_manifest: Path,
    destination: Path,
    rollback_dir: Path,
    *,
    obsolete_paths: Sequence[Path] = (),
) -> tuple[dict[str, Path], Path]:
    """Publish a report set with best-effort transactional rollback.

    Windows refuses to replace a DOCX/PDF while Word or a PDF viewer holds a
    restrictive file lock.  Back up every existing destination *before* the first
    replacement so a later failure cannot leave a new HTML file beside an old
    DOCX/PDF and stale manifest.  The manifest is always the final commit marker.
    """
    entries = [
        (fmt, source, destination / source.name)
        for fmt, source in temp_outputs.items()
    ]
    final_manifest = destination / temp_manifest.name
    obsolete = tuple(
        path for path in dict.fromkeys(Path(item) for item in obsolete_paths)
        if path not in {target for _kind, _source, target in entries}
    )
    commit_entries = [*entries, ("manifest", temp_manifest, final_manifest)]

    rollback_dir.mkdir(parents=True, exist_ok=True)
    backups: dict[Path, Path] = {}
    for target in dict.fromkeys([
            *(target for _kind, _source, target in commit_entries),
            *obsolete,
    ]):
        if not target.exists():
            continue
        if not target.is_file():
            raise OSError(f"report destination is not a file: {target}")
        backup = rollback_dir / target.name
        shutil.copyfile(target, backup)
        if _sha256_file(backup) != _sha256_file(target):
            raise OSError(f"report rollback copy verification failed: {target}")
        backups[target] = backup

    published: list[Path] = []
    deleted: list[Path] = []
    manifest_invalidated = False
    try:
        # Withdraw the old commit marker before touching any member file.  If a
        # viewer holds it with a restrictive Windows lock, publication aborts
        # here while every old artifact is still untouched.  If a later step
        # fails, the old marker is restored only after all members roll back.
        if final_manifest.exists():
            if final_manifest not in backups:
                raise OSError(f"report manifest backup is unavailable: {final_manifest}")
            final_manifest.unlink()
            manifest_invalidated = True
        for _kind, source, target in entries:
            os.replace(source, target)
            published.append(target)
        for target in obsolete:
            if target.exists():
                target.unlink()
                deleted.append(target)
        os.replace(temp_manifest, final_manifest)
        published.append(final_manifest)
    except Exception as publish_error:
        rollback_errors: list[str] = []
        for target in reversed(published):
            try:
                backup = backups.get(target)
                if backup is None:
                    target.unlink(missing_ok=True)
                else:
                    os.replace(backup, target)
            except Exception as rollback_error:  # noqa: BLE001
                rollback_errors.append(f"{target}: {rollback_error}")
        for target in reversed(deleted):
            try:
                backup = backups.get(target)
                if backup is not None:
                    os.replace(backup, target)
            except Exception as rollback_error:  # noqa: BLE001
                rollback_errors.append(f"{target}: {rollback_error}")
        if manifest_invalidated and not rollback_errors:
            try:
                os.replace(backups[final_manifest], final_manifest)
            except Exception as rollback_error:  # noqa: BLE001
                rollback_errors.append(f"{final_manifest}: {rollback_error}")
        if rollback_errors:
            backup_records = []
            recovery_record_errors = []
            for target, backup in backups.items():
                if not backup.exists():
                    continue
                record = {"target": str(target), "backup": str(backup)}
                try:
                    record.update({
                        "sha256": _sha256_file(backup),
                        "size": backup.stat().st_size,
                    })
                except Exception as evidence_error:  # noqa: BLE001
                    recovery_record_errors.append(
                        f"could not fingerprint {backup}: {evidence_error}"
                    )
                backup_records.append(record)
            recovery_record = {
                "schema": "vcstudio.paper-report.recovery/v1",
                "artifact_status": "recovery_required",
                "created_at_utc": datetime.now(timezone.utc).replace(
                    microsecond=0).isoformat(),
                "destination": str(destination),
                "manifest_path": str(final_manifest),
                "canonical_manifest_withheld": bool(manifest_invalidated),
                "publication_error": str(publish_error),
                "rollback_errors": rollback_errors,
                "recovery_record_errors": recovery_record_errors,
                "backups": backup_records,
            }
            recovery_record_path = rollback_dir / "RECOVERY.json"
            record_error = None
            try:
                recovery_record_path.write_text(
                    json.dumps(
                        recovery_record,
                        ensure_ascii=False,
                        indent=2,
                        sort_keys=True,
                        allow_nan=False,
                    ) + "\n",
                    encoding="utf-8",
                )
            except Exception as recovery_error:  # noqa: BLE001
                record_error = str(recovery_error)
            detail = "; ".join(rollback_errors)
            message = (
                "report publication failed and rollback was incomplete; "
                f"recovery evidence retained at {rollback_dir}: {detail}"
            )
            if record_error:
                message += f"; RECOVERY.json could not be written: {record_error}"
            raise ReportRecoveryError(
                message, recovery_dir=rollback_dir
            ) from publish_error
        raise

    output_paths = {
        fmt: destination / source.name
        for fmt, source in temp_outputs.items()
    }
    output_paths["manifest"] = final_manifest
    return output_paths, final_manifest


def _fingerprint_model(model: dict) -> dict:
    payload = {}
    for key, value in model.items():
        if key in {"report_spec", "report_snapshot", "validation"}:
            # Full contracts are retained for audit, while their semantic
            # digests in contract_refs bind the model without timestamps or
            # machine-specific locators polluting reproducibility.
            continue
        if key == "revision":
            payload[key] = _reference_semantic_value(value, remove_times=True)
            continue
        if key in {"template_ref", "policy_refs"}:
            payload[key] = _reference_semantic_value(value)
            continue
        if key in {"extensions", "contract_refs", "claims", "claim_graph",
                   "comparison_context"}:
            payload[key] = _json_safe(value)
            continue
        if key != "figures":
            payload[key] = value
            continue
        payload["figures"] = [
            _json_safe({
                key: value
                for key, value in figure.items()
                if not key.startswith("_") and key not in {"asset_path", "asset_name"}
            } | {
                "asset_sha256": figure["asset_sha256"],
            })
            for figure in value
        ]
    return _json_safe(payload)


def _content_fingerprint(model: Mapping[str, Any]) -> dict:
    """Return only visible/scientific report content for validation binding."""
    payload = {}
    for key in _CONTENT_FINGERPRINT_KEYS:
        value = model.get(key)
        if key == "template_ref":
            payload[key] = _reference_semantic_value(value)
            continue
        if key != "figures":
            payload[key] = _json_safe(value)
            continue
        figures = []
        for figure in value or []:
            digest = _text(figure.get("asset_sha256")).lower()
            if not _SHA256_HEX_RE.fullmatch(digest):
                raise ValueError("report figure is missing a content SHA-256")
            figures.append(_json_safe({
                item_key: item_value
                for item_key, item_value in figure.items()
                if not str(item_key).startswith("_")
                and item_key not in {"asset_path", "asset_name"}
            } | {"asset_sha256": digest}))
        payload[key] = figures
    return _json_safe(payload)


def _render_html(model: dict, output: Path) -> None:
    output.write_text(_html_document(model), encoding="utf-8")


def _report_theme_id(model: Mapping[str, Any]) -> str:
    reference = model.get("template_ref")
    raw = _text(reference.get("id")) if isinstance(reference, Mapping) else ""
    if raw.startswith("vcstudio-"):
        raw = raw[len("vcstudio-"):]
    return raw if raw in _REPORT_THEMES else "academic-a4"


def _report_theme(model: Mapping[str, Any]) -> dict[str, Any]:
    return dict(_REPORT_THEMES[_report_theme_id(model)])


def _hex_rgb(value: str) -> tuple[int, int, int]:
    text = str(value or "").lstrip("#")
    if not re.fullmatch(r"[0-9A-Fa-f]{6}", text):
        raise ValueError(f"invalid report theme color: {value!r}")
    return tuple(int(text[index:index + 2], 16) for index in (0, 2, 4))


def _html_document(model: dict) -> str:
    theme_id = _report_theme_id(model)
    theme = _report_theme(model)
    sections = _section_plan(model)
    metadata = "".join(
        "<div class='meta-row'><dt>{}</dt><dd>{}</dd></div>".format(
            html.escape(row["label"]),
            html.escape(row["value"]),
        )
        for row in model["metadata"]
    )
    body = []
    table_number = 0
    figure_number = 0
    for section_number, (key, heading) in enumerate(sections, 1):
        body.append(f"<section><h2><span>{section_number:02d}</span>{html.escape(heading)}</h2>")
        if key == "executive_summary":
            body.append(_html_blocks(model[key], ordered=False))
        elif key == "key_findings":
            body.append(_html_blocks(model[key], ordered=True))
        elif key in {"candidate_evaluations", "adsorption_table", "comparison_table"}:
            table_number += 1
            body.append(_html_table(model[key], table_number, model["locale"]))
        elif key == "figures":
            for figure in model["figures"]:
                figure_number += 1
                caption = _caption_text(figure["title"], figure["caption"])
                body.append(
                    "<figure>"
                    f"<img src='{html.escape(figure['asset_path'], quote=True)}' "
                    f"alt='{html.escape(figure['alt'], quote=True)}'>"
                    f"<figcaption>{html.escape(_figure_caption(model['locale'], figure_number, caption))}"
                    "</figcaption>"
                    "</figure>"
                )
        else:
            body.append(_html_blocks(model[key], ordered=key == "recommendations"))
        body.append("</section>")

    title = html.escape(model["title"])
    subtitle = html.escape(model["subtitle"])
    kicker = html.escape(model["kicker"])
    meta_description = html.escape(
        model["subtitle"] or model["title"], quote=True)
    report_kind = html.escape(model["report_kind"], quote=True)
    report_status = html.escape(_report_kind_label(model))
    document = f"""<!doctype html>
<html lang="{model['locale']}">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<meta name="author" content="VASP Catalyst Studio">
<meta name="description" content="{meta_description}">
<title>{title}</title>
<style>
@page {{
  size: A4 portrait;
  margin: 22mm {theme['page_margin_mm']}mm 20mm;
  @top-left {{ content: "{_css_string(model['kicker'])}"; color: #6b7280; font-size: 8pt; }}
  @bottom-right {{ content: {_page_counter_css(model['locale'])}; color: #6b7280; font-size: 8pt; }}
}}
:root {{ --ink:{theme['ink']}; --muted:{theme['muted']};
  --accent:{theme['accent']}; --highlight:{theme['highlight']};
  --rule:{theme['rule']}; --title:{theme['title']};
  --subtitle:{theme['subtitle']}; --caption:{theme['caption']}; }}
* {{ box-sizing:border-box; }}
body {{
  margin:0; color:var(--ink); background:{theme['screen']};
  font-family:"Noto Serif CJK SC","Songti SC",SimSun,"Times New Roman",serif;
  font-size:10.5pt; line-height:1.68;
}}
.paper {{ width:210mm; min-height:297mm; margin:18px auto; background:#fff;
  padding:22mm {theme['page_margin_mm']}mm 20mm; box-shadow:0 8px 30px rgba(30,41,59,.12); }}
.cover {{ min-height:248mm; display:flex; flex-direction:column; justify-content:center;
  text-align:center; page-break-after:always; }}
.kicker {{ margin:0 0 16mm; color:var(--highlight); font:700 9pt/1.2 Arial,sans-serif;
  letter-spacing:.18em; text-transform:uppercase; }}
.report-status {{ align-self:center; margin:0 0 6mm; padding:2.2mm 5mm;
  border:1px solid #9ca3af; border-radius:999px; color:#4b5563;
  font:700 8.5pt/1.2 Arial,"Microsoft YaHei",sans-serif; letter-spacing:.04em; }}
.report-status[data-kind="final"] {{ border-color:#15803d; background:#f0fdf4; color:#166534; }}
.report-status[data-kind="diagnostic"] {{ border-color:#b45309; background:#fffbeb; color:#92400e; }}
.report-status[data-kind="draft"] {{ border-color:#64748b; background:#f8fafc; color:#475569; }}
h1 {{ margin:0; color:var(--title); font-size:29pt; line-height:1.22; font-weight:600; }}
.subtitle {{ margin:5mm auto 0; max-width:135mm; color:var(--subtitle); font-size:14pt; line-height:1.5; }}
.metadata {{ width:min(125mm,100%); margin:25mm auto 0; padding-top:6mm;
  border-top:.8pt solid #b7c0c7; }}
.meta-row {{ display:grid; grid-template-columns:35mm 1fr; gap:5mm; text-align:left;
  padding:1.2mm 0; }}
.meta-row dt {{ color:var(--muted); font:700 8.5pt/1.5 Arial,sans-serif;
  text-transform:uppercase; letter-spacing:.05em; }}
.meta-row dd {{ margin:0; }}
section {{ margin:0 0 10mm; break-inside:auto; }}
h2 {{ display:flex; align-items:baseline; gap:4mm; margin:12mm 0 4mm;
  color:var(--title); font-size:16pt; line-height:1.25; font-weight:600;
  border-bottom:.6pt solid #c4ccd1; padding-bottom:2mm; }}
h2 span {{ color:var(--highlight); font:700 8.5pt/1 Arial,sans-serif; letter-spacing:.08em; }}
h3 {{ margin:5mm 0 1.5mm; color:var(--accent); font-size:11.5pt; }}
p {{ margin:0 0 3.2mm; text-align:justify; }}
ol.findings {{ margin:0; padding-left:7mm; }}
ol.findings li {{ margin:0 0 3mm; padding-left:2mm; }}
ol.findings li strong {{ color:var(--accent); }}
table.three-line {{ width:100%; border-collapse:collapse; table-layout:auto;
  margin:1mm 0 2mm; border-top:1.4pt solid var(--rule); border-bottom:1.4pt solid var(--rule);
  font-size:9.2pt; }}
table.three-line caption {{ caption-side:top; text-align:left; margin:0 0 2mm;
  font-size:9pt; font-weight:600; color:#303c44; }}
table.three-line thead {{ border-bottom:.8pt solid var(--rule); }}
table.three-line th, table.three-line td {{ border:0; padding:2mm 2.2mm;
  vertical-align:middle; overflow-wrap:anywhere; }}
table.three-line th {{ font-family:Arial,sans-serif; font-size:8.5pt; text-align:left; }}
table.three-line td.numeric, table.three-line th.numeric {{ text-align:right; }}
.table-note {{ margin:1mm 0 0; color:var(--muted); font-size:8.5pt; text-align:left; }}
figure {{ margin:6mm auto 9mm; break-inside:avoid; text-align:center; }}
figure img {{ display:block; max-width:100%; max-height:150mm; margin:0 auto; object-fit:contain; }}
figcaption {{ margin-top:2.5mm; color:var(--caption); font-size:9pt; line-height:1.45; text-align:center; }}
.screen-running {{ display:none; }}
@media print {{
  body {{ background:#fff; }}
  .paper {{ width:auto; min-height:auto; margin:0; padding:0; box-shadow:none; }}
}}
@media screen and (max-width:800px) {{
  .paper {{ width:100%; min-height:0; margin:0; padding:18mm 8vw; box-shadow:none; }}
  .cover {{ min-height:80vh; }}
}}
</style>
</head>
<body data-report-theme="{theme_id}">
<main class="paper">
<header class="cover">
  <p class="kicker">{kicker}</p>
  <p class="report-status" data-kind="{report_kind}">{report_status}</p>
  <h1>{title}</h1>
  {f'<p class="subtitle">{subtitle}</p>' if subtitle else ''}
  {f'<dl class="metadata">{metadata}</dl>' if metadata else ''}
</header>
{''.join(body)}
</main>
</body>
</html>
"""
    return document


def _html_blocks(blocks: Sequence[dict], *, ordered: bool) -> str:
    if ordered:
        items = []
        for block in blocks:
            title = f"<strong>{html.escape(block['title'])}.</strong> " if block["title"] else ""
            items.append(f"<li>{title}{html.escape(block['text'])}</li>")
        return f"<ol class='findings'>{''.join(items)}</ol>"
    result = []
    for block in blocks:
        if block["title"]:
            result.append(f"<h3>{html.escape(block['title'])}</h3>")
        result.extend(f"<p>{html.escape(part)}</p>" for part in _paragraphs(block["text"]))
    return "".join(result)


def _html_table(table: dict, number: int, locale: str = "en-US") -> str:
    headers = "".join(
        f"<th class='{_html_align_class(column, table['rows'], index)}'>"
        f"{html.escape(column['label'])}</th>"
        for index, column in enumerate(table["columns"])
    )
    rows = []
    for row in table["rows"]:
        cells = "".join(
            f"<td class='{_html_align_class(table['columns'][index], table['rows'], index)}'>"
            f"{html.escape(value)}</td>"
            for index, value in enumerate(row)
        )
        rows.append(f"<tr>{cells}</tr>")
    caption = html.escape(_table_caption(locale, number, table["title"]))
    note = (
        f"<p class='table-note'>{html.escape(table['caption'])}</p>"
        if table["caption"]
        else ""
    )
    return (
        f"<table class='three-line'><caption>{caption}</caption>"
        f"<thead><tr>{headers}</tr></thead><tbody>{''.join(rows)}</tbody></table>{note}"
    )


def _render_docx(model: dict, output: Path) -> None:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.image.image import Image as DocxImage
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, Pt, RGBColor

    theme = _report_theme(model)
    title_color = _hex_rgb(theme["title"])
    accent_color = _hex_rgb(theme["accent"])
    muted_color = _hex_rgb(theme["muted"])
    caption_color = _hex_rgb(theme["caption"])
    highlight_color = _hex_rgb(theme["highlight"])
    ink_color = _hex_rgb(theme["ink"])
    doc = Document()
    _set_docx_document_language(doc, model["locale"], qn, OxmlElement)
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(22)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(theme["page_margin_mm"])
    section.right_margin = Mm(theme["page_margin_mm"])
    section.header_distance = Mm(10)
    section.footer_distance = Mm(10)

    def set_font(run, *, name="Times New Roman", east_asia="SimSun", size=None,
                 bold=None, italic=None, color=None):
        run.font.name = name
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
        if size is not None:
            run.font.size = Pt(size)
        if bold is not None:
            run.bold = bold
        if italic is not None:
            run.italic = italic
        if color is not None:
            run.font.color.rgb = RGBColor(*color)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, before, after, color in (
        ("Heading 1", 16, 18, 8, title_color),
        ("Heading 2", 13, 12, 6, accent_color),
        ("Heading 3", 11.5, 8, 4, accent_color),
    ):
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(*color)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "SimSun")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    caption_style = styles["Caption"]
    caption_style.font.name = "Times New Roman"
    caption_style.font.size = Pt(9)
    caption_style.font.italic = False
    caption_style.font.color.rgb = RGBColor(*caption_color)
    caption_style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "SimSun")
    caption_style.paragraph_format.space_before = Pt(4)
    caption_style.paragraph_format.space_after = Pt(5)
    caption_style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run(f"{model['kicker'].upper()}  ·  {model['title']}")
    set_font(run, name="Arial", east_asia="Microsoft YaHei", size=8, color=muted_color)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    page_prefix = "第 " if model["locale"] == "zh-CN" else "Page "
    run = footer.add_run(page_prefix)
    set_font(run, name="Arial", east_asia="Microsoft YaHei", size=8, color=muted_color)
    _append_word_field(run, "PAGE", OxmlElement)
    if model["locale"] == "zh-CN":
        suffix = footer.add_run(" 页")
        set_font(
            suffix,
            name="Arial",
            east_asia="Microsoft YaHei",
            size=8,
            color=muted_color,
        )

    doc.core_properties.title = model["title"]
    doc.core_properties.subject = model["subtitle"]
    doc.core_properties.author = "VASP Catalyst Studio"
    doc.core_properties.keywords = "scientific report; adsorption energy; catalysis"
    doc.core_properties.language = model["locale"]

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(theme["cover_spacer_mm"] * 2.2)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(20)
    set_font(
        kicker.add_run(model["kicker"].upper()),
        name="Arial",
        east_asia="Microsoft YaHei",
        size=9,
        bold=True,
        color=highlight_color,
    )
    status = doc.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status.paragraph_format.space_after = Pt(16)
    status_color = {
        "final": (22, 101, 52),
        "diagnostic": (146, 64, 14),
        "draft": (71, 85, 105),
    }[model["report_kind"]]
    set_font(
        status.add_run(_report_kind_label(model)),
        name="Arial",
        east_asia="Microsoft YaHei",
        size=9,
        bold=True,
        color=status_color,
    )
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    set_font(title.add_run(model["title"]), size=29, bold=True, color=title_color)
    if model["subtitle"]:
        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.space_after = Pt(28)
        set_font(subtitle.add_run(model["subtitle"]), size=14, color=_hex_rgb(theme["subtitle"]))
    for item in model["metadata"]:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(2)
        set_font(
            paragraph.add_run(f"{item['label'].upper()}  "),
            name="Arial",
            east_asia="Microsoft YaHei",
            size=8,
            bold=True,
            color=muted_color,
        )
        set_font(paragraph.add_run(item["value"]), size=9.5, color=ink_color)
    doc.add_page_break()

    sections = _section_plan(model)
    table_number = 0
    figure_number = 0
    for section_number, (key, heading) in enumerate(sections, 1):
        doc.add_heading(f"{section_number}. {heading}", level=1)
        if key == "executive_summary":
            _docx_add_blocks(doc, model[key], set_font, ordered=False)
        elif key == "key_findings":
            _docx_add_blocks(doc, model[key], set_font, ordered=True)
        elif key in {"candidate_evaluations", "adsorption_table", "comparison_table"}:
            table_number += 1
            _docx_add_table(
                doc,
                model[key],
                table_number,
                model["locale"],
                set_font,
                qn,
                OxmlElement,
                WD_ALIGN_PARAGRAPH,
                WD_CELL_VERTICAL_ALIGNMENT,
                theme["rule"].lstrip("#"),
            )
        elif key == "figures":
            for figure in model["figures"]:
                figure_number += 1
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(5)
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.keep_with_next = True
                image = DocxImage.from_file(str(figure["_render_path"]))
                max_width = Inches(6.25)
                max_height = Inches(7.4)
                width, height = image.scaled_dimensions(width=max_width)
                if height > max_height:
                    width, height = image.scaled_dimensions(height=max_height)
                inline_shape = paragraph.add_run().add_picture(
                    str(figure["_render_path"]),
                    width=width,
                    height=height,
                )
                _set_docx_picture_accessibility(
                    inline_shape,
                    alt=figure["alt"],
                    title=figure["title"],
                )
                caption = doc.add_paragraph(style="Caption")
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.paragraph_format.keep_with_next = False
                caption.add_run(
                    _figure_caption(
                        model["locale"],
                        figure_number,
                        _caption_text(figure["title"], figure["caption"]),
                    )
                )
        else:
            _docx_add_blocks(
                doc,
                model[key],
                set_font,
                ordered=key == "recommendations",
            )

    # Save from the temporary directory; the caller commits only complete files.
    doc.save(output)


def _set_docx_document_language(doc, locale, qn, element_factory) -> None:
    """Set Word document defaults and report styles to one BCP-47 locale."""

    def set_language(run_properties) -> None:
        language = run_properties.find(qn("w:lang"))
        if language is None:
            language = element_factory("w:lang")
            run_properties.append(language)
        language.set(qn("w:val"), locale)
        language.set(qn("w:eastAsia"), locale)

    styles = doc.styles.element
    defaults = styles.find(qn("w:docDefaults"))
    if defaults is None:
        defaults = element_factory("w:docDefaults")
        styles.insert(0, defaults)
    run_defaults = defaults.find(qn("w:rPrDefault"))
    if run_defaults is None:
        run_defaults = element_factory("w:rPrDefault")
        defaults.append(run_defaults)
    run_properties = run_defaults.find(qn("w:rPr"))
    if run_properties is None:
        run_properties = element_factory("w:rPr")
        run_defaults.append(run_properties)
    set_language(run_properties)

    for style_name in (
        "Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3",
        "Caption",
    ):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        set_language(style._element.get_or_add_rPr())


def _set_docx_picture_accessibility(inline_shape, *, alt, title) -> None:
    """Write image description/title to the standard ``wp:docPr`` element."""

    properties = inline_shape._inline.docPr
    properties.set("descr", _text(alt))
    display_title = _text(title)
    if not display_title or _GENERIC_FIGURE_ALT_RE.fullmatch(display_title):
        display_title = _text(alt)
    if display_title:
        properties.set("title", display_title)


def _append_word_field(run, instruction: str, element_factory) -> None:
    from docx.oxml.ns import qn

    begin = element_factory("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = element_factory("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = element_factory("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = element_factory("w:t")
    text.text = "1"
    end = element_factory("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))


def _docx_add_blocks(doc, blocks, set_font, *, ordered: bool) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    numbering_id = _new_decimal_numbering(doc) if ordered else None
    for block in blocks:
        if ordered:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(7)
            _apply_numbering(paragraph, numbering_id)
            if block["title"]:
                set_font(paragraph.add_run(f"{block['title']}. "), size=10.5, bold=True)
            set_font(paragraph.add_run(block["text"]), size=10.5)
        else:
            if block["title"]:
                doc.add_heading(block["title"], level=2)
            for text in _paragraphs(block["text"]):
                paragraph = doc.add_paragraph(text)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _new_decimal_numbering(doc) -> int:
    """Create one real, restartable Word decimal-list definition."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    root = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in root.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in root.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("w:start", "1"), ("w:numFmt", "decimal"), ("w:lvlText", "%1.")):
        node = OxmlElement(tag)
        node.set(qn("w:val"), value)
        level.append(node)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    paragraph_properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    paragraph_properties.append(tabs)
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), "540")
    indentation.set(qn("w:hanging"), "270")
    paragraph_properties.append(indentation)
    level.append(paragraph_properties)
    abstract.append(level)
    root.append(abstract)

    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), str(abstract_id))
    number.append(reference)
    root.append(number)
    return num_id


def _apply_numbering(paragraph, num_id: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = paragraph._p.get_or_add_pPr()
    numbering = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    reference = OxmlElement("w:numId")
    reference.set(qn("w:val"), str(num_id))
    numbering.extend((level, reference))
    properties.append(numbering)


def _docx_add_table(
    doc,
    table_model,
    number,
    locale,
    set_font,
    qn,
    element_factory,
    paragraph_alignment,
    vertical_alignment,
    rule_color,
) -> None:
    from docx.shared import Pt

    caption = doc.add_paragraph(style="Caption")
    caption.alignment = paragraph_alignment.LEFT
    caption.add_run(_table_caption(locale, number, table_model["title"]))
    columns = table_model["columns"]
    widths = _column_widths_dxa(table_model)
    table = doc.add_table(rows=1, cols=len(columns))
    table.autofit = False
    _set_table_geometry(table, widths, qn, element_factory)
    _set_three_line_table(table, qn, element_factory, rule_color)
    _repeat_table_header(table.rows[0], qn, element_factory)
    for index, column in enumerate(columns):
        cell = table.rows[0].cells[index]
        cell.vertical_alignment = vertical_alignment.CENTER
        _set_cell_width(cell, widths[index], qn, element_factory)
        _set_cell_margins(cell, qn, element_factory)
        _set_cell_border(cell, "bottom", 8, rule_color, qn, element_factory)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = _docx_alignment(column, table_model["rows"], index, paragraph_alignment)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        set_font(
            paragraph.add_run(column["label"]),
            name="Arial",
            east_asia="Microsoft YaHei",
            size=8.5,
            bold=True,
            color=_hex_rgb(f"#{rule_color}"),
        )
    for row_values in table_model["rows"]:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cell = cells[index]
            cell.vertical_alignment = vertical_alignment.CENTER
            _set_cell_width(cell, widths[index], qn, element_factory)
            _set_cell_margins(cell, qn, element_factory)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = _docx_alignment(
                columns[index], table_model["rows"], index, paragraph_alignment
            )
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            set_font(paragraph.add_run(value), size=8.8)
    if table_model["caption"]:
        note = doc.add_paragraph(style="Caption")
        note.alignment = paragraph_alignment.LEFT
        note.paragraph_format.keep_with_next = False
        set_font(note.add_run(table_model["caption"]), size=8.5, color=(103, 113, 121))
    else:
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(1)


def _set_table_geometry(table, widths, qn, element_factory) -> None:
    table_xml = table._tbl
    properties = table_xml.tblPr
    width = properties.find(qn("w:tblW"))
    if width is None:
        width = element_factory("w:tblW")
        properties.append(width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(sum(widths)))
    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = element_factory("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), "120")
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = element_factory("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table_xml.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        col = element_factory("w:gridCol")
        col.set(qn("w:w"), str(value))
        grid.append(col)


def _set_three_line_table(table, qn, element_factory, rule_color) -> None:
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = element_factory("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = element_factory(f"w:{edge}")
            borders.append(node)
        if edge in {"top", "bottom"}:
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), "12")
            node.set(qn("w:color"), rule_color)
        else:
            node.set(qn("w:val"), "nil")


def _set_cell_border(cell, edge, size, color, qn, element_factory) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = element_factory("w:tcBorders")
        properties.append(borders)
    node = borders.find(qn(f"w:{edge}"))
    if node is None:
        node = element_factory(f"w:{edge}")
        borders.append(node)
    node.set(qn("w:val"), "single")
    node.set(qn("w:sz"), str(size))
    node.set(qn("w:color"), color)


def _set_cell_width(cell, width, qn, element_factory) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = properties.find(qn("w:tcW"))
    if node is None:
        node = element_factory("w:tcW")
        properties.append(node)
    node.set(qn("w:type"), "dxa")
    node.set(qn("w:w"), str(width))


def _set_cell_margins(cell, qn, element_factory) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.find(qn("w:tcMar"))
    if margins is None:
        margins = element_factory("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = element_factory(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _repeat_table_header(row, qn, element_factory) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = element_factory("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _register_pdf_fonts(reportlab_package, pdfmetrics, reportlab_ttfont, model) -> None:
    """Register embedded fonts with full scientific and Simplified-Chinese coverage."""
    with _PDF_FONT_LOCK:
        registered = set(pdfmetrics.getRegisteredFontNames())
        if "PaperSans" not in registered:
            font_files = _latin_pdf_font_files(reportlab_package)
            for font_name, path in font_files.items():
                pdfmetrics.registerFont(reportlab_ttfont(font_name, str(path)))
            pdfmetrics.registerFontFamily(
                "PaperSans",
                normal="PaperSans",
                bold="PaperSansBold",
                italic="PaperSansItalic",
                boldItalic="PaperSansBoldItalic",
            )

        required_text = _pdf_required_text(model)
        if any(_is_cjk(char) for char in required_text):
            registered = set(pdfmetrics.getRegisteredFontNames())
            if _PDF_CJK_REGULAR not in registered:
                font_root = _pdf_font_root()
                sources = {
                    _PDF_CJK_REGULAR: font_root / "noto-sans-sc-400.ttf",
                    _PDF_CJK_BOLD: font_root / "noto-sans-sc-700.ttf",
                }
                for font_name, source in sources.items():
                    pdfmetrics.registerFont(reportlab_ttfont(font_name, str(source)))
                pdfmetrics.registerFontFamily(
                    "PaperCJK",
                    normal=_PDF_CJK_REGULAR,
                    bold=_PDF_CJK_BOLD,
                    italic=_PDF_CJK_REGULAR,
                    boldItalic=_PDF_CJK_BOLD,
                )
        _validate_pdf_glyph_coverage(pdfmetrics, model)


def _pdf_font_root() -> Path:
    """Resolve bundled fonts in source installs and one-file PyInstaller builds."""
    roots = []
    frozen_root = getattr(sys, "_MEIPASS", None)
    if frozen_root:
        roots.extend(
            (
                Path(frozen_root) / "vcstudio_assets" / "fonts",
                Path(frozen_root) / "vcstudio" / "gui_web" / "assets" / "fonts",
                Path(frozen_root) / "gui_web" / "assets" / "fonts",
            )
        )
    roots.append(Path(__file__).resolve().parents[1] / "gui_web" / "assets" / "fonts")
    required = ("noto-sans-sc-400.ttf", "noto-sans-sc-700.ttf")
    for root in roots:
        if all((root / name).is_file() for name in required):
            return root
    raise ReportDependencyError(
        "CJK PDF generation requires bundled noto-sans-sc-400.ttf and "
        "noto-sans-sc-700.ttf"
    )


def _latin_pdf_font_files(reportlab_package) -> dict[str, Path]:
    """Use bundled DejaVu for scientific symbols without a matplotlib dependency."""
    try:
        root = _pdf_font_root()
        candidates = {
            "PaperSans": root / "dejavu-sans-400.ttf",
            "PaperSansBold": root / "dejavu-sans-700.ttf",
            "PaperSansItalic": root / "dejavu-sans-400-italic.ttf",
            "PaperSansBoldItalic": root / "dejavu-sans-700-italic.ttf",
        }
        if all(path.is_file() for path in candidates.values()):
            return candidates
    except ReportDependencyError:
        pass

    # Backward compatibility for source checkouts made before the standalone
    # scientific fonts were bundled.
    try:
        matplotlib = importlib.import_module("matplotlib")
        root = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
        candidates = {
            "PaperSans": root / "DejaVuSans.ttf",
            "PaperSansBold": root / "DejaVuSans-Bold.ttf",
            "PaperSansItalic": root / "DejaVuSans-Oblique.ttf",
            "PaperSansBoldItalic": root / "DejaVuSans-BoldOblique.ttf",
        }
        if all(path.is_file() for path in candidates.values()):
            return candidates
    except (ImportError, AttributeError):
        pass
    root = Path(reportlab_package.__file__).resolve().parent / "fonts"
    candidates = {
        "PaperSans": root / "Vera.ttf",
        "PaperSansBold": root / "VeraBd.ttf",
        "PaperSansItalic": root / "VeraIt.ttf",
        "PaperSansBoldItalic": root / "VeraBI.ttf",
    }
    if not all(path.is_file() for path in candidates.values()):
        raise ReportDependencyError("ReportLab's bundled Vera fonts are missing")
    return candidates


def _pdf_required_text(model: dict) -> str:
    payload = json.dumps(_fingerprint_model(model), ensure_ascii=False, sort_keys=True)
    labels = "".join(_SECTION_LABELS[model["locale"]].values())
    furniture = "Table Figure Page" if model["locale"] == "en-US" else "表图第页"
    return payload + labels + furniture


def _validate_pdf_glyph_coverage(pdfmetrics, model: dict) -> None:
    text = _pdf_required_text(model)
    latin = pdfmetrics.getFont("PaperSans").face.charToGlyph
    cjk = None
    if any(_is_cjk(char) for char in text):
        cjk = pdfmetrics.getFont(_PDF_CJK_REGULAR).face.charToGlyph
    missing = []
    for char in text:
        if char.isspace() or ord(char) < 32:
            continue
        coverage = cjk if _is_cjk(char) else latin
        if coverage is None or ord(char) not in coverage:
            if char not in missing:
                missing.append(char)
    if missing:
        preview = " ".join(repr(char) for char in missing[:10])
        raise ReportDependencyError(
            "PDF fonts cannot render all report characters "
            f"({preview}); install matplotlib for scientific glyph coverage"
        )


def _render_pdf(model: dict, output: Path) -> None:
    import reportlab as reportlab_package
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        BaseDocTemplate,
        CondPageBreak,
        Frame,
        Image,
        KeepTogether,
        PageBreak,
        PageTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
    )

    _register_pdf_fonts(reportlab_package, pdfmetrics, TTFont, model)
    font = "PaperSans"
    bold_font = "PaperSansBold"
    theme = _report_theme(model)
    palette = {
        key: colors.HexColor(theme[key])
        for key in (
            "ink", "muted", "accent", "highlight", "rule", "title",
            "subtitle", "caption",
        )
    }
    base = getSampleStyleSheet()
    styles = {
        "body": ParagraphStyle(
            "PaperBody",
            parent=base["BodyText"],
            fontName=font,
            fontSize=10,
            leading=14.2,
            textColor=palette["ink"],
            alignment=TA_JUSTIFY,
            spaceAfter=6,
        ),
        "cover_kicker": ParagraphStyle(
            "CoverKicker",
            fontName=font,
            fontSize=9,
            leading=12,
            textColor=palette["highlight"],
            alignment=TA_CENTER,
            spaceAfter=18 * mm,
        ),
        "cover_title": ParagraphStyle(
            "CoverTitle",
            fontName=bold_font,
            fontSize=28,
            leading=35,
            textColor=palette["title"],
            alignment=TA_CENTER,
            spaceAfter=5 * mm,
        ),
        "cover_status": ParagraphStyle(
            "CoverStatus",
            fontName=bold_font,
            fontSize=8.5,
            leading=12,
            textColor=palette["muted"],
            alignment=TA_CENTER,
            spaceAfter=5 * mm,
        ),
        "cover_subtitle": ParagraphStyle(
            "CoverSubtitle",
            fontName=font,
            fontSize=14,
            leading=21,
            textColor=palette["subtitle"],
            alignment=TA_CENTER,
            spaceAfter=20 * mm,
        ),
        "metadata": ParagraphStyle(
            "Metadata",
            fontName=font,
            fontSize=9,
            leading=13,
            textColor=palette["muted"],
            alignment=TA_CENTER,
            spaceAfter=2,
        ),
        "h1": ParagraphStyle(
            "SectionHeading",
            fontName=bold_font,
            fontSize=15.5,
            leading=20,
            textColor=palette["title"],
            spaceBefore=11 * mm,
            spaceAfter=4 * mm,
            keepWithNext=True,
        ),
        "h2": ParagraphStyle(
            "Subheading",
            fontName=bold_font,
            fontSize=11.5,
            leading=15,
            textColor=palette["accent"],
            spaceBefore=4 * mm,
            spaceAfter=1.5 * mm,
            keepWithNext=True,
        ),
        "number": ParagraphStyle(
            "NumberedPoint",
            parent=base["BodyText"],
            fontName=font,
            fontSize=10,
            leading=14,
            textColor=palette["ink"],
            leftIndent=7 * mm,
            firstLineIndent=-7 * mm,
            spaceAfter=7,
        ),
        "table_caption": ParagraphStyle(
            "TableCaption",
            fontName=bold_font,
            fontSize=8.8,
            leading=12,
            textColor=palette["caption"],
            spaceBefore=2 * mm,
            spaceAfter=2 * mm,
            keepWithNext=True,
        ),
        "table_head": ParagraphStyle(
            "TableHead",
            fontName=bold_font,
            fontSize=8,
            leading=10,
            textColor=palette["rule"],
        ),
        "table_cell": ParagraphStyle(
            "TableCell",
            fontName=font,
            fontSize=8.3,
            leading=10.5,
            textColor=palette["ink"],
        ),
        "table_note": ParagraphStyle(
            "TableNote",
            fontName=font,
            fontSize=8.2,
            leading=11,
            textColor=palette["muted"],
            spaceAfter=3 * mm,
        ),
        "figure_caption": ParagraphStyle(
            "FigureCaption",
            fontName=font,
            fontSize=8.8,
            leading=12,
            textColor=palette["caption"],
            alignment=TA_CENTER,
            spaceBefore=2.5 * mm,
            spaceAfter=5 * mm,
        ),
    }

    document = BaseDocTemplate(
        str(output),
        pagesize=A4,
        rightMargin=theme["page_margin_mm"] * mm,
        leftMargin=theme["page_margin_mm"] * mm,
        topMargin=22 * mm,
        bottomMargin=24 * mm,
        title=model["title"],
        author="VASP Catalyst Studio",
        subject=model["subtitle"],
        creator="VASP Catalyst Studio paper_report",
        keywords=["scientific report", "adsorption energy", "catalysis"],
        lang=model["locale"],
    )

    story = [Spacer(1, theme["cover_spacer_mm"] * mm)]
    story.append(
        Paragraph(_pdf_markup(model["kicker"].upper(), _PDF_CJK_BOLD), styles["cover_kicker"])
    )
    story.append(
        Paragraph(_pdf_markup(_report_kind_label(model), _PDF_CJK_BOLD), styles["cover_status"])
    )
    story.append(Paragraph(_pdf_markup(model["title"], _PDF_CJK_BOLD), styles["cover_title"]))
    if model["subtitle"]:
        story.append(Paragraph(_pdf_markup(model["subtitle"]), styles["cover_subtitle"]))
    else:
        story.append(Spacer(1, 13 * mm))
    for item in model["metadata"]:
        text = _pdf_markup(f"{item['label'].upper()}  ·  {item['value']}")
        story.append(Paragraph(text, styles["metadata"]))
    story.append(PageBreak())

    sections = _section_plan(model)
    table_number = 0
    figure_number = 0
    for section_number, (key, heading) in enumerate(sections, 1):
        minimum_space = 165 * mm if key == "figures" else 35 * mm
        story.append(CondPageBreak(minimum_space))
        story.append(
            Paragraph(
                _pdf_markup(f"{section_number:02d}  {heading}", _PDF_CJK_BOLD),
                styles["h1"],
            )
        )
        if key == "executive_summary":
            _pdf_add_blocks(story, model[key], styles, ordered=False)
        elif key == "key_findings":
            _pdf_add_blocks(story, model[key], styles, ordered=True)
        elif key in {"candidate_evaluations", "adsorption_table", "comparison_table"}:
            table_number += 1
            story.extend(
                _pdf_table(
                    model[key],
                    table_number,
                    model["locale"],
                    styles,
                    document.width,
                    colors,
                    Table,
                    TableStyle,
                    Paragraph,
                    TA_LEFT,
                    TA_RIGHT,
                    TA_CENTER,
                    palette["rule"],
                )
            )
        elif key == "figures":
            for figure in model["figures"]:
                figure_number += 1
                image = Image(str(figure["_render_path"]))
                scale = min(
                    document.width / image.drawWidth,
                    (132 * mm) / image.drawHeight,
                    1.0,
                )
                image.drawWidth *= scale
                image.drawHeight *= scale
                caption = Paragraph(
                    _pdf_markup(
                        _figure_caption(
                            model["locale"],
                            figure_number,
                            _caption_text(figure["title"], figure["caption"]),
                        )
                    ),
                    styles["figure_caption"],
                )
                story.append(KeepTogether([image, caption]))
        else:
            _pdf_add_blocks(
                story,
                model[key],
                styles,
                ordered=key == "recommendations",
            )

    def draw_page(canvas, doc):
        canvas.saveState()
        canvas.setTitle(model["title"])
        canvas.setAuthor("VASP Catalyst Studio")
        canvas.setSubject(model["subtitle"])
        canvas.setCreator("VASP Catalyst Studio paper_report")
        canvas.setKeywords(["scientific report", "adsorption energy", "catalysis"])
        canvas.setFont(font, 7.8)
        canvas.setFillColor(palette["muted"])
        if doc.page > 1:
            _draw_pdf_mixed_text(
                canvas,
                model["kicker"].upper(),
                document.leftMargin,
                A4[1] - 11 * mm,
                font,
                _PDF_CJK_REGULAR,
                7.8,
                align="left",
            )
            title_text = model["title"]
            if len(title_text) > 42:
                title_text = title_text[:39] + "..."
            _draw_pdf_mixed_text(
                canvas,
                title_text,
                A4[0] - document.rightMargin,
                A4[1] - 11 * mm,
                font,
                _PDF_CJK_REGULAR,
                7.8,
                align="right",
            )
        page_label = (
            f"第 {doc.page} 页" if model["locale"] == "zh-CN" else f"Page {doc.page}"
        )
        _draw_pdf_mixed_text(
            canvas,
            page_label,
            A4[0] / 2,
            10 * mm,
            font,
                _PDF_CJK_REGULAR,
            7.8,
            align="center",
        )
        canvas.restoreState()

    frame = Frame(
        document.leftMargin,
        document.bottomMargin,
        document.width,
        document.height,
        id="report-body",
        leftPadding=0,
        rightPadding=0,
        topPadding=0,
        bottomPadding=0,
    )
    # ``onPage`` runs before story flowables; an opaque plot that ever reaches
    # a margin can then cover the running header.  Draw furniture in
    # ``onPageEnd`` so page number and provenance remain visible on top.
    document.addPageTemplates(
        PageTemplate(id="report", frames=[frame], onPageEnd=draw_page)
    )
    document.build(story)


def _pdf_add_blocks(story, blocks, styles, *, ordered: bool) -> None:
    from reportlab.platypus import Paragraph

    for index, block in enumerate(blocks, 1):
        if ordered:
            label = f"{index}. "
            if block["title"]:
                label += f"{block['title']}. "
            story.append(Paragraph(_pdf_markup(label + block["text"]), styles["number"]))
        else:
            if block["title"]:
                story.append(
                    Paragraph(_pdf_markup(block["title"], _PDF_CJK_BOLD), styles["h2"])
                )
            for text in _paragraphs(block["text"]):
                story.append(Paragraph(_pdf_markup(text), styles["body"]))


def _pdf_table(
    table_model,
    number,
    locale,
    styles,
    content_width,
    colors,
    table_class,
    table_style_class,
    paragraph_class,
    align_left,
    align_right,
    align_center,
    rule_color,
):
    caption = paragraph_class(
        _pdf_markup(_table_caption(locale, number, table_model["title"]), _PDF_CJK_BOLD),
        styles["table_caption"],
    )
    data = [
        [
            paragraph_class(
                _pdf_markup(column["label"], _PDF_CJK_BOLD),
                styles["table_head"],
            )
            for column in table_model["columns"]
        ]
    ]
    for row in table_model["rows"]:
        data.append(
            [paragraph_class(_pdf_markup(value), styles["table_cell"]) for value in row]
        )
    dxa_widths = _column_widths_dxa(table_model)
    widths = [content_width * value / sum(dxa_widths) for value in dxa_widths]
    table = table_class(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    commands = [
        ("LINEABOVE", (0, 0), (-1, 0), 1.1, rule_color),
        ("LINEBELOW", (0, 0), (-1, 0), 0.7, rule_color),
        ("LINEBELOW", (0, -1), (-1, -1), 1.1, rule_color),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    for index, column in enumerate(table_model["columns"]):
        align = _resolved_align(column, table_model["rows"], index)
        pdf_align = {"left": align_left, "right": align_right, "center": align_center}[align]
        commands.append(("ALIGN", (index, 0), (index, -1), {
            align_left: "LEFT", align_right: "RIGHT", align_center: "CENTER"
        }[pdf_align]))
    table.setStyle(table_style_class(commands))
    result = [caption, table]
    if table_model["caption"]:
        result.append(
            paragraph_class(_pdf_markup(table_model["caption"]), styles["table_note"])
        )
    return result


def _section_plan(model: dict) -> list[tuple[str, str]]:
    labels = _SECTION_LABELS[model["locale"]]
    return [
        (key, labels[key])
        for key in model.get("outline") or labels
        if key in labels and model.get(key)
    ]


def _column_widths_dxa(table: dict) -> list[int]:
    weights = []
    for index, column in enumerate(table["columns"]):
        lengths = [_visual_length(column["label"])]
        lengths.extend(_visual_length(row[index]) for row in table["rows"])
        weight = max(5.0, min(36.0, max(lengths, default=5)))
        if _resolved_align(column, table["rows"], index) == "right":
            weight = min(weight, 14.0)
        weights.append(weight)
    minimum = min(900, _A4_CONTENT_WIDTH_DXA // max(len(weights), 1))
    remaining = _A4_CONTENT_WIDTH_DXA - minimum * len(weights)
    total_weight = sum(weights) or 1
    result = [minimum + round(remaining * weight / total_weight) for weight in weights]
    result[-1] += _A4_CONTENT_WIDTH_DXA - sum(result)
    return result


def _docx_alignment(column, rows, index, alignment):
    resolved = _resolved_align(column, rows, index)
    return {
        "left": alignment.LEFT,
        "right": alignment.RIGHT,
        "center": alignment.CENTER,
    }[resolved]


def _html_align_class(column, rows, index) -> str:
    return "numeric" if _resolved_align(column, rows, index) == "right" else ""


def _resolved_align(column: dict, rows: Sequence[Sequence[str]], index: int) -> str:
    if column["align"] in {"left", "right", "center"}:
        return column["align"]
    values = [row[index].strip() for row in rows if index < len(row) and row[index].strip()]
    if values and all(_looks_numeric(value) for value in values):
        return "right"
    return "left"


def _valid_align(value: str) -> str:
    return value if value in {"left", "right", "center"} else "auto"


def _looks_numeric(value: str) -> bool:
    compact = value.replace(",", "").replace("−", "-").strip()
    return bool(re.fullmatch(r"[+\-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+\-]?\d+)?(?:\s*\S+)?", compact))


def _visual_length(value: str) -> float:
    return sum(2 if ord(char) > 127 else 1 for char in value)


def _caption_text(title: str, caption: str) -> str:
    if caption and caption != title:
        return f"{title}. {caption}"
    return title or caption


def _table_caption(locale: str, number: int, title: str) -> str:
    if locale == "zh-CN":
        return f"表 {number}  {title}"
    return f"Table {number}. {title}"


def _figure_caption(locale: str, number: int, caption: str) -> str:
    if locale == "zh-CN":
        return f"图 {number}  {caption}"
    return f"Figure {number}. {caption}"


def _page_counter_css(locale: str) -> str:
    return '"第 " counter(page) " 页"' if locale == "zh-CN" else '"Page " counter(page)'


def _paragraphs(text: str) -> list[str]:
    return [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()] or [""]


def _human_label(value: str) -> str:
    return value.replace("_", " ").strip().title()


def _safe_asset_stem(value: str) -> str:
    sanitized = re.sub(r"[^\w.-]+", "-", value, flags=re.UNICODE).strip("-.")
    return (sanitized or "figure")[:48]


def _display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("report model contains a non-finite number")
        return f"{value:g}"
    if isinstance(value, (str, int)):
        return str(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Mapping):
        return "; ".join(f"{_text(key)}: {_display(val)}" for key, val in value.items())
    if isinstance(value, Sequence) and not isinstance(value, (str, bytes)):
        return "; ".join(_display(item) for item in value)
    if isinstance(value, os.PathLike):
        return os.fspath(value)
    raise TypeError(
        f"report model contains unsupported display value type: {type(value).__name__}"
    )


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, bool):
        return str(value).strip()
    if isinstance(value, int):
        return str(value).strip()
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("report model contains a non-finite number")
        return str(value).strip()
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, os.PathLike):
        return os.fspath(value).strip()
    raise TypeError(
        f"report model contains unsupported text value type: {type(value).__name__}"
    )


def _json_safe(value: Any) -> Any:
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("report model contains a non-finite number")
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, os.PathLike):
        return os.fspath(value)
    if isinstance(value, Mapping):
        result = {}
        for key, val in value.items():
            if not isinstance(key, str):
                raise TypeError("report model JSON object keys must be strings")
            result[key] = _json_safe(val)
        return result
    if isinstance(value, Sequence) and not isinstance(value, (str, bytes)):
        return [_json_safe(item) for item in value]
    raise TypeError(
        f"report model contains a non-JSON value of type {type(value).__name__}"
    )


def _canonical_json_bytes(value: Any) -> bytes:
    return json.dumps(
        _json_safe(value),
        ensure_ascii=False,
        allow_nan=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")


def _sha256_json(value: Any) -> str:
    return hashlib.sha256(_canonical_json_bytes(value)).hexdigest()


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _file_record(path: Path, *, relative_path: str) -> dict:
    return {
        "path": relative_path,
        "sha256": _sha256_file(path),
        "size": path.stat().st_size,
    }


def _css_string(value: str) -> str:
    return (
        str(value)
        .replace("\\", "\\\\")
        .replace('"', '\\"')
        .replace("<", "\\3C ")
        .replace(">", "\\3E ")
        .replace("&", "\\26 ")
        .replace("\n", " ")
    )


def _pdf_markup(value: str, cjk_font: str = _PDF_CJK_REGULAR) -> str:
    """Escape Paragraph markup and select a CJK fallback without harming Latin spacing."""
    pieces = []
    for is_cjk, text in _pdf_text_runs(str(value)):
        escaped = xml_escape(text)
        if is_cjk:
            pieces.append(f'<font name="{cjk_font}">{escaped}</font>')
        else:
            pieces.append(escaped)
    return "".join(pieces)


def _pdf_text_runs(value: str) -> list[tuple[bool, str]]:
    runs: list[tuple[bool, str]] = []
    for char in value:
        is_cjk = _is_cjk(char)
        if runs and runs[-1][0] == is_cjk:
            runs[-1] = (is_cjk, runs[-1][1] + char)
        else:
            runs.append((is_cjk, char))
    return runs


def _is_cjk(char: str) -> bool:
    codepoint = ord(char)
    return (
        0x2E80 <= codepoint <= 0x9FFF
        or 0xF900 <= codepoint <= 0xFAFF
        or 0x3040 <= codepoint <= 0x30FF
        or 0xAC00 <= codepoint <= 0xD7AF
        or 0xFF00 <= codepoint <= 0xFFEF
    )


def _draw_pdf_mixed_text(
    canvas,
    value: str,
    x: float,
    y: float,
    latin_font: str,
    cjk_font: str,
    size: float,
    *,
    align: str,
) -> None:
    runs = _pdf_text_runs(str(value))
    widths = [
        canvas.stringWidth(text, cjk_font if is_cjk else latin_font, size)
        for is_cjk, text in runs
    ]
    total_width = sum(widths)
    if align == "right":
        cursor = x - total_width
    elif align == "center":
        cursor = x - total_width / 2
    else:
        cursor = x
    for (is_cjk, text), width in zip(runs, widths):
        canvas.setFont(cjk_font if is_cjk else latin_font, size)
        canvas.drawString(cursor, y, text)
        cursor += width
