"""收尾流水线 Draft-Ready —— 把项目全体成员的真实产物装订成投稿级 SI/三线表/方法学/口径稽核。

铁律(设计宪法第一条):**只组织磁盘上真实存在的数据,绝不生成任何数值**。
缺文件/缺能量 → 写 ``[待确认:...]`` 占位并入 issues,绝不用默认值冒充;引文缺失只告警不编造。
POTCAR 本体为 VASP 版权材料,一律**不打包**(只导出身份表供溯源与复算对齐)。

复用既有公开函数(不碰私有名):
- adsorption.delta_e_rows / export_csv(ΔE/ΔΔE/最稳/续算,单一真相源);
- methods_text.extract_facts / render_zh / render_en / render_bibtex(双语方法学 + BibTeX);
- incar_builder.parse_incar / dispersion_audit(参数解析 + 色散一致性);
- screening.screening_table_csv(描述符表 CSV);manifest.load_manifest(状态/能量/溯源)。

时间戳一律由调用方经 ``timestamp`` 传入(缺省回落本机时间),便于可复现测试。
中文注释允许,英文标识符。
"""
from __future__ import annotations

import csv
import io
import os
import tempfile
import time
import zipfile
from pathlib import Path

import yaml

from vcstudio import __version__
from vcstudio.generate import methods_text
from vcstudio.generate.incar_builder import dispersion_audit, parse_incar
from vcstudio.project import adsorption
from vcstudio.shared import manifest as manifest_mod

# 逐成员打包的真实输入文件与"是否复现必需"(POTCAR 本体绝不在列 —— 版权)。
_MEMBER_FILES = (('INCAR', True), ('KPOINTS', False), ('POSCAR', True),
                 ('CONTCAR', False), ('job.yaml', True))


# ── 公共小工具 ────────────────────────────────────────────────────────────────

def _now() -> str:
    return time.strftime('%Y-%m-%dT%H:%M:%S')


def _ts(timestamp: str | None) -> str:
    """生成时间:调用方给了就用(可复现),否则回落本机时间(仅作脚注,非科学数值)。"""
    return timestamp or _now()


def _read_file(job_dir, name: str) -> str | None:
    """读作业目录下某文本文件;缺失/不可读 → None(按"无此文件"降级)。"""
    try:
        with open(os.path.join(str(job_dir), name), 'r', encoding='utf-8',
                  errors='replace') as f:
            return f.read()
    except OSError:
        return None


def _member_records(project: dict) -> list:
    """project.members → [{'role','name','dir'}](仅保留已设的成员目录,保序)。"""
    mem = (project or {}).get('members') or {}
    recs: list = []

    def _add(role, d):
        if d:
            recs.append({'role': role, 'dir': str(d),
                         'name': os.path.basename(os.path.normpath(str(d)))})

    _add('clean_slab', mem.get('clean_slab'))
    _add('gas_ref', mem.get('gas_ref'))
    for c in (mem.get('configs') or []):
        _add('config', c)
    return recs


def _continuation_rounds(m: dict | None) -> int:
    """续算轮数 = manifest.attempts 里 result==continued 的次数(与 report_full 同口径)。"""
    return sum(1 for a in ((m or {}).get('attempts') or [])
               if a.get('result') == 'continued')


def _potcar_provenance_of(m: dict | None) -> list:
    """取 manifest 记录的赝势身份表(兼容 inputs.potcar_provenance 与 inputs.potcar 两键)。"""
    inputs = (m or {}).get('inputs') or {}
    return list(inputs.get('potcar_provenance') or inputs.get('potcar') or [])


def _richest_potcar_provenance(recs: list) -> list:
    """跨成员取元素最全的一份赝势身份(构型含吸附质,元素多于清洁表面)。"""
    best: list = []
    for r in recs:
        cand = _potcar_provenance_of(manifest_mod.load_manifest(r['dir']))
        if len(cand) > len(best):
            best = cand
    return best


def _incar_val(text: str | None, key: str):
    """从 INCAR 文本取某键(已大写化);缺文本/缺键 → None。"""
    if text is None:
        return None
    return parse_incar(text).get(key)


def _group_by_val(pairs: list) -> dict:
    """[(name, value)] → {value: [name,...]}(保出现序)。"""
    grp: dict = {}
    for n, v in pairs:
        grp.setdefault(v, []).append(n)
    return grp


# ── 1. SI 装订机 ─────────────────────────────────────────────────────────────

def _potcar_identity_csv(prov: list) -> str:
    """赝势身份表 CSV 文本(无 BOM;写入时统一 utf-8-sig)。"""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(['元素', 'variant', 'TITEL', 'ENMAX/eV'])
    for p in prov:
        w.writerow([p.get('element', ''), p.get('variant', ''),
                    p.get('titel', ''), p.get('enmax', '')])
    return buf.getvalue()


def _energy_summary_csv(project: dict, summary: dict) -> str:
    """复用 adsorption.export_csv 生成能量汇总(真实 ΔE,零再生成),读回为字符串。"""
    with tempfile.TemporaryDirectory() as td:
        p = adsorption.export_csv(project, summary, os.path.join(td, 'e.csv'))
        return Path(p).read_text(encoding='utf-8-sig')


def _si_readme(name: str, ts: str, members: list, has_potcar_id: bool) -> str:
    """SI 包 README(中文逐文件说明)。"""
    lines = [
        'VASP Catalyst Studio —— 支撑信息(SI)数据包',
        f'项目:{name}',
        f'生成时间:{ts}',
        f'生成工具:vcstudio {__version__}',
        '',
        '── 目录/文件说明 ──',
        '  members/<成员名>/   每个计算成员的真实输入(可直接复算):',
        '    INCAR      VASP 参数文件',
        '    KPOINTS    K 点网格(缺失表示按 INCAR 自动/Γ 点)',
        '    POSCAR     初始结构',
        '    CONTCAR    弛豫后结构(若存在)',
        '    job.yaml   vcstudio 任务清单(状态/能量/输入指纹)',
        '  POTCAR_identity.csv   赝势身份表(元素/variant/TITEL/ENMAX)。',
        '      注:POTCAR 本体为 VASP 版权材料,不随 SI 分发;此表仅供溯源与复算对齐。',
        '  energy_summary.csv    吸附能汇总(ΔE/ΔΔE/是否最稳/状态);数值全部取自各 job.yaml 真实结果。',
        '  metadata.yaml         vcstudio 版本与输入指纹(POSCAR/POTCAR sha256),供复现核对。',
        '',
        '── 成员清单 ──',
    ]
    for mr in members:
        tail = ''
        if mr['energy_e0_eV'] is not None:
            tail += f",E0={mr['energy_e0_eV']} eV"
        else:
            tail += ',能量待确认'
        if mr['continuation_rounds']:
            tail += f",续算×{mr['continuation_rounds']}"
        lines.append(f"  - {mr['name']}({mr['role']}):{mr['state']}{tail}")
    if not has_potcar_id:
        lines += ['', '[待确认:未找到 POTCAR 身份信息,请补全各 job.yaml 的 inputs.potcar]']
    lines += ['', '铁律:本包只搬运真实文件与结构化导出,不生成任何数值;缺项一律以 [待确认] 标注。']
    return '\n'.join(lines) + '\n'


def build_si_package(project: dict, out_zip, *, config: dict | None = None,
                     timestamp: str | None = None) -> dict:
    """SI 装订机:把项目全体成员的真实输入 + 能量汇总 + 赝势身份 + 元数据 + README 打成 zip。

    绝不打包 POTCAR 本体(版权);缺文件记 issues 不中断。返回
    ``{'ok','zip_path','contents':[...],'issues':[...]}``,ok=有成员且无关键输入文件缺失。
    """
    config = config or {}
    ts = _ts(timestamp)
    recs = _member_records(project)
    contents: list = []
    issues: list = []
    critical_missing = False
    name = str((project or {}).get('name') or 'project')

    if not recs:
        issues.append('[待确认:项目无任何成员目录,SI 包为空]')

    out_zip = str(out_zip)
    os.makedirs(os.path.dirname(os.path.abspath(out_zip)) or '.', exist_ok=True)

    meta_members: list = []
    with zipfile.ZipFile(out_zip, 'w', zipfile.ZIP_DEFLATED) as zf:
        for r in recs:
            m = manifest_mod.load_manifest(r['dir']) or {}
            for fname, required in _MEMBER_FILES:
                src = os.path.join(r['dir'], fname)
                if os.path.isfile(src):
                    arc = f"members/{r['name']}/{fname}"
                    zf.write(src, arc)
                    contents.append(arc)
                elif required:
                    critical_missing = True
                    issues.append(f"[待确认:成员 {r['name']} 缺 {fname}(SI 复现需要)]")
            e = (m.get('results') or {}).get('energy_e0_eV')
            k = m.get('kpoints') or (m.get('inputs') or {}).get('kpoints')
            meta_members.append({
                'name': r['name'], 'role': r['role'], 'state': m.get('state', '?'),
                'energy_e0_eV': (float(e) if isinstance(e, (int, float)) else None),
                'kpoints': (list(k) if k else None),
                'poscar_sha256': (m.get('inputs') or {}).get('poscar_sha256'),
                'potcar_sha256': (m.get('inputs') or {}).get('potcar_sha256'),
                'continuation_rounds': _continuation_rounds(m),
            })

        # POTCAR 身份表(不含本体)
        prov = _richest_potcar_provenance(recs)
        if prov:
            zf.writestr('POTCAR_identity.csv',
                        _potcar_identity_csv(prov).encode('utf-8-sig'))
            contents.append('POTCAR_identity.csv')
        else:
            issues.append('[待确认:无 POTCAR 身份信息(job.yaml 缺 inputs.potcar),赝势溯源不可用]')

        # 能量汇总 CSV(真实 ΔE);缺能量的构型入 issues,绝不冒充
        try:
            summary = adsorption.delta_e_rows(project)
            zf.writestr('energy_summary.csv',
                        _energy_summary_csv(project, summary).encode('utf-8-sig'))
            contents.append('energy_summary.csv')
            for row in summary.get('rows', []):
                if row.get('delta_e') is None:
                    issues.append(f"[待确认:{row.get('name')} 无 ΔE —— "
                                  f"{row.get('note') or '成员未完成'}]")
        except Exception as e:                       # noqa: BLE001 汇总失败降级不中断打包
            issues.append(f'[待确认:能量汇总失败:{e}]')

        # 元数据:vcstudio 版本 + 输入指纹(sha256 均取自 manifest 真实记录)
        meta = {
            'vcstudio_version': __version__,
            'generated_at': ts,
            'project': {'name': name, 'root': (project or {}).get('root')},
            'members': meta_members,
            'note': ('POTCAR 本体未包含(版权);poscar_sha256/potcar_sha256 为输入指纹,'
                     '供复算对齐。'),
        }
        zf.writestr('metadata.yaml',
                    yaml.safe_dump(meta, allow_unicode=True, sort_keys=False))
        contents.append('metadata.yaml')

        zf.writestr('README.txt', _si_readme(name, ts, meta_members, bool(prov)))
        contents.append('README.txt')

    return {'ok': bool(recs) and not critical_missing, 'zip_path': out_zip,
            'contents': contents, 'issues': issues}


# ── 2. 三线表工厂 ─────────────────────────────────────────────────────────────

def _html_three_line(title: str, headers: list, rows: list, footer: str) -> str:
    """极简三线表(顶/底粗线 + 表头下细线,无竖线;内联 CSS,可直接截图/复制进 Word)。"""
    import html as _h
    th = ''.join(
        f'<th style="padding:4px 12px;text-align:center;'
        f'border-bottom:1.4px solid #222;">{_h.escape(str(x))}</th>' for x in headers)
    body = []
    for r in rows:
        tds = ''.join(f'<td style="padding:3px 12px;text-align:center;">'
                      f'{_h.escape(str(c))}</td>' for c in r)
        body.append(f'<tr>{tds}</tr>')
    return (
        '<table style="border-collapse:collapse;'
        "font-family:'Times New Roman',SimSun,serif;font-size:13px;"
        'border-top:2px solid #222;border-bottom:2px solid #222;margin:10px 0;">'
        f'<caption style="caption-side:top;font-weight:bold;padding:5px;">'
        f'{_h.escape(title)}</caption>'
        f'<thead><tr>{th}</tr></thead><tbody>{"".join(body)}</tbody></table>'
        f'<div style="font-size:11px;color:#555;margin-top:4px;">'
        f'{_h.escape(footer)}</div>')


def _write_xlsx(path: str, headers: list, rows: list, title: str,
                footer: str) -> str | None:
    """openpyxl 可用则写 xlsx 返回路径;不可用返回 None(调用方以 CSV+HTML 兜底)。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return None
    wb = Workbook()
    ws = wb.active
    ws.append([title])
    ws['A1'].font = Font(bold=True)
    ws.append(list(headers))
    for cell in ws[2]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(list(r))
    ws.append([])
    ws.append([footer])
    wb.save(path)
    return path


def _emit_table(base: str, headers: list, rows: list, title: str,
                footer: str) -> list:
    """一张表 → CSV(总在)+ HTML 三线表(总在)+ xlsx(openpyxl 在时)。返回文件路径列表。"""
    out: list = []
    csv_path = base + '.csv'
    with open(csv_path, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f)
        w.writerow([title])
        w.writerow(headers)
        for r in rows:
            w.writerow(r)
        w.writerow([])
        w.writerow([footer])
    out.append(csv_path)
    html_path = base + '.html'
    Path(html_path).write_text(_html_three_line(title, headers, rows, footer),
                               encoding='utf-8')
    out.append(html_path)
    xlsx = _write_xlsx(base + '.xlsx', headers, rows, title, footer)
    if xlsx:
        out.append(xlsx)
    return out


def _fmt4(v) -> str:
    """数值 → 4 位小数;None/非数 → '[待确认]'(缺项绝不留空冒充)。"""
    return f'{v:.4f}' if isinstance(v, (int, float)) and not isinstance(v, bool) \
        else '[待确认]'


def _dash4(v) -> str:
    """数值 → 4 位小数;None/非数 → '—'(描述符缺失属正常留白,非阻塞)。"""
    return f'{v:.4f}' if isinstance(v, (int, float)) and not isinstance(v, bool) \
        else '—'


def _adsorption_table(project: dict, out_dir: str, footer: str) -> tuple:
    """吸附能三线表:构型×(状态/ΔE/ΔΔE/是否最稳/续算轮数/备注)。缺 ΔE 入 issues。"""
    files: list = []
    issues: list = []
    summary = adsorption.delta_e_rows(project)
    configs = ((project.get('members') or {}).get('configs')) or []
    cont = {}
    for d in configs:
        if not d:
            continue
        key = os.path.basename(os.path.normpath(str(d)))
        cont[key] = _continuation_rounds(manifest_mod.load_manifest(d))

    headers = ['构型', '状态', 'ΔE_ads/eV', 'ΔΔE/eV', '是否最稳', '续算轮数', '备注']
    rows: list = []
    for r in summary.get('rows', []):
        if r.get('delta_e') is None:
            issues.append(f"[待确认:{r.get('name')} 无 ΔE —— {r.get('note') or '成员未完成'}]")
        n = cont.get(r.get('name'), 0)
        rows.append([r.get('name', ''), r.get('state', ''), _fmt4(r.get('delta_e')),
                     _fmt4(r.get('dd_e')), '是' if r.get('is_most_stable') else '',
                     str(n) if n else '', r.get('note', '')])
    files += _emit_table(os.path.join(out_dir, 'adsorption_table'), headers, rows,
                         '吸附能汇总表', footer)
    return files, issues


def _free_energy_table(fed: dict, out_dir: str, footer: str) -> tuple:
    """自由能三线表:物种/步 ×(ΔG_step/G 累计/决速步)。fed 为 freeenergy 路径结果。"""
    files: list = []
    issues: list = []
    steps = (fed or {}).get('steps') or []
    if not steps:
        issues.append('[待确认:自由能路径无台阶数据(fed.steps 为空)]')
        return files, issues

    extra = []
    if fed.get('u_l') is not None:
        extra.append(f"U_L={fed['u_l']} V")
    if fed.get('eta') is not None:
        extra.append(f"η={fed['eta']} V")
    extra.append('含 ZPE−TS 热校正' if fed.get('thermo_corrected')
                 else '纯电子能口径(未含 ZPE/熵)')
    ft = footer + '  ' + ';'.join(extra)

    pds = fed.get('pds_index')
    headers = ['步骤', 'ΔG_step/eV', 'G(累计)/eV', '决速步']
    rows: list = []
    prev = None
    for i, st in enumerate(steps):
        g = st.get('G')
        if g is None:
            issues.append(f"[待确认:步骤 {st.get('label')} 无 G]")
        if i == 0:
            dstep = '—'
        elif isinstance(g, (int, float)) and isinstance(prev, (int, float)):
            dstep = f'{g - prev:+.4f}'
        else:
            dstep = '[待确认]'
        # pds_index 为逐电子决速步(st→st+1 的过渡),对应"进入 st+1"那一行
        is_pds = (isinstance(pds, int) and i == pds + 1)
        rows.append([st.get('label', ''), dstep, _fmt4(g), '是' if is_pds else ''])
        prev = g
    files += _emit_table(os.path.join(out_dir, 'free_energy_table'), headers, rows,
                         '自由能台阶表', ft)
    return files, issues


def _descriptor_table(descriptors, out_dir: str, footer: str) -> tuple:
    """描述符三线表:催化剂×(各物种 ΔE/d 带中心/ICOHP/U_L)。CSV 复用 screening 公开函数。"""
    from vcstudio.project import screening
    files: list = []
    issues: list = []
    if isinstance(descriptors, list):                # 容错:直接给 rows 列表
        descriptors = {'rows': descriptors, 'missing': []}
    rows_in = (descriptors or {}).get('rows')
    if not rows_in:
        issues.append('[待确认:描述符表无数据行]')
        return files, issues

    base = os.path.join(out_dir, 'descriptor_table')
    screening.screening_table_csv(descriptors, base + '.csv')   # 复用公开函数(带数据来源列)
    files.append(base + '.csv')

    species = sorted({sp for r in rows_in for sp in (r.get('de') or {})})
    headers = (['催化剂'] + [f'ΔE({sp})/eV' for sp in species]
               + ['d带中心/eV', 'ICOHP_ms/eV', 'U_L/V'])
    trows: list = []
    for r in rows_in:
        de = r.get('de') or {}
        trows.append([r.get('catalyst', '')] + [_dash4(de.get(sp)) for sp in species]
                     + [_dash4(r.get('d_band_center')), _dash4(r.get('icohp_ms')),
                        _dash4(r.get('u_l'))])
    Path(base + '.html').write_text(
        _html_three_line('描述符汇总表', headers, trows, footer), encoding='utf-8')
    files.append(base + '.html')
    xlsx = _write_xlsx(base + '.xlsx', headers, trows, '描述符汇总表', footer)
    if xlsx:
        files.append(xlsx)
    for miss in (descriptors.get('missing') or []):
        issues.append(f"[待确认:{miss.get('catalyst')} 缺 {miss.get('field')} —— "
                      f"{miss.get('reason', '')}]")
    return files, issues


def make_tables(project: dict, out_dir, *, fed: dict | None = None,
                descriptors=None, timestamp: str | None = None) -> dict:
    """三线表工厂:吸附能表(总产)+ 自由能表(fed 提供时)+ 描述符表(descriptors 提供时)。

    每表恒产 CSV + HTML 三线表,openpyxl 在时另产 xlsx。每表带来源脚注(项目路径 +
    调用方传入的 timestamp)。返回 ``{'files':[...],'issues':[...]}``。
    """
    out_dir = str(out_dir)
    os.makedirs(out_dir, exist_ok=True)
    ts = _ts(timestamp)
    name = str((project or {}).get('name') or 'project')
    root = (project or {}).get('root') or ''
    footer = (f'来源:项目 {name}({root});生成时间 {ts};vcstudio {__version__}。'
              f'数值取自 job.yaml 真实结果,未做任何再生成。')

    files: list = []
    issues: list = []
    a_files, a_issues = _adsorption_table(project, out_dir, footer)
    files += a_files
    issues += a_issues
    if fed is not None:
        f_files, f_issues = _free_energy_table(fed, out_dir, footer)
        files += f_files
        issues += f_issues
    if descriptors is not None:
        d_files, d_issues = _descriptor_table(descriptors, out_dir, footer)
        files += d_files
        issues += d_issues
    return {'files': files, 'issues': issues}


# ── 3. 口径稽核 ──────────────────────────────────────────────────────────────

def _audit_encut(incars: dict) -> dict:
    pairs = [(n, _incar_val(t, 'ENCUT')) for n, t in incars.items()]
    if not pairs:
        return {'name': 'ENCUT 一致性', 'ok': False, 'detail': '无成员 INCAR 可读。'}
    missing = [n for n, v in pairs if v is None]
    if missing:
        return {'name': 'ENCUT 一致性', 'ok': False,
                'detail': f'成员 {"、".join(missing)} 的 INCAR 未设 ENCUT(截断能未锁定,基组不可比)。'}
    vals = {v for _, v in pairs}
    if len(vals) == 1:
        return {'name': 'ENCUT 一致性', 'ok': True,
                'detail': f'全部成员 ENCUT={next(iter(vals)):g} eV,基组一致。'}
    grp = _group_by_val(pairs)
    detail = ';'.join(f'{k:g} eV(成员:{"、".join(v)})' for k, v in grp.items())
    return {'name': 'ENCUT 一致性', 'ok': False,
            'detail': f'ENCUT 不一致(ΔE 大数相减被不同基组污染):{detail}。'}


def _func_label(key: tuple) -> str:
    gga, metagga, _ = key
    if metagga and metagga != '-':
        return f'meta-GGA {metagga}'
    if gga and gga != '-':
        return methods_text.GGA_MAP.get(gga, f'GGA={gga}')
    return 'PBE(未显式设 GGA,由 PAW_PBE 赝势隐含)'


def _audit_functional(incars: dict) -> dict:
    pairs: list = []
    for n, t in incars.items():
        if t is None:
            pairs.append((n, None))
            continue
        inc = parse_incar(t)
        pairs.append((n, (str(inc.get('GGA', '')).upper() or '-',
                          str(inc.get('METAGGA', '')).upper() or '-',
                          str(inc.get('LIBXC', '')).upper() or '-')))
    present = [(n, v) for n, v in pairs if v is not None]
    if not present:
        return {'name': '泛函一致性', 'ok': False, 'detail': '无成员 INCAR 可读。'}
    missing = [n for n, v in pairs if v is None]
    if missing:
        return {'name': '泛函一致性', 'ok': False,
                'detail': f'成员 {"、".join(missing)} 无 INCAR,泛函无法核对。'}
    vals = {v for _, v in present}
    if len(vals) == 1:
        key = next(iter(vals))
        return {'name': '泛函一致性', 'ok': True,
                'detail': f'全部成员泛函一致(GGA/METAGGA/LIBXC={key[0]}/{key[1]}/{key[2]},'
                          f'即 {_func_label(key)})。'}
    grp = _group_by_val(present)
    detail = ';'.join(f'{k[0]}/{k[1]}/{k[2]}(成员:{"、".join(v)})' for k, v in grp.items())
    return {'name': '泛函一致性', 'ok': False,
            'detail': f'泛函(GGA/METAGGA/LIBXC)混用,能量不可比:{detail}。'}


def _audit_dispersion(incars: dict) -> dict:
    names = list(incars.keys())
    texts = [incars[n] or '' for n in names]
    if not texts:
        return {'name': '色散校正一致性', 'ok': False, 'detail': '无成员 INCAR 可读。'}
    res = dispersion_audit(texts)                     # 复用 incar_builder 的色散审计
    detail = res['detail']
    if not res['ok']:
        detail += '(成员序:' + '、'.join(f'第{i + 1}份={n}' for i, n in enumerate(names)) + ')'
    return {'name': '色散校正一致性', 'ok': res['ok'], 'detail': detail}


def _audit_reference(project: dict) -> dict:
    mem = (project or {}).get('members') or {}
    has_ref = bool(mem.get('gas_ref'))
    species_refs = bool((project or {}).get('species_refs'))
    if has_ref and species_refs:
        return {'name': '参考态口径', 'ok': False,
                'detail': ('同时设了单一气相参考(gas_ref)与逐物种参考(species_refs),'
                           '两套口径混用会使 ΔE 参照系不一致,须二选一。')}
    if species_refs:
        return {'name': '参考态口径', 'ok': True,
                'detail': '采用逐物种气相参考(species_refs),口径一致。'}
    if has_ref:
        return {'name': '参考态口径', 'ok': True,
                'detail': '采用单一气相参考(gas_ref),口径一致。'}
    return {'name': '参考态口径', 'ok': True,
            'detail': '未设气相参考:ΔE 记为 E(slab+ads)−E(slab)(表面吸附能口径),报告须明示。'}


def _audit_kmesh(recs: list) -> dict:
    if not recs:
        return {'name': 'K 网格档位记录', 'ok': False, 'detail': '无成员可读。'}
    got, missing = [], []
    for r in recs:
        m = manifest_mod.load_manifest(r['dir']) or {}
        k = m.get('kpoints') or (m.get('inputs') or {}).get('kpoints')
        if k:
            got.append(f"{r['name']}={'×'.join(str(x) for x in k)}")
        else:
            missing.append(r['name'])
    if missing:
        return {'name': 'K 网格档位记录', 'ok': False,
                'detail': (f'成员 {"、".join(missing)} 的 job.yaml 未记录 K 网格(复现缺项)。'
                           f'已记录:{"; ".join(got) or "无"}。')}
    return {'name': 'K 网格档位记录', 'ok': True,
            'detail': '全部成员 K 网格已记录:' + '; '.join(got) + '。'}


def _audit_ispin(incars: dict, roles: dict | None = None) -> dict:
    roles = dict(roles or {})
    pairs: list = []
    for n, t in incars.items():
        if t is None:
            pairs.append((n, None, roles.get(n)))
            continue
        v = parse_incar(t).get('ISPIN')
        # 缺 ISPIN → VASP 默认 1
        pairs.append((n, int(v) if isinstance(v, (int, float)) else 1, roles.get(n)))
    present = [(n, v, role) for n, v, role in pairs if v is not None]
    if not present:
        return {'name': 'ISPIN 一致性', 'ok': False, 'detail': '无成员 INCAR 可读。'}
    missing = [n for n, v, _role in pairs if v is None]
    if missing:
        return {'name': 'ISPIN 一致性', 'ok': False,
                'detail': f'成员 {"、".join(missing)} 无 INCAR,自旋设置无法核对。'}
    invalid = [(n, v) for n, v, _role in present if v not in {1, 2}]
    if invalid:
        return {'name': 'ISPIN 一致性', 'ok': False,
                'detail': '存在非法 ISPIN：' + '、'.join(f'{n}={v}' for n, v in invalid) + '。'}

    periodic = [(n, v) for n, v, role in present if role in {'clean_slab', 'config'}]
    references = [(n, v) for n, v, role in present if role == 'gas_ref']
    if periodic:
        periodic_vals = {v for _, v in periodic}
        if len(periodic_vals) != 1:
            grp = _group_by_val(periodic)
            detail = ';'.join(f'ISPIN={k}(成员:{"、".join(v)})'
                              for k, v in grp.items())
            return {'name': 'ISPIN 一致性', 'ok': False,
                    'detail': f'clean slab 与吸附构型自旋口径不一致:{detail}。'}
        value = next(iter(periodic_vals))
        detail = f'周期能量项统一使用 ISPIN={value}'
        if references:
            refs = '、'.join(f'{name}=ISPIN {spin}' for name, spin in references)
            detail += (f'；分子/气相参考 {refs}。不同体系可按各自基态采用'
                       '不同 ISPIN，需保留自旋极化与多初态核对证据。')
        else:
            detail += '，clean slab 与全部吸附构型一致。'
        return {'name': 'ISPIN 一致性', 'ok': True, 'detail': detail}

    vals = {v for _, v, _role in present}
    if len(vals) == 1:
        v = next(iter(vals))
        return {'name': 'ISPIN 一致性', 'ok': True,
                'detail': f'全部成员 ISPIN={v}({"自旋极化" if v == 2 else "非自旋极化"}),一致。'}
    grp = _group_by_val([(n, v) for n, v, _role in present])
    detail = ';'.join(f'ISPIN={k}(成员:{"、".join(v)})' for k, v in grp.items())
    return {'name': 'ISPIN 一致性', 'ok': False,
            'detail': f'ISPIN 混用(自旋口径不一致,能量不可比):{detail}。'}


def consistency_audit(project: dict, *, config: dict | None = None) -> dict:
    """口径稽核:ENCUT/泛函/色散/参考态/K 网格/ISPIN 六项一致性。

    每项 ``{'name','ok','detail'}``(中文);summary 一句话给"可否进稿"。
    返回 ``{'ok','checks':[...],'summary'}``,ok=六项全过。
    """
    recs = _member_records(project)
    incars = {r['name']: _read_file(r['dir'], 'INCAR') for r in recs}
    roles = {r['name']: r['role'] for r in recs}
    checks = [
        _audit_encut(incars),
        _audit_functional(incars),
        _audit_dispersion(incars),
        _audit_reference(project),
        _audit_kmesh(recs),
        _audit_ispin(incars, roles),
    ]
    ok = all(c['ok'] for c in checks)
    if ok:
        summary = '六项口径全部一致,可进入投稿流程。'
    else:
        failed = '、'.join(c['name'] for c in checks if not c['ok'])
        summary = (f'口径稽核未通过({failed}存在不一致/缺失),进稿前必须核对以上项,'
                   f'否则各成员能量不可比、结论不可靠。')
    return {'ok': ok, 'checks': checks, 'summary': summary}


# ── 4. 方法学装订 ─────────────────────────────────────────────────────────────

def _synth_potcar_text(m: dict | None) -> str | None:
    """POTCAR 本体缺失时,用 manifest 记录的真实 TITEL 合成最小文本供 extract_facts 解析。"""
    prov = _potcar_provenance_of(m)
    lines = [f"   TITEL  = {p.get('titel', '')}" for p in prov if p.get('titel')]
    return '\n'.join(lines) if lines else None


def methods_bundle(project: dict, out_dir, *, config: dict | None = None,
                   timestamp: str | None = None) -> dict:
    """方法学装订:双语 Methods 散文 + BibTeX,写 methods_zh.md/methods_en.md/references.bib。

    INCAR 参数从**第一个 DONE 成员真实读取**;成员间不一致时 issues 点名并在文本头部注
    ``[待确认]`` 行。引文缺失(泛函/赝势未知)只告警不编造。返回 ``{'files','issues'}``。
    """
    out_dir = str(out_dir)
    os.makedirs(out_dir, exist_ok=True)
    ts = _ts(timestamp)
    issues: list = []
    recs = _member_records(project)
    done = [r for r in recs
            if (manifest_mod.load_manifest(r['dir']) or {}).get('state') == 'DONE']

    zh_header = en_header = None
    if not done:
        issues.append('[待确认:项目无 DONE 成员,方法学无法从真实参数生成]')
        zh_body = '[待确认:无 DONE 成员,方法学段待计算完成后依据真实 INCAR/KPOINTS/POTCAR 生成]'
        en_body = ('[TO CONFIRM: no DONE member; the methods paragraph will be generated '
                   'from real INCAR/KPOINTS/POTCAR once calculations complete]')
        bib = methods_text.render_bibtex({})          # 仅 VASP 基础引文(恒引,非编造)
    else:
        src = done[0]
        incar_text = _read_file(src['dir'], 'INCAR')
        kpoints_text = _read_file(src['dir'], 'KPOINTS')
        potcar_text = _read_file(src['dir'], 'POTCAR')
        if potcar_text is None:                       # 本体缺 → 退回 manifest 真实 TITEL
            potcar_text = _synth_potcar_text(manifest_mod.load_manifest(src['dir']))
        facts = methods_text.extract_facts(incar_text or '', kpoints_text, potcar_text)
        f = facts['facts']
        for w in facts['warnings']:                   # 缺件降级告警入 issues,绝不编造
            issues.append(f'[待确认:{src["name"]} —— {w}]')
        zh_body = methods_text.render_zh(f)
        en_body = methods_text.render_en(f)
        bib = methods_text.render_bibtex(f)
        audit = consistency_audit(project, config=config)
        if not audit['ok']:
            failed = '、'.join(c['name'] for c in audit['checks'] if not c['ok'])
            zh_header = (f'[待确认:成员间参数不一致({failed});本方法学取自首个 DONE 成员 '
                         f'{src["name"]},请人工核对后定稿]')
            en_header = (f'[TO CONFIRM: parameters differ across members ({failed}); methods '
                         f'taken from first DONE member {src["name"]}, verify before submission]')
            issues.append(zh_header)

    zh_lines = ([zh_header, ''] if zh_header else []) + [
        f'<!-- 生成:vcstudio {__version__} @ {ts} -->', '', '## 方法(Methods)', '', zh_body]
    en_lines = ([en_header, ''] if en_header else []) + [
        f'<!-- generated: vcstudio {__version__} @ {ts} -->', '', '## Methods', '', en_body]

    zh_path = os.path.join(out_dir, 'methods_zh.md')
    en_path = os.path.join(out_dir, 'methods_en.md')
    bib_path = os.path.join(out_dir, 'references.bib')
    Path(zh_path).write_text('\n'.join(zh_lines) + '\n', encoding='utf-8')
    Path(en_path).write_text('\n'.join(en_lines) + '\n', encoding='utf-8')
    Path(bib_path).write_text(bib + '\n', encoding='utf-8')
    return {'files': [zh_path, en_path, bib_path], 'issues': issues}


# ── 5. 一键 Draft-Ready ───────────────────────────────────────────────────────

def _draft_summary_md(name: str, ok: bool, summary: str, report: dict,
                      issues: list, ts: str) -> str:
    """Draft-Ready 总结(首行醒目;稽核不过时一眼可见)。"""
    lines = [('# ✅ ' if ok else '# ⚠️ ')
             + f'{name} Draft-Ready 稽核{"通过" if ok else "未通过"}', '', summary, '',
             '## 口径稽核']
    for c in report['audit']['checks']:
        lines.append(f"- {'✅' if c['ok'] else '❌'} {c['name']}:{c['detail']}")
    lines += ['', '## 产出',
              f"- SI 包:{report['si_package']['zip_path']}"
              f"(ok={report['si_package']['ok']},{len(report['si_package']['contents'])} 项)",
              f"- 表格:{len(report['tables']['files'])} 个文件",
              f"- 方法学:{len(report['methods']['files'])} 个文件", '']
    if issues:
        lines.append(f'## 待确认(共 {len(issues)} 项)')
        lines += [f'- {it}' for it in issues]
        lines.append('')
    lines.append(f'生成时间 {ts};vcstudio {__version__}。'
                 f'铁律:只组织真实数据,缺项以 [待确认] 标注,绝不编造。')
    return '\n'.join(lines) + '\n'


def draft_ready(project: dict, out_dir, **kw) -> dict:
    """一键全套:SI zip + 三表 + 口径稽核 + Methods。

    稽核不通过(或 SI 缺关键输入)仍产出全部工件,但 ok=False,且 DRAFT_READY.md 首行醒目。
    支持 kw:timestamp / config / fed / descriptors。返回
    ``{'ok','report':{...},'issues_total','issues','summary','summary_path'}``。
    """
    out_dir = str(out_dir)
    os.makedirs(out_dir, exist_ok=True)
    timestamp = kw.get('timestamp')
    config = kw.get('config')
    name = str((project or {}).get('name') or 'project')
    report: dict = {}
    issues: list = []

    report['audit'] = consistency_audit(project, config=config)
    si = build_si_package(project, os.path.join(out_dir, f'{name}_SI.zip'),
                          config=config, timestamp=timestamp)
    report['si_package'] = si
    issues += si.get('issues', [])
    tables = make_tables(project, os.path.join(out_dir, 'tables'),
                         fed=kw.get('fed'), descriptors=kw.get('descriptors'),
                         timestamp=timestamp)
    report['tables'] = tables
    issues += tables.get('issues', [])
    methods = methods_bundle(project, os.path.join(out_dir, 'methods'),
                             config=config, timestamp=timestamp)
    report['methods'] = methods
    issues += methods.get('issues', [])

    ok = bool(report['audit']['ok'] and si.get('ok'))
    if ok:
        summary = (f'✅ Draft-Ready 通过:口径稽核一致、SI 包完整,{name} 可进入投稿流程。')
    else:
        reasons = []
        if not report['audit']['ok']:
            reasons.append('口径稽核未通过')
        if not si.get('ok'):
            reasons.append('SI 包缺关键输入文件')
        summary = ('⚠️ 进稿前必须解决以下问题(' + '、'.join(reasons) + '):'
                   + report['audit']['summary'])

    spath = os.path.join(out_dir, 'DRAFT_READY.md')
    Path(spath).write_text(
        _draft_summary_md(name, ok, summary, report, issues, _ts(timestamp)),
        encoding='utf-8')
    return {'ok': ok, 'report': report, 'issues_total': len(issues),
            'issues': issues, 'summary': summary, 'summary_path': spath}
