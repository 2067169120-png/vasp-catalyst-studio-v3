"""Generate publication-style DOCX and PDF reports from one structured model.

The module is deliberately independent from the HTML report pipeline.  Both
renderers consume the same normalized model, embed images directly (so source
data and report files may live on different drives), and publish through a
temporary file in the destination directory followed by :func:`os.replace`.
Microsoft Office is not required.

Canonical model schema::

    {
        "title": "Required report title",
        "subtitle": "Optional subtitle",
        "report_label": "Optional masthead kicker",
        "running_title": "Optional running-header text; defaults to title",
        "footer_label": "Optional footer text",
        "base_dir": "/optional/base/for/relative/figure/paths",
        "metadata": {
            "Project": "Catalyst A",
            "Generated": "2026-07-23",
        },
        # ``metadata`` may instead be [{"label": "...", "value": "..."}].
        "abstract": "A string or a sequence of paragraphs.",
        "sections": [
            {
                "title": "Results",
                "level": 1,  # 1..3, default 1
                "blocks": [
                    {"type": "paragraph", "text": "..."},
                    {"type": "note", "text": "..."},
                    {"type": "bullets", "items": ["...", "..."]},
                    {"type": "numbered", "items": ["...", "..."]},
                    {"type": "heading", "text": "...", "level": 2},
                    {
                        "type": "table",
                        "caption": "Adsorption energies",
                        "columns": ["Site", "Delta E / eV"],
                        "rows": [["top", -1.23], ["bridge", -0.98]],
                        # Optional positive weights, normalized to full width.
                        "column_widths": [2, 1],
                        "note": "Optional table note.",
                    },
                    {
                        "type": "figure",
                        "path": "figures/ladder.png",
                        "caption": "Free-energy pathway",
                        "alt_text": "Optional accessibility description",
                    },
                    {"type": "page_break"},
                ],
            },
        ],
    }

Table rows may also be mappings keyed by the column labels.  Figure paths may
be absolute or relative to ``base_dir``.  Missing, empty, unsupported, or
corrupt images never abort a report; the renderer inserts a clear
``Figure unavailable`` annotation and retains the caption.

The visual system resolves the ``standard_business_brief`` preset with the
named ``scientific_serif`` override: Cambria/SimSun in DOCX and Times with an
embedded Windows CJK font in PDF when available (falling back to the standard
``STSong-Light`` CID font). Geometry, headings, table widths, captions,
headers, footers, and page numbers are all explicit.
"""

from __future__ import annotations

import io
import os
import re
import tempfile
import threading
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any, Callable


STYLE_PRESET = "standard_business_brief"
STYLE_OVERRIDE = "scientific_serif"

_DOCX_LATIN_FONT = "Cambria"
_DOCX_CJK_FONT = "SimSun"
_PDF_CID_FONT = "STSong-Light"
_PDF_CJK_FONT = _PDF_CID_FONT
_PDF_FONT_LOCK = threading.RLock()
_PDF_LATIN_FONT = "Times-Roman"
_PDF_LATIN_FONT_BOLD = "Times-Bold"
_PDF_LATIN_FONT_ITALIC = "Times-Italic"

_CONTENT_WIDTH_DXA = 9360
_TABLE_INDENT_DXA = 120
_CELL_TOP_BOTTOM_DXA = 80
_CELL_START_END_DXA = 120

_HEADING_PREFIX_RE = {
    "Figure": re.compile(r"^(?:figure|fig[.]?|图)\s*\d+", re.IGNORECASE),
    "Table": re.compile(r"^(?:table|表)\s*\d+", re.IGNORECASE),
}
_NUMERIC_RE = re.compile(
    r"^\s*[+-]?(?:\d+(?:[.]\d*)?|[.]\d+)(?:[eE][+-]?\d+)?(?:\s*[%a-zA-Z/]+)?\s*$"
)


class ReportModelError(ValueError):
    """Raised when the structured report model is internally inconsistent."""


def generate_report_documents(
    model: Mapping[str, Any],
    docx_path: str | os.PathLike[str],
    pdf_path: str | os.PathLike[str],
) -> tuple[Path, Path]:
    """Generate DOCX and PDF from one model and return ``(docx, pdf)`` paths."""

    normalized = _normalize_model(model)
    written_docx = _generate_docx_normalized(normalized, Path(docx_path))
    written_pdf = _generate_pdf_normalized(normalized, Path(pdf_path))
    return written_docx, written_pdf


def generate_docx(
    model: Mapping[str, Any],
    output_path: str | os.PathLike[str],
) -> Path:
    """Generate one DOCX atomically and return its destination path."""

    return _generate_docx_normalized(_normalize_model(model), Path(output_path))


def generate_pdf(
    model: Mapping[str, Any],
    output_path: str | os.PathLike[str],
) -> Path:
    """Generate one PDF atomically and return its destination path."""

    return _generate_pdf_normalized(_normalize_model(model), Path(output_path))


def _normalize_model(model: Mapping[str, Any]) -> dict[str, Any]:
    if not isinstance(model, Mapping):
        raise ReportModelError("report model must be a mapping")
    title = _required_text(model.get("title"), "title")
    metadata = _normalize_metadata(model.get("metadata"))
    abstract = _normalize_paragraphs(model.get("abstract"), "abstract")

    raw_sections = model.get("sections", [])
    if raw_sections is None:
        raw_sections = []
    if not _is_sequence(raw_sections):
        raise ReportModelError("sections must be a sequence")

    sections: list[dict[str, Any]] = []
    for index, raw_section in enumerate(raw_sections, 1):
        if not isinstance(raw_section, Mapping):
            raise ReportModelError(f"section {index} must be a mapping")
        level = _heading_level(raw_section.get("level", 1), f"section {index} level")
        raw_blocks = raw_section.get("blocks")
        if raw_blocks is None:
            raw_blocks = _legacy_section_blocks(raw_section)
        if not _is_sequence(raw_blocks):
            raise ReportModelError(f"section {index} blocks must be a sequence")
        blocks = [
            _normalize_block(block, f"section {index} block {block_index}")
            for block_index, block in enumerate(raw_blocks, 1)
        ]
        section_title = _optional_text(raw_section.get("title"))
        sections.append({"title": section_title, "level": level, "blocks": blocks})

    base_dir_raw = model.get("base_dir")
    base_dir = Path(os.fspath(base_dir_raw)) if base_dir_raw not in (None, "") else None
    return {
        "title": title,
        "subtitle": _optional_text(model.get("subtitle")),
        "report_label": _optional_text(model.get("report_label"))
        or "SCIENTIFIC COMPUTATION REPORT",
        "running_title": _optional_text(model.get("running_title")) or title,
        "footer_label": _optional_text(model.get("footer_label"))
        or "VASP Catalyst Studio | Scientific report",
        "metadata": metadata,
        "abstract": abstract,
        "sections": sections,
        "base_dir": base_dir,
    }


def _legacy_section_blocks(section: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Accept a small convenience shape while keeping ``blocks`` canonical."""

    blocks: list[dict[str, Any]] = []
    for text in _normalize_paragraphs(section.get("paragraphs"), "section paragraphs"):
        blocks.append({"type": "paragraph", "text": text})
    for table in section.get("tables") or []:
        if not isinstance(table, Mapping):
            raise ReportModelError("section table must be a mapping")
        blocks.append({"type": "table", **dict(table)})
    for figure in section.get("figures") or []:
        if not isinstance(figure, Mapping):
            raise ReportModelError("section figure must be a mapping")
        blocks.append({"type": "figure", **dict(figure)})
    return blocks


def _normalize_metadata(value: Any) -> list[tuple[str, str]]:
    if value in (None, ""):
        return []
    if isinstance(value, Mapping):
        return [(_as_text(label), _as_text(item)) for label, item in value.items()]
    if not _is_sequence(value):
        raise ReportModelError("metadata must be a mapping or a sequence")

    result: list[tuple[str, str]] = []
    for index, item in enumerate(value, 1):
        if isinstance(item, Mapping):
            label = _required_text(item.get("label"), f"metadata item {index} label")
            result.append((label, _as_text(item.get("value"))))
        elif _is_sequence(item) and len(item) == 2:
            result.append((_as_text(item[0]), _as_text(item[1])))
        else:
            raise ReportModelError(
                f"metadata item {index} must contain label and value"
            )
    return result


def _normalize_paragraphs(value: Any, field: str) -> list[str]:
    if value in (None, ""):
        return []
    if isinstance(value, str):
        return [value]
    if not _is_sequence(value):
        raise ReportModelError(f"{field} must be text or a sequence of text")
    return [_as_text(item) for item in value if item not in (None, "")]


def _normalize_block(block: Any, field: str) -> dict[str, Any]:
    if not isinstance(block, Mapping):
        raise ReportModelError(f"{field} must be a mapping")
    block_type = str(block.get("type") or "paragraph").strip().lower()

    if block_type in {"paragraph", "note"}:
        return {"type": block_type, "text": _as_text(block.get("text"))}
    if block_type in {"bullets", "numbered"}:
        items = block.get("items") or []
        if not _is_sequence(items):
            raise ReportModelError(f"{field} items must be a sequence")
        return {"type": block_type, "items": [_as_text(item) for item in items]}
    if block_type == "heading":
        return {
            "type": "heading",
            "text": _required_text(block.get("text"), f"{field} text"),
            "level": _heading_level(block.get("level", 2), f"{field} level"),
        }
    if block_type == "table":
        return _normalize_table(block, field)
    if block_type == "figure":
        return {
            "type": "figure",
            "path": block.get("path"),
            "caption": _optional_text(block.get("caption")),
            "alt_text": _optional_text(block.get("alt_text")),
        }
    if block_type == "page_break":
        return {"type": "page_break"}
    raise ReportModelError(f"{field} has unsupported type {block_type!r}")


def _normalize_table(block: Mapping[str, Any], field: str) -> dict[str, Any]:
    columns_raw = block.get("columns", block.get("headers"))
    if not _is_sequence(columns_raw) or not columns_raw:
        raise ReportModelError(f"{field} columns must be a non-empty sequence")
    columns = [_as_text(value) for value in columns_raw]

    rows_raw = block.get("rows") or []
    if not _is_sequence(rows_raw):
        raise ReportModelError(f"{field} rows must be a sequence")
    rows: list[list[str]] = []
    for row_index, raw_row in enumerate(rows_raw, 1):
        if isinstance(raw_row, Mapping):
            row = [_as_text(raw_row.get(column)) for column in columns]
        elif _is_sequence(raw_row):
            if len(raw_row) != len(columns):
                raise ReportModelError(
                    f"{field} row {row_index} has {len(raw_row)} cells; "
                    f"expected {len(columns)}"
                )
            row = [_as_text(value) for value in raw_row]
        else:
            raise ReportModelError(f"{field} row {row_index} must be a mapping or sequence")
        rows.append(row)

    weights_raw = block.get("column_widths")
    if weights_raw is None:
        weights = _default_column_weights(columns, rows)
    else:
        if not _is_sequence(weights_raw) or len(weights_raw) != len(columns):
            raise ReportModelError(
                f"{field} column_widths must match the number of columns"
            )
        try:
            weights = [float(value) for value in weights_raw]
        except (TypeError, ValueError) as exc:
            raise ReportModelError(f"{field} column_widths must be numeric") from exc
        if any(value <= 0 for value in weights):
            raise ReportModelError(f"{field} column_widths must all be positive")

    return {
        "type": "table",
        "caption": _optional_text(block.get("caption")),
        "columns": columns,
        "rows": rows,
        "widths_dxa": _normalize_widths(weights),
        "note": _optional_text(block.get("note")),
    }


def _default_column_weights(columns: list[str], rows: list[list[str]]) -> list[float]:
    """Favor descriptive columns without letting long values consume everything."""

    weights: list[float] = []
    for column_index, heading in enumerate(columns):
        lengths = [len(heading)]
        lengths.extend(len(row[column_index]) for row in rows[:100])
        representative = max(lengths, default=8)
        weights.append(float(min(32, max(8, representative))))
    return weights


def _normalize_widths(weights: Sequence[float]) -> list[int]:
    total = sum(weights)
    widths = [max(1, round(_CONTENT_WIDTH_DXA * weight / total)) for weight in weights]
    widths[-1] += _CONTENT_WIDTH_DXA - sum(widths)
    if widths[-1] <= 0:
        raise ReportModelError("column widths cannot be represented at page width")
    return widths


def _column_alignments(rows: list[list[str]], count: int) -> list[str]:
    alignments: list[str] = []
    for column_index in range(count):
        values = [row[column_index] for row in rows if row[column_index].strip()]
        if column_index > 0 and values and all(_NUMERIC_RE.match(value) for value in values):
            alignments.append("right")
        elif column_index > 0 and values and max(map(len, values)) <= 12:
            alignments.append("center")
        else:
            alignments.append("left")
    return alignments


def _heading_level(value: Any, field: str) -> int:
    try:
        level = int(value)
    except (TypeError, ValueError) as exc:
        raise ReportModelError(f"{field} must be 1, 2, or 3") from exc
    if level not in (1, 2, 3):
        raise ReportModelError(f"{field} must be 1, 2, or 3")
    return level


def _required_text(value: Any, field: str) -> str:
    text = _as_text(value).strip()
    if not text:
        raise ReportModelError(f"{field} is required")
    return text


def _optional_text(value: Any) -> str:
    return "" if value is None else _as_text(value).strip()


def _as_text(value: Any) -> str:
    if value is None:
        return "-"
    return str(value)


def _is_sequence(value: Any) -> bool:
    return isinstance(value, Sequence) and not isinstance(value, (str, bytes, bytearray))


def _generate_docx_normalized(model: dict[str, Any], output_path: Path) -> Path:
    return _atomic_build(output_path, lambda temporary: _build_docx(model, temporary))


def _generate_pdf_normalized(model: dict[str, Any], output_path: Path) -> Path:
    return _atomic_build(output_path, lambda temporary: _build_pdf(model, temporary))


def _atomic_build(target: Path, builder: Callable[[Path], None]) -> Path:
    """Build beside ``target`` and replace it only after a complete write."""

    target = target.expanduser()
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(
        prefix=f".{target.name}.",
        suffix=".tmp",
        dir=str(target.parent),
    )
    os.close(fd)
    temporary = Path(temporary_name)
    try:
        builder(temporary)
        if not temporary.is_file() or temporary.stat().st_size == 0:
            raise OSError(f"document builder produced no data for {target.name}")
        with temporary.open("rb+") as stream:
            os.fsync(stream.fileno())
        os.replace(temporary, target)
        return target
    finally:
        try:
            temporary.unlink()
        except FileNotFoundError:
            pass


def _build_docx(model: dict[str, Any], path: Path) -> None:
    try:
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:  # pragma: no cover - dependency error is environment-specific
        raise RuntimeError(
            "DOCX export requires the optional dependency 'python-docx'"
        ) from exc

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    _configure_docx_styles(document, WD_STYLE_TYPE, WD_LINE_SPACING, Pt, RGBColor, qn)
    _configure_docx_header_footer(
        document,
        model,
        WD_ALIGN_PARAGRAPH,
        WD_TAB_ALIGNMENT,
        Inches,
        Pt,
        OxmlElement,
        qn,
    )
    _enable_docx_field_updates(document, OxmlElement, qn)

    document.core_properties.title = model["title"]
    document.core_properties.subject = "Scientific computation report"
    document.core_properties.category = f"{STYLE_PRESET}:{STYLE_OVERRIDE}"

    document.add_paragraph(model["report_label"], style="Scientific Kicker")
    document.add_paragraph(model["title"], style="Title")
    if model["subtitle"]:
        document.add_paragraph(model["subtitle"], style="Scientific Subtitle")
    for label, value in model["metadata"]:
        paragraph = document.add_paragraph(style="Scientific Metadata")
        label_run = paragraph.add_run(f"{label}: ")
        _set_docx_run_font(label_run, 10, bold=True)
        value_run = paragraph.add_run(value)
        _set_docx_run_font(value_run, 10)

    if model["abstract"]:
        document.add_heading("摘要 / Abstract", level=1)
        for text in model["abstract"]:
            document.add_paragraph(text, style="Normal")

    table_number = 0
    figure_number = 0
    for section_model in model["sections"]:
        if section_model["title"]:
            document.add_heading(section_model["title"], level=section_model["level"])
        for block in section_model["blocks"]:
            block_type = block["type"]
            if block_type == "paragraph":
                document.add_paragraph(block["text"], style="Normal")
            elif block_type == "note":
                document.add_paragraph(block["text"], style="Scientific Note")
            elif block_type in {"bullets", "numbered"}:
                style = "List Bullet" if block_type == "bullets" else "List Number"
                for item in block["items"]:
                    document.add_paragraph(item, style=style)
            elif block_type == "heading":
                document.add_heading(block["text"], level=block["level"])
            elif block_type == "table":
                table_number += 1
                _add_docx_table(document, block, table_number)
            elif block_type == "figure":
                figure_number += 1
                _add_docx_figure(
                    document,
                    block,
                    figure_number,
                    model["base_dir"],
                    WD_ALIGN_PARAGRAPH,
                    Inches,
                )
            elif block_type == "page_break":
                document.add_page_break()

    document.save(str(path))


def _configure_docx_styles(
    document: Any,
    WD_STYLE_TYPE: Any,
    WD_LINE_SPACING: Any,
    Pt: Any,
    RGBColor: Any,
    qn: Any,
) -> None:
    from docx.oxml import OxmlElement
    from docx.shared import Inches

    styles = document.styles
    style_specs = {
        "Normal": (11, "000000", False, 0, 6, 1.10),
        "Title": (24, "0B2545", True, 0, 6, 1.0),
        "Heading 1": (16, "2E74B5", True, 16, 8, 1.0),
        "Heading 2": (13, "2E74B5", True, 12, 6, 1.0),
        "Heading 3": (12, "1F4D78", True, 8, 4, 1.0),
        "Caption": (9, "4B5563", False, 4, 8, 1.0),
        "Header": (8.5, "6B7280", False, 0, 0, 1.0),
        "Footer": (8.5, "6B7280", False, 0, 0, 1.0),
        "List Bullet": (11, "000000", False, 0, 8, 1.167),
        "List Number": (11, "000000", False, 0, 8, 1.167),
    }
    custom_specs = {
        "Scientific Kicker": (9, "2E74B5", True, 0, 3, 1.0),
        "Scientific Subtitle": (12, "4B5563", False, 0, 14, 1.0),
        "Scientific Metadata": (10, "374151", False, 0, 2, 1.0),
        "Scientific Note": (9.5, "4B5563", False, 4, 6, 1.10),
    }
    for name, spec in custom_specs.items():
        if name not in styles:
            styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style_specs[name] = spec

    for name, (size, color, bold, before, after, line_spacing) in style_specs.items():
        style = styles[name]
        style.font.name = _DOCX_LATIN_FONT
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        _set_docx_style_fonts(style, qn)
        paragraph_format = style.paragraph_format
        paragraph_format.space_before = Pt(before)
        paragraph_format.space_after = Pt(after)
        paragraph_format.line_spacing = line_spacing
        paragraph_format.widow_control = True

    for name in ("Heading 1", "Heading 2", "Heading 3", "Scientific Kicker"):
        styles[name].paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.keep_with_next = True
    styles["Scientific Subtitle"].paragraph_format.keep_with_next = True
    styles["Caption"].paragraph_format.keep_together = True
    styles["Scientific Note"].paragraph_format.keep_together = True

    for name in ("List Bullet", "List Number"):
        paragraph_format = styles[name].paragraph_format
        paragraph_format.left_indent = Inches(0.5)
        paragraph_format.first_line_indent = Inches(-0.25)
        paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE

    note_ppr = styles["Scientific Note"].element.get_or_add_pPr()
    shading = note_ppr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        note_ppr.append(shading)
    shading.set(qn("w:fill"), "F4F6F9")


def _set_docx_style_fonts(style: Any, qn: Any) -> None:
    from docx.oxml import OxmlElement

    run_properties = style.element.get_or_add_rPr()
    fonts = run_properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), _DOCX_LATIN_FONT)
    fonts.set(qn("w:hAnsi"), _DOCX_LATIN_FONT)
    fonts.set(qn("w:cs"), _DOCX_LATIN_FONT)
    fonts.set(qn("w:eastAsia"), _DOCX_CJK_FONT)


def _set_docx_run_font(
    run: Any,
    size: float | None = None,
    *,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = _DOCX_LATIN_FONT
    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), _DOCX_LATIN_FONT)
    fonts.set(qn("w:hAnsi"), _DOCX_LATIN_FONT)
    fonts.set(qn("w:cs"), _DOCX_LATIN_FONT)
    fonts.set(qn("w:eastAsia"), _DOCX_CJK_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _configure_docx_header_footer(
    document: Any,
    model: dict[str, Any],
    WD_ALIGN_PARAGRAPH: Any,
    WD_TAB_ALIGNMENT: Any,
    Inches: Any,
    Pt: Any,
    OxmlElement: Any,
    qn: Any,
) -> None:
    section = document.sections[0]
    header_paragraph = section.header.paragraphs[0]
    header_paragraph.style = document.styles["Header"]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header_paragraph.add_run(model["running_title"])
    _set_docx_run_font(header_run, 8.5)

    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.style = document.styles["Footer"]
    footer_paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.5), WD_TAB_ALIGNMENT.RIGHT
    )
    label_run = footer_paragraph.add_run(model["footer_label"])
    _set_docx_run_font(label_run, 8.5)
    page_label_run = footer_paragraph.add_run("\tPage ")
    _set_docx_run_font(page_label_run, 8.5)
    _append_docx_page_field(
        footer_paragraph,
        OxmlElement,
        qn,
    )


def _append_docx_page_field(paragraph: Any, OxmlElement: Any, qn: Any) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    _set_docx_run_font(run, 8.5)
    run._r.extend((begin, instruction, separate, display, end))


def _enable_docx_field_updates(document: Any, OxmlElement: Any, qn: Any) -> None:
    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def _add_docx_table(document: Any, block: dict[str, Any], number: int) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    caption = _numbered_caption("Table", number, block["caption"])
    caption_paragraph = document.add_paragraph(caption, style="Caption")
    caption_paragraph.paragraph_format.keep_with_next = True

    columns = block["columns"]
    rows = block["rows"]
    widths = block["widths_dxa"]
    alignments = _column_alignments(rows, len(columns))
    table = document.add_table(rows=1, cols=len(columns))
    table.autofit = False
    table.style = "Table Grid"
    _set_docx_table_geometry(table, widths, OxmlElement, qn)

    for column_index, heading in enumerate(columns):
        cell = table.rows[0].cells[column_index]
        _set_docx_cell_text(
            cell,
            heading,
            "center",
            header=True,
            WD_ALIGN_PARAGRAPH=WD_ALIGN_PARAGRAPH,
        )
        _set_docx_cell_fill(cell, "F2F4F7", OxmlElement, qn)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _repeat_docx_table_header(table.rows[0], OxmlElement, qn)

    for values in rows:
        cells = table.add_row().cells
        for column_index, value in enumerate(values):
            cell = cells[column_index]
            _set_docx_cell_text(
                cell,
                value,
                alignments[column_index],
                header=False,
                WD_ALIGN_PARAGRAPH=WD_ALIGN_PARAGRAPH,
            )
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_docx_table_geometry(table, widths, OxmlElement, qn)

    if block["note"]:
        document.add_paragraph(block["note"], style="Scientific Note")


def _set_docx_table_geometry(
    table: Any,
    widths: Sequence[int],
    OxmlElement: Any,
    qn: Any,
) -> None:
    table_properties = table._tbl.tblPr
    table_width = table_properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.insert(0, table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(_CONTENT_WIDTH_DXA))

    table_indent = table_properties.find(qn("w:tblInd"))
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        table_properties.append(table_indent)
    table_indent.set(qn("w:type"), "dxa")
    table_indent.set(qn("w:w"), str(_TABLE_INDENT_DXA))

    layout = table_properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    borders = table_properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "B7C1CC")

    grid_columns = list(table._tbl.tblGrid)
    while len(grid_columns) < len(widths):
        column = OxmlElement("w:gridCol")
        table._tbl.tblGrid.append(column)
        grid_columns.append(column)
    while len(grid_columns) > len(widths):
        table._tbl.tblGrid.remove(grid_columns.pop())
    for column, width in zip(grid_columns, widths):
        column.set(qn("w:w"), str(width))

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width = cell_properties.get_or_add_tcW()
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(width))
            _set_docx_cell_margins(cell_properties, OxmlElement, qn)


def _set_docx_cell_margins(cell_properties: Any, OxmlElement: Any, qn: Any) -> None:
    margins = cell_properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)
    values = {
        "top": _CELL_TOP_BOTTOM_DXA,
        "bottom": _CELL_TOP_BOTTOM_DXA,
        "start": _CELL_START_END_DXA,
        "end": _CELL_START_END_DXA,
    }
    for side, value in values.items():
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _set_docx_cell_fill(cell: Any, fill: str, OxmlElement: Any, qn: Any) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _repeat_docx_table_header(row: Any, OxmlElement: Any, qn: Any) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = properties.find(qn("w:tblHeader"))
    if repeat is None:
        repeat = OxmlElement("w:tblHeader")
        properties.append(repeat)
    repeat.set(qn("w:val"), "true")


def _set_docx_cell_text(
    cell: Any,
    text: str,
    alignment: str,
    *,
    header: bool,
    WD_ALIGN_PARAGRAPH: Any,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = 0
    paragraph.paragraph_format.space_after = 0
    paragraph.paragraph_format.line_spacing = 1.05
    paragraph.alignment = {
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
    }.get(alignment, WD_ALIGN_PARAGRAPH.LEFT)
    run = paragraph.add_run(text)
    _set_docx_run_font(run, 9.5, bold=header)


def _add_docx_figure(
    document: Any,
    block: dict[str, Any],
    number: int,
    base_dir: Path | None,
    WD_ALIGN_PARAGRAPH: Any,
    Inches: Any,
) -> None:
    from docx.image.image import Image as DocxImage
    from docx.shared import Emu

    image_path = _figure_path(block.get("path"), base_dir)
    image_bytes = _read_figure_bytes(image_path)
    embedded = False
    if image_bytes is not None:
        try:
            image = DocxImage.from_blob(image_bytes)
            native_width = int(image.width)
            native_height = int(image.height)
            if native_width <= 0 or native_height <= 0:
                raise ValueError("image has no usable dimensions")
            maximum_width = int(Inches(6.25))
            maximum_height = int(Inches(5.8))
            scale = min(
                maximum_width / native_width,
                maximum_height / native_height,
                1.0,
            )
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_together = True
            run = paragraph.add_run()
            shape = run.add_picture(
                io.BytesIO(image_bytes),
                width=Emu(round(native_width * scale)),
                height=Emu(round(native_height * scale)),
            )
            alt_text = block["alt_text"] or block["caption"]
            if alt_text:
                shape._inline.docPr.set("descr", alt_text)
            embedded = True
        except Exception:
            embedded = False
    if not embedded:
        document.add_paragraph(
            _figure_unavailable_text(image_path),
            style="Scientific Note",
        )

    caption = _numbered_caption("Figure", number, block["caption"])
    document.add_paragraph(caption, style="Caption")


def _register_pdf_cjk_font(pdfmetrics: Any, UnicodeCIDFont: Any, TTFont: Any) -> str:
    """Prefer an embedded Windows CJK font; fall back to the standard CID font.

    The desktop build targets Windows, where these fonts normally exist.  A
    user-supplied path can be provided through ``VCSTUDIO_REPORT_CJK_FONT``.
    Embedding the selected TrueType/TTC subset makes the resulting PDF portable
    to readers that do not install Adobe's Asian font pack.
    """
    global _PDF_CJK_FONT

    with _PDF_FONT_LOCK:
        try:
            pdfmetrics.getFont("VCS-CJK")
        except KeyError:
            pass
        else:
            _PDF_CJK_FONT = "VCS-CJK"
            return _PDF_CJK_FONT

        windows_root = os.environ.get("WINDIR") or os.environ.get("SystemRoot") or ""
        font_root = Path(windows_root) / "Fonts" if windows_root else None
        configured = str(os.environ.get("VCSTUDIO_REPORT_CJK_FONT") or "").strip()
        packaged = Path(__file__).resolve().parents[1] / "gui_web" / "assets" / "fonts"
        candidates = [
            Path(configured).expanduser() if configured else None,
            packaged / "noto-serif-sc.ttf",
            packaged / "noto-sans-sc.ttf",
        ]
        if font_root is not None:
            candidates.extend(
                font_root / name
                for name in ("simsun.ttc", "msyh.ttc", "simhei.ttf", "Deng.ttf")
            )
        for candidate in candidates:
            if candidate is None or not candidate.is_file():
                continue
            try:
                pdfmetrics.registerFont(
                    TTFont("VCS-CJK", str(candidate), subfontIndex=0))
            except Exception:                             # noqa: BLE001 尝试下一字体
                continue
            _PDF_CJK_FONT = "VCS-CJK"
            return _PDF_CJK_FONT

        try:
            pdfmetrics.getFont(_PDF_CID_FONT)
        except KeyError:
            pdfmetrics.registerFont(UnicodeCIDFont(_PDF_CID_FONT))
        _PDF_CJK_FONT = _PDF_CID_FONT
        return _PDF_CJK_FONT


def _build_pdf(model: dict[str, Any], path: Path) -> None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.lib.utils import ImageReader
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import (
            Image,
            KeepTogether,
            ListFlowable,
            ListItem,
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Table,
            TableStyle,
        )
    except ImportError as exc:  # pragma: no cover - dependency error is environment-specific
        raise RuntimeError(
            "PDF export requires the optional dependency 'reportlab'"
        ) from exc

    _register_pdf_cjk_font(pdfmetrics, UnicodeCIDFont, TTFont)

    palette = {
        "blue": colors.HexColor("#2E74B5"),
        "dark_blue": colors.HexColor("#1F4D78"),
        "ink": colors.HexColor("#0B2545"),
        "muted": colors.HexColor("#4B5563"),
        "grid": colors.HexColor("#B7C1CC"),
        "table_fill": colors.HexColor("#F2F4F7"),
        "note_fill": colors.HexColor("#F4F6F9"),
    }
    styles = _pdf_styles(ParagraphStyle, TA_LEFT, TA_CENTER, TA_RIGHT, palette)
    document = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        leftMargin=inch,
        rightMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
        title=model["title"],
        subject="Scientific computation report",
        author="VASP Catalyst Studio",
        creator=f"{STYLE_PRESET}:{STYLE_OVERRIDE}",
    )

    story: list[Any] = [
        Paragraph(_pdf_text(model["report_label"]), styles["Kicker"]),
        Paragraph(_pdf_text(model["title"]), styles["Title"]),
    ]
    if model["subtitle"]:
        story.append(Paragraph(_pdf_text(model["subtitle"]), styles["Subtitle"]))
    for label, value in model["metadata"]:
        metadata_markup = (
            f'<font color="#1F4D78">{_pdf_text(label)}:</font> {_pdf_text(value)}'
        )
        story.append(Paragraph(metadata_markup, styles["Metadata"]))

    if model["abstract"]:
        story.append(Paragraph(_pdf_text("摘要 / Abstract"), styles["Heading1"]))
        for text in model["abstract"]:
            story.append(Paragraph(_pdf_text(text), styles["Body"]))

    table_number = 0
    figure_number = 0
    for section_model in model["sections"]:
        if section_model["title"]:
            story.append(
                Paragraph(
                    _pdf_text(section_model["title"]),
                    styles[f"Heading{section_model['level']}"],
                )
            )
        for block in section_model["blocks"]:
            block_type = block["type"]
            if block_type == "paragraph":
                story.append(Paragraph(_pdf_text(block["text"]), styles["Body"]))
            elif block_type == "note":
                story.append(Paragraph(_pdf_text(block["text"]), styles["Note"]))
            elif block_type in {"bullets", "numbered"}:
                items = [
                    ListItem(
                        Paragraph(_pdf_text(item), styles["ListBody"]),
                        leftIndent=0,
                    )
                    for item in block["items"]
                ]
                if items:
                    list_options = {
                        "bulletType": "bullet" if block_type == "bullets" else "1",
                        "leftIndent": 36,
                        "bulletFontName": _PDF_LATIN_FONT,
                        "bulletFontSize": 9,
                        "bulletOffsetY": 1,
                        "spaceAfter": 8,
                    }
                    if block_type == "numbered":
                        list_options["start"] = 1
                    story.append(
                        ListFlowable(
                            items,
                            **list_options,
                        )
                    )
            elif block_type == "heading":
                story.append(
                    Paragraph(
                        _pdf_text(block["text"]),
                        styles[f"Heading{block['level']}"],
                    )
                )
            elif block_type == "table":
                table_number += 1
                story.extend(
                    _pdf_table_flowables(
                        block,
                        table_number,
                        styles,
                        colors,
                        Table,
                        TableStyle,
                    )
                )
            elif block_type == "figure":
                figure_number += 1
                story.extend(
                    _pdf_figure_flowables(
                        block,
                        figure_number,
                        model["base_dir"],
                        styles,
                        ImageReader,
                        Image,
                        Paragraph,
                        KeepTogether,
                    )
                )
            elif block_type == "page_break":
                story.append(PageBreak())

    def draw_page(canvas: Any, _doc: Any) -> None:
        canvas.saveState()
        canvas.setFillColor(palette["muted"])
        header = _fit_pdf_text(
            model["running_title"],
            6.5 * inch,
            8.5,
            pdfmetrics,
        )
        _draw_pdf_mixed_string(
            canvas,
            inch,
            letter[1] - 0.492 * inch,
            header,
            8.5,
            pdfmetrics,
        )
        footer = _fit_pdf_text(
            model["footer_label"],
            5.2 * inch,
            8.5,
            pdfmetrics,
        )
        _draw_pdf_mixed_string(
            canvas,
            inch,
            0.492 * inch,
            footer,
            8.5,
            pdfmetrics,
        )
        canvas.setFont(_PDF_LATIN_FONT, 8.5)
        canvas.drawRightString(
            letter[0] - inch,
            0.492 * inch,
            f"Page {canvas.getPageNumber()}",
        )
        canvas.restoreState()

    document.build(story, onFirstPage=draw_page, onLaterPages=draw_page)


def _pdf_styles(
    ParagraphStyle: Any,
    TA_LEFT: Any,
    TA_CENTER: Any,
    TA_RIGHT: Any,
    palette: Mapping[str, Any],
) -> dict[str, Any]:
    return {
        "Body": ParagraphStyle(
            "ScientificBody",
            fontName=_PDF_LATIN_FONT,
            fontSize=11,
            leading=12.1,
            alignment=TA_LEFT,
            spaceBefore=0,
            spaceAfter=6,
            textColor=palette["ink"],
            allowWidows=0,
            allowOrphans=0,
        ),
        "Title": ParagraphStyle(
            "ScientificTitle",
            fontName=_PDF_LATIN_FONT_BOLD,
            fontSize=24,
            leading=28,
            alignment=TA_LEFT,
            spaceBefore=0,
            spaceAfter=6,
            textColor=palette["ink"],
            keepWithNext=1,
        ),
        "Subtitle": ParagraphStyle(
            "ScientificSubtitle",
            fontName=_PDF_LATIN_FONT_ITALIC,
            fontSize=12,
            leading=15,
            alignment=TA_LEFT,
            spaceBefore=0,
            spaceAfter=14,
            textColor=palette["muted"],
            keepWithNext=1,
        ),
        "Kicker": ParagraphStyle(
            "ScientificKicker",
            fontName=_PDF_LATIN_FONT_BOLD,
            fontSize=9,
            leading=11,
            alignment=TA_LEFT,
            spaceBefore=0,
            spaceAfter=3,
            textColor=palette["blue"],
            keepWithNext=1,
        ),
        "Metadata": ParagraphStyle(
            "ScientificMetadata",
            fontName=_PDF_LATIN_FONT,
            fontSize=10,
            leading=12,
            alignment=TA_LEFT,
            spaceBefore=0,
            spaceAfter=2,
            textColor=palette["muted"],
        ),
        "Heading1": ParagraphStyle(
            "ScientificHeading1",
            fontName=_PDF_LATIN_FONT_BOLD,
            fontSize=16,
            leading=19,
            alignment=TA_LEFT,
            spaceBefore=16,
            spaceAfter=8,
            textColor=palette["blue"],
            keepWithNext=1,
        ),
        "Heading2": ParagraphStyle(
            "ScientificHeading2",
            fontName=_PDF_LATIN_FONT_BOLD,
            fontSize=13,
            leading=16,
            alignment=TA_LEFT,
            spaceBefore=12,
            spaceAfter=6,
            textColor=palette["blue"],
            keepWithNext=1,
        ),
        "Heading3": ParagraphStyle(
            "ScientificHeading3",
            fontName=_PDF_LATIN_FONT_BOLD,
            fontSize=12,
            leading=15,
            alignment=TA_LEFT,
            spaceBefore=8,
            spaceAfter=4,
            textColor=palette["dark_blue"],
            keepWithNext=1,
        ),
        "Caption": ParagraphStyle(
            "ScientificCaption",
            fontName=_PDF_LATIN_FONT_ITALIC,
            fontSize=9,
            leading=11,
            alignment=TA_CENTER,
            spaceBefore=4,
            spaceAfter=8,
            textColor=palette["muted"],
        ),
        "TableCaption": ParagraphStyle(
            "ScientificTableCaption",
            fontName=_PDF_LATIN_FONT_ITALIC,
            fontSize=9,
            leading=11,
            alignment=TA_LEFT,
            spaceBefore=4,
            spaceAfter=4,
            textColor=palette["muted"],
            keepWithNext=1,
        ),
        "Note": ParagraphStyle(
            "ScientificNote",
            fontName=_PDF_LATIN_FONT,
            fontSize=9.5,
            leading=11.5,
            alignment=TA_LEFT,
            spaceBefore=4,
            spaceAfter=6,
            leftIndent=6,
            rightIndent=6,
            borderPadding=(5, 6, 5, 6),
            backColor=palette["note_fill"],
            textColor=palette["muted"],
        ),
        "ListBody": ParagraphStyle(
            "ScientificListBody",
            fontName=_PDF_LATIN_FONT,
            fontSize=11,
            leading=14,
            alignment=TA_LEFT,
            spaceBefore=0,
            spaceAfter=0,
            textColor=palette["ink"],
        ),
        "TableHeader": ParagraphStyle(
            "ScientificTableHeader",
            fontName=_PDF_LATIN_FONT_BOLD,
            fontSize=9.5,
            leading=11,
            alignment=TA_CENTER,
            spaceBefore=0,
            spaceAfter=0,
            textColor=palette["ink"],
        ),
        "TableBodyLeft": ParagraphStyle(
            "ScientificTableBodyLeft",
            fontName=_PDF_LATIN_FONT,
            fontSize=9.5,
            leading=11,
            alignment=TA_LEFT,
            spaceBefore=0,
            spaceAfter=0,
            textColor=palette["ink"],
        ),
        "TableBodyCenter": ParagraphStyle(
            "ScientificTableBodyCenter",
            fontName=_PDF_LATIN_FONT,
            fontSize=9.5,
            leading=11,
            alignment=TA_CENTER,
            spaceBefore=0,
            spaceAfter=0,
            textColor=palette["ink"],
        ),
        "TableBodyRight": ParagraphStyle(
            "ScientificTableBodyRight",
            fontName=_PDF_LATIN_FONT,
            fontSize=9.5,
            leading=11,
            alignment=TA_RIGHT,
            spaceBefore=0,
            spaceAfter=0,
            textColor=palette["ink"],
        ),
    }


def _pdf_table_flowables(
    block: dict[str, Any],
    number: int,
    styles: Mapping[str, Any],
    colors: Any,
    Table: Any,
    TableStyle: Any,
) -> list[Any]:
    from reportlab.platypus import Paragraph

    caption = Paragraph(
        _pdf_text(_numbered_caption("Table", number, block["caption"])),
        styles["TableCaption"],
    )
    alignments = _column_alignments(block["rows"], len(block["columns"]))
    header = [
        Paragraph(_pdf_text(value), styles["TableHeader"])
        for value in block["columns"]
    ]
    body: list[list[Any]] = []
    for row in block["rows"]:
        rendered_row = []
        for column_index, value in enumerate(row):
            style_name = {
                "right": "TableBodyRight",
                "center": "TableBodyCenter",
            }.get(alignments[column_index], "TableBodyLeft")
            rendered_row.append(Paragraph(_pdf_text(value), styles[style_name]))
        body.append(rendered_row)

    widths = [width / 20 for width in block["widths_dxa"]]
    table = Table(
        [header, *body],
        colWidths=widths,
        repeatRows=1,
        hAlign="LEFT",
        splitByRow=1,
    )
    commands: list[tuple[Any, ...]] = [
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B7C1CC")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F4F7")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    table.setStyle(TableStyle(commands))
    result: list[Any] = [caption, table]
    if block["note"]:
        result.append(Paragraph(_pdf_text(block["note"]), styles["Note"]))
    return result


def _pdf_figure_flowables(
    block: dict[str, Any],
    number: int,
    base_dir: Path | None,
    styles: Mapping[str, Any],
    ImageReader: Any,
    Image: Any,
    Paragraph: Any,
    KeepTogether: Any,
) -> list[Any]:
    image_path = _figure_path(block.get("path"), base_dir)
    image_bytes = _read_figure_bytes(image_path)
    caption = Paragraph(
        _pdf_text(_numbered_caption("Figure", number, block["caption"])),
        styles["Caption"],
    )
    if image_bytes is None:
        return [
            Paragraph(_pdf_text(_figure_unavailable_text(image_path)), styles["Note"]),
            caption,
        ]

    try:
        probe_stream = io.BytesIO(image_bytes)
        reader = ImageReader(probe_stream)
        width, height = reader.getSize()
        reader.getRGBData()
        if width <= 0 or height <= 0:
            raise ValueError("image has no usable dimensions")
        maximum_width = 6.25 * 72
        maximum_height = 5.8 * 72
        scale = min(maximum_width / width, maximum_height / height, 1.0)
        image = Image(
            io.BytesIO(image_bytes),
            width=width * scale,
            height=height * scale,
        )
        image.hAlign = "CENTER"
        return [KeepTogether([image, caption])]
    except Exception:
        return [
            Paragraph(_pdf_text(_figure_unavailable_text(image_path)), styles["Note"]),
            caption,
        ]


def _figure_path(value: Any, base_dir: Path | None) -> Path | None:
    if value in (None, ""):
        return None
    try:
        path = Path(os.fspath(value)).expanduser()
    except TypeError:
        return None
    if not path.is_absolute() and base_dir is not None:
        path = base_dir / path
    return path


def _read_figure_bytes(path: Path | None) -> bytes | None:
    if path is None:
        return None
    try:
        if not path.is_file():
            return None
        data = path.read_bytes()
    except OSError:
        return None
    return data or None


def _figure_unavailable_text(path: Path | None) -> str:
    name = path.name if path is not None and path.name else "unspecified image"
    return (
        f'Figure unavailable: "{name}" is missing, empty, unsupported, or unreadable.'
    )


def _numbered_caption(kind: str, number: int, caption: str) -> str:
    clean_caption = caption.strip()
    if clean_caption and _HEADING_PREFIX_RE[kind].match(clean_caption):
        return clean_caption
    if clean_caption:
        return f"{kind} {number}. {clean_caption}"
    return f"{kind} {number}."


def _pdf_text(value: Any) -> str:
    from xml.sax.saxutils import escape

    text = _as_text(value).replace("\r\n", "\n").replace("\r", "\n")
    fragments: list[str] = []
    for font_name, chunk in _pdf_font_runs(text):
        escaped = escape(chunk).replace("\n", "<br/>")
        if font_name == _PDF_CJK_FONT:
            fragments.append(f'<font name="{_PDF_CJK_FONT}">{escaped}</font>')
        else:
            fragments.append(escaped)
    return "".join(fragments)


def _fit_pdf_text(
    text: str,
    maximum_width: float,
    font_size: float,
    pdfmetrics: Any,
) -> str:
    if _pdf_mixed_width(text, font_size, pdfmetrics) <= maximum_width:
        return text
    suffix = "..."
    low, high = 0, len(text)
    while low < high:
        middle = (low + high + 1) // 2
        candidate = text[:middle].rstrip() + suffix
        if _pdf_mixed_width(candidate, font_size, pdfmetrics) <= maximum_width:
            low = middle
        else:
            high = middle - 1
    return text[:low].rstrip() + suffix


def _pdf_font_runs(text: str) -> list[tuple[str, str]]:
    """Split text so CJK/non-WinAnsi glyphs use the registered CID font."""

    if not text:
        return [(_PDF_LATIN_FONT, "")]
    runs: list[tuple[str, str]] = []
    current_font = _PDF_CJK_FONT if _requires_pdf_cid(text[0]) else _PDF_LATIN_FONT
    current: list[str] = []
    for character in text:
        font = _PDF_CJK_FONT if _requires_pdf_cid(character) else _PDF_LATIN_FONT
        if font != current_font and current:
            runs.append((current_font, "".join(current)))
            current = []
            current_font = font
        current.append(character)
    if current:
        runs.append((current_font, "".join(current)))
    return runs


def _requires_pdf_cid(character: str) -> bool:
    if character in "\r\n\t":
        return False
    try:
        character.encode("cp1252")
    except UnicodeEncodeError:
        return True
    return False


def _pdf_mixed_width(text: str, font_size: float, pdfmetrics: Any) -> float:
    return sum(
        pdfmetrics.stringWidth(chunk, font_name, font_size)
        for font_name, chunk in _pdf_font_runs(text)
    )


def _draw_pdf_mixed_string(
    canvas: Any,
    x: float,
    y: float,
    text: str,
    font_size: float,
    pdfmetrics: Any,
) -> None:
    cursor = x
    for font_name, chunk in _pdf_font_runs(text):
        canvas.setFont(font_name, font_size)
        canvas.drawString(cursor, y, chunk)
        cursor += pdfmetrics.stringWidth(chunk, font_name, font_size)
