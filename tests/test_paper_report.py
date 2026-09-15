"""Unified paper report renderer tests (offline, no Office dependency)."""
from __future__ import annotations

import base64
import copy
import hashlib
import json
import os
import re
import subprocess
import sys
import time
import zipfile
from pathlib import Path
from types import MappingProxyType

import pytest

from vcstudio.project import paper_report


def _png(path: Path, color=(37, 99, 235)) -> Path:
    from PIL import Image

    image = Image.new("RGB", (640, 360), color)
    image.save(path)
    return path


def _model(tmp_path: Path) -> dict:
    figure = _png(tmp_path / "step plot.png")
    return {
        "schema": "vcstudio.research-report/v1",
        "report_id": "report-screening-001",
        "report_kind": "diagnostic",
        "scientific_qualification": "diagnostic",
        "input_fingerprint": "input-fingerprint-001",
        "preset_id": "scientific-review",
        "title": "Adsorption Screening Report",
        "subtitle": "A reproducible multi-catalyst comparison",
        "kicker": "Computational Catalysis",
        "metadata": {
            "Project": "Li-S catalyst screen",
            "Date": "2026-07-24",
            "Status": "Validated data",
        },
        "executive_summary": (
            "The shared report model keeps numerical statements aligned across all formats."
        ),
        "key_findings": [
            {"title": "Balanced binding", "text": "Candidate A merits further calculations."},
            "Near ties below 0.2 eV should not be over-ranked.",
        ],
        "candidate_evaluations": [
            {
                "candidate": "Candidate A",
                "status": "Advance",
                "evidence": "Moderate long-chain binding",
            },
            {
                "candidate": "Candidate B",
                "status": "Conditional",
                "evidence": "Check Li2S decomposition",
            },
        ],
        "adsorption_table": {
            "title": "Stable adsorption configurations",
            "caption": "Negative values indicate exothermic adsorption.",
            "columns": [
                {"key": "species", "label": "Species"},
                {"key": "energy", "label": "Eads (eV)", "align": "right"},
                {"key": "decision", "label": "Screening decision"},
            ],
            "rows": [
                {"species": "Li2S8", "energy": -1.04, "decision": "Advance"},
                {"species": "Li2S", "energy": -2.81, "decision": "Check kinetics"},
            ],
        },
        "comparison_table": [
            {"catalyst": "Candidate A", "UL (V)": 1.55, "PDS": "Li2S2 -> Li2S"},
            {"catalyst": "Candidate B", "UL (V)": 1.31, "PDS": "Li2S4 -> Li2S2"},
        ],
        "figures": [
            {
                "path": figure,
                "title": "Free-energy pathways",
                "caption": "All candidate paths use the same reaction-state order.",
                "alt": "Overlaid free-energy pathways",
            }
        ],
        "methods": [
            {
                "title": "Electronic structure",
                "text": "Method evidence is read from each managed calculation.",
            }
        ],
        "limitations": ["Solvation and activation barriers are not included."],
        "recommendations": [
            "Run Li2S decomposition NEB for Candidate A.",
            "Retain the raw method-consistency evidence.",
        ],
    }


def _validated_model(
    tmp_path: Path,
    *,
    formats=("html", "docx", "pdf"),
    requested_kind="final",
    effective_kind="final",
    locale="en-US",
    outline=None,
    theme_id="academic-a4",
    preset_id="scientific-review",
) -> dict:
    """Return a render model with a real, fully bound validation chain."""
    from vcstudio.project.report_contracts import (
        ReportSnapshot,
        ReportSpec,
        ValidationCheck,
        ValidationResult,
    )

    model = _model(tmp_path)
    model["locale"] = locale
    fingerprint = model["input_fingerprint"]
    spec_kwargs = {}
    if outline is not None:
        spec_kwargs["outline"] = tuple(outline)
    spec = ReportSpec(
        preset_id=preset_id,
        requested_kind=requested_kind,
        locale=locale,
        formats=tuple(formats),
        scope={"kind": "project", "project_ids": ["project-1"], "job_ids": []},
        theme_id=theme_id,
        template_ref={"id": f"vcstudio-{theme_id}", "version": "1"},
        **spec_kwargs,
    )
    snapshot = ReportSnapshot(
        spec_sha256=spec.semantic_sha256,
        input_fingerprint=fingerprint,
        resolved_scope={"project_ids": ["project-1"], "job_ids": []},
        sources=({"source_id": "project:1", "locator": "C:/work/project.yaml"},),
    )
    is_final = effective_kind == "final"
    qualification = "adsorption_result_verified" if is_final else "diagnostic"
    model.update(
        locale=locale,
        preset_id=preset_id,
        report_kind=effective_kind,
        scientific_qualification=qualification,
        template_ref=spec.template_ref,
    )
    report_model_sha256 = paper_report.report_content_sha256(
        model, outline=spec.outline)
    validation = ValidationResult(
        spec_sha256=spec.semantic_sha256,
        snapshot_sha256=snapshot.semantic_sha256,
        validator={"id": "project-final-report-gate", "version": "1"},
        status="passed",
        effective_kind=effective_kind,
        final_allowed=is_final,
        scientific_qualification=qualification,
        report_model_sha256=report_model_sha256,
        checks=(ValidationCheck(id="delivery-gate", status="pass"),),
    )
    model.update(
        report_spec=spec.to_dict(),
        report_snapshot=snapshot.to_dict(),
        validation=validation.to_dict(),
    )
    return model


def _rebind_report_content(model: dict) -> dict:
    from vcstudio.project.report_contracts import (
        ReportSnapshot,
        ReportSpec,
        ValidationResult,
    )

    spec = ReportSpec.from_mapping(model["report_spec"])
    snapshot = ReportSnapshot.from_mapping(model["report_snapshot"], spec=spec)
    validation = copy.deepcopy(model["validation"])
    validation["report_model_sha256"] = paper_report.report_content_sha256(
        model, outline=spec.outline)
    model["validation"] = ValidationResult.from_mapping(
        validation, spec=spec, snapshot=snapshot
    ).to_dict()
    model.pop("contract_refs", None)
    return model


def test_render_all_formats_from_one_model_and_write_manifest(tmp_path):
    model = _validated_model(tmp_path)
    result = paper_report.render_report_bundle(model, tmp_path / "bundle", stem="screening")

    assert result["schema"] == paper_report.BUNDLE_SCHEMA
    assert set(result["files"]) == {"html", "docx", "pdf", "manifest"}
    assert all(path.is_file() and path.stat().st_size > 100 for path in result["files"].values())
    assert result["model_file"].is_file()
    assert len(result["assets"]) == 1 and result["assets"][0].parent.name == "assets"

    manifest = json.loads(result["manifest"].read_text(encoding="utf-8"))
    assert manifest["schema"] == paper_report.BUNDLE_SCHEMA
    assert manifest["model_schema"] == paper_report.MODEL_SCHEMA
    assert manifest["source_model_schema"] == "vcstudio.research-report/v1"
    assert manifest["artifact_status"] == "complete"
    assert manifest["report_kind"] == "final"
    assert manifest["scientific_status"] == "final"
    assert manifest["scientific_qualification"] == "adsorption_result_verified"
    assert manifest["contract_status"] == "bound"
    assert result["contract_status"] == "bound"
    assert manifest["report_model_sha256"] == result["report_model_sha256"]
    assert manifest["report_model_sha256"] == model["validation"]["report_model_sha256"]
    assert manifest["contracts"]["validation"]["report_model_sha256"] == (
        manifest["report_model_sha256"]
    )
    assert set(result["contract_files"]) == {"spec", "snapshot", "validation"}


def test_manifest_governs_real_archive_inputs_and_html_relative_assets(tmp_path):
    rights = {
        "source_kind": "project",
        "third_party": False,
        "redistributable": True,
        "license": "CC-BY-4.0",
        "attribution": "VCS paper report fixture",
    }
    model = _validated_model(tmp_path)
    model["extensions"] = {"rights": copy.deepcopy(rights)}
    model["figures"][0]["rights"] = copy.deepcopy(rights)
    _rebind_report_content(model)

    result = paper_report.render_report_bundle(
        model, tmp_path / "governed-bundle", stem="governed")
    manifest = json.loads(result["manifest"].read_text(encoding="utf-8"))

    for fmt, media_type in {
        "html": "text/html; charset=utf-8",
        "docx": (
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        "pdf": "application/pdf",
    }.items():
        record = manifest["files"][fmt]
        artifact = result["files"][fmt]
        assert record["logical_role"] == "rendered_report"
        assert record["format"] == fmt
        assert record["media_type"] == media_type
        assert record["authority"] == "report_service_frozen_revision"
        assert record["sha256"] == hashlib.sha256(artifact.read_bytes()).hexdigest()
        assert record["size"] == artifact.stat().st_size
        assert record["rights"] == {
            "schema": paper_report.ARTIFACT_RIGHTS_SCHEMA,
            "license": "CC-BY-4.0",
            "license_status": "declared",
            "attribution": "VCS paper report fixture",
            "source_kind": "project",
            "third_party": False,
            "redistributable": True,
            "redistributable_status": "declared",
        }

    asset = manifest["assets"][0]
    assert asset["logical_role"] == "report_figure"
    assert asset["authority"] == "report_service_frozen_revision"
    assert asset["rights"] == manifest["files"]["html"]["rights"]
    html_refs = manifest["files"]["html"]["relative_refs"]
    assert html_refs == [{
        "path": asset["path"],
        "sha256": asset["sha256"],
        "size": asset["size"],
        "logical_role": "report_figure",
    }]
    referenced_asset = result["files"]["html"].parent / html_refs[0]["path"]
    assert referenced_asset.read_bytes() == result["assets"][0].read_bytes()
    assert result["files"]["docx"].read_bytes().startswith(b"PK")
    assert result["files"]["pdf"].read_bytes().startswith(b"%PDF")
    assert manifest["model_sha256"] == result["model_sha256"]
    model_payload = json.loads(result["model_file"].read_text(encoding="utf-8"))
    assert paper_report._sha256_json(model_payload) == manifest["model_sha256"]
    assert paper_report.frozen_report_content_sha256(model_payload) == (
        manifest["report_model_sha256"]
    )
    assert result["model_file"].read_bytes() == paper_report._canonical_json_bytes(
        model_payload)
    assert manifest["model_file"] == {
        "path": result["model_file"].name,
        "sha256": result["model_sha256"],
        "size": result["model_file"].stat().st_size,
    }
    assert manifest["accessibility"] == result["accessibility"]
    assert manifest["accessibility"]["html"]["status"] == "conditional"
    assert manifest["accessibility"]["docx"]["status"] == "conditional"
    assert manifest["accessibility"]["pdf"]["status"] == "partial"
    assert manifest["accessibility"]["pdf"]["visual"] is True
    assert manifest["accessibility"]["pdf"]["searchable"] is True
    assert manifest["accessibility"]["pdf"]["tagged"] is False
    assert manifest["accessibility"]["pdf"]["pdf_ua"] is False
    assert set(manifest["files"]) == {"html", "docx", "pdf"}
    for record in manifest["files"].values():
        path = result["manifest"].parent / record["path"]
        assert hashlib.sha256(path.read_bytes()).hexdigest() == record["sha256"]
        assert path.stat().st_size == record["size"]
    for key, path in result["contract_files"].items():
        record = manifest["contracts"][key]
        assert path == result["manifest"].parent / record["path"]
        assert hashlib.sha256(path.read_bytes()).hexdigest() == record["file_sha256"]
        assert path.stat().st_size == record["size"]


def test_html_preview_is_self_contained_deterministic_and_side_effect_free(tmp_path):
    model = _validated_model(tmp_path)
    figure = Path(model["figures"][0]["path"]).resolve()
    before = {
        path.relative_to(tmp_path).as_posix() for path in tmp_path.rglob("*")
    }

    first = paper_report.render_report_html_preview(model)
    second = paper_report.render_report_html_preview(model)

    after = {
        path.relative_to(tmp_path).as_posix() for path in tmp_path.rglob("*")
    }
    assert first == second
    assert before == after
    assert first["schema"] == "vcstudio.paper-report.html-preview/v1"
    assert first["artifact_status"] == "preview"
    assert first["report_model_sha256"] == model["validation"]["report_model_sha256"]
    assert first["report_kind"] == first["scientific_status"] == "final"
    assert first["scientific_qualification"] == "adsorption_result_verified"
    assert first["input_fingerprint"] == model["input_fingerprint"]
    assert first["contract_status"] == "bound"
    assert "files" not in first and "manifest" not in first and "revision" not in first

    document = first["html"]
    match = re.search(r"<img src='data:image/png;base64,([^']+)'", document)
    assert match is not None
    assert base64.b64decode(match.group(1)) == figure.read_bytes()
    assert str(figure) not in document
    assert "assets/" not in document
    assert "locator" not in json.dumps(first["contract_refs"], ensure_ascii=False)
    assert not list(tmp_path.rglob("*.manifest.json"))
    assert not list(tmp_path.rglob("*.model.json"))


def test_html_preview_uses_bound_outline_locale_and_matches_formal_hashes(tmp_path):
    model = _validated_model(
        tmp_path,
        formats=("html",),
        locale="zh-CN",
        outline=(
            "limitations", "methods", "adsorption_table", "executive_summary",
        ),
    )

    preview = paper_report.render_report_html_preview(model)
    published = paper_report.render_report_bundle(
        model, tmp_path / "published", formats=("html",)
    )
    formal_html = published["files"]["html"].read_text(encoding="utf-8")

    assert '<html lang="zh-CN">' in preview["html"]
    assert '<meta name="author" content="VASP Catalyst Studio">' in preview["html"]
    assert "<main class=\"paper\">" in preview["html"]
    assert "<section><h2>" in preview["html"]
    assert "<figure>" not in preview["html"]  # outline intentionally omits figures.
    assert preview["accessibility"]["status"] == "conditional"
    assert preview["html"].index("局限性") < preview["html"].index("执行摘要")
    assert "核心结论" not in preview["html"]
    assert formal_html.index("局限性") < formal_html.index("执行摘要")
    assert "核心结论" not in formal_html
    assert preview["report_model_sha256"] == published["report_model_sha256"]
    assert preview["model_sha256"] == published["model_sha256"]
    published_refs = {
        key: {
            field: value
            for field, value in record.items()
            if field not in {"path", "file_sha256", "size"}
        }
        for key, record in published["contracts"].items()
    }
    assert preview["contract_refs"] == published_refs


@pytest.mark.parametrize(
    ("theme_id", "accent"),
    [
        ("academic-a4", "#1F4F64"),
        ("compact-brief", "#0F766E"),
        ("diagnostic-a4", "#9A3412"),
    ],
)
def test_html_preview_applies_and_binds_server_owned_theme(
        tmp_path, theme_id, accent):
    model = _validated_model(
        tmp_path, formats=("html",), theme_id=theme_id)

    preview = paper_report.render_report_html_preview(model)

    assert f'data-report-theme="{theme_id}"' in preview["html"]
    assert f"--accent:{accent}" in preview["html"]
    assert preview["report_model_sha256"] == (
        model["validation"]["report_model_sha256"])


def test_report_content_digest_changes_with_visible_theme(tmp_path):
    model = _model(tmp_path)
    model["template_ref"] = {
        "id": "vcstudio-academic-a4", "version": "1"}
    academic = paper_report.report_content_sha256(model)
    model["template_ref"] = {
        "id": "vcstudio-diagnostic-a4", "version": "1"}

    assert paper_report.report_content_sha256(model) != academic


def test_html_preview_rejects_contract_mismatch_and_validated_content_tamper(tmp_path):
    mismatched = _validated_model(tmp_path, formats=("html",))
    mismatched["report_snapshot"]["spec_sha256"] = "0" * 64
    with pytest.raises(ValueError, match="snapshot spec binding mismatch"):
        paper_report.render_report_html_preview(mismatched)

    tampered = _validated_model(tmp_path, formats=("html",))
    tampered["executive_summary"] = "Validated content was replaced after the gate."
    with pytest.raises(ValueError, match="report content conflicts"):
        paper_report.render_report_html_preview(tampered)


def test_html_preview_rejects_figure_bytes_changed_after_validation(tmp_path):
    model = _validated_model(tmp_path, formats=("html",))
    _png(Path(model["figures"][0]["path"]), color=(220, 38, 38))

    with pytest.raises(ValueError, match="report content conflicts"):
        paper_report.render_report_html_preview(model)


def test_model_hash_and_visible_label_change_with_report_kind(tmp_path):
    final_model = _validated_model(tmp_path, formats=("html",))
    final = paper_report.render_report_bundle(
        final_model, tmp_path / "final", stem="report", formats=("html",)
    )
    diagnostic_model = _validated_model(
        tmp_path,
        formats=("html",),
        requested_kind="diagnostic",
        effective_kind="diagnostic",
    )
    diagnostic = paper_report.render_report_bundle(
        diagnostic_model, tmp_path / "diagnostic", stem="report", formats=("html",)
    )

    assert final["model_sha256"] != diagnostic["model_sha256"]
    assert final["scientific_status"] == "final"
    assert diagnostic["scientific_status"] == "diagnostic"
    html_text = diagnostic["files"]["html"].read_text(encoding="utf-8")
    assert "诊断报告" in html_text or "DIAGNOSTIC REPORT" in html_text
    manifest = json.loads(diagnostic["manifest"].read_text(encoding="utf-8"))
    assert manifest["artifact_status"] == "complete"
    assert manifest["report_kind"] == "diagnostic"


def test_report_content_digest_binds_visible_scientific_state(tmp_path):
    diagnostic = _model(tmp_path)
    baseline = paper_report.report_content_sha256(diagnostic)

    draft = copy.deepcopy(diagnostic)
    draft["report_kind"] = "draft"
    final = copy.deepcopy(diagnostic)
    final["report_kind"] = "final"
    qualified = copy.deepcopy(diagnostic)
    qualified["scientific_qualification"] = "adsorption_result_verified"
    ceiling = copy.deepcopy(diagnostic)
    ceiling["claim_ceiling"] = "electronic_adsorption_screen"

    assert paper_report.report_content_sha256(draft) != baseline
    assert paper_report.report_content_sha256(final) != baseline
    assert paper_report.report_content_sha256(qualified) != baseline
    assert paper_report.report_content_sha256(ceiling) != baseline


def test_legacy_flat_model_defaults_to_diagnostic_fail_closed(tmp_path):
    model = _model(tmp_path)
    for key in (
        "schema",
        "report_id",
        "report_kind",
        "scientific_qualification",
        "input_fingerprint",
        "preset_id",
        "contract_refs",
    ):
        model.pop(key, None)

    result = paper_report.render_report_bundle(
        model, tmp_path / "legacy", stem="legacy", formats=("html",)
    )
    manifest = json.loads(result["manifest"].read_text(encoding="utf-8"))
    assert result["report_kind"] == "diagnostic"
    assert manifest["scientific_qualification"] == "diagnostic"
    assert manifest["source_model_schema"] == "vcstudio.paper-report.model/v1"


def test_diagnostic_label_is_visible_in_html_docx_and_pdf(tmp_path):
    from docx import Document
    from pypdf import PdfReader

    model = _model(tmp_path)
    model.update(
        locale="zh-CN",
        report_kind="diagnostic",
        scientific_qualification="diagnostic",
    )
    result = paper_report.render_report_bundle(
        model, tmp_path / "diagnostic-all", stem="diagnostic"
    )

    html_text = result["files"]["html"].read_text(encoding="utf-8")
    docx_text = "\n".join(
        paragraph.text for paragraph in Document(result["files"]["docx"]).paragraphs
    )
    pdf_text = "\n".join(
        page.extract_text() or "" for page in PdfReader(result["files"]["pdf"]).pages
    )
    for text in (html_text, docx_text, pdf_text):
        assert "诊断报告" in text
        assert "不构成最终科学结论" in text


def test_assets_are_hash_staged_and_never_use_cross_drive_relpath(monkeypatch, tmp_path):
    model = _model(tmp_path)
    duplicate = tmp_path / "different-name.png"
    duplicate.write_bytes(Path(model["figures"][0]["path"]).read_bytes())
    model["figures"].append({"path": duplicate, "title": "Same pixels"})

    def forbidden_relpath(*_args, **_kwargs):
        raise AssertionError("os.path.relpath must not be used for report assets")

    monkeypatch.setattr(paper_report.os.path, "relpath", forbidden_relpath)
    result = paper_report.render_report_bundle(
        model,
        tmp_path / "other-volume-output",
        formats=("html",),
    )
    assert len(result["assets"]) == 1  # identical bytes are deduplicated
    html = result["files"]["html"].read_text(encoding="utf-8")
    assert html.count("assets/") == 2
    assert str(tmp_path) not in html
    assert model["figures"][0]["path"].is_file()  # sources are never moved or rewritten


def test_html_escapes_content_and_uses_paper_tables(tmp_path):
    model = _model(tmp_path)
    model["title"] = "<script>alert(1)</script>"
    result = paper_report.render_report_bundle(
        model,
        tmp_path / "html-only",
        formats=("html",),
    )
    text = result["files"]["html"].read_text(encoding="utf-8")
    assert "<script>alert(1)</script>" not in text
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in text
    assert "@page" in text and "size: A4 portrait" in text
    assert "table class='three-line'" in text
    assert "Figure 1. Free-energy pathways" in text


def test_docx_is_a4_with_header_footer_page_field_and_three_line_table(tmp_path):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Mm
    from xml.etree import ElementTree as ET

    result = paper_report.render_report_bundle(
        _model(tmp_path),
        tmp_path / "docx-only",
        formats=("docx",),
    )
    path = result["files"]["docx"]
    doc = Document(path)
    section = doc.sections[0]
    assert abs(section.page_width - Mm(210)) < 1000  # OOXML rounds to whole twips.
    assert abs(section.page_height - Mm(297)) < 1000
    assert "COMPUTATIONAL CATALYSIS" in section.header.paragraphs[0].text
    with zipfile.ZipFile(path) as archive:
        footer_xml = archive.read("word/footer1.xml").decode("utf-8")
        document_xml = archive.read("word/document.xml").decode("utf-8")
        styles_xml = archive.read("word/styles.xml").decode("utf-8")
        core_xml = archive.read("docProps/core.xml").decode("utf-8")
    assert "PAGE" in footer_xml
    assert "w:numPr" in document_xml  # findings/recommendations use real Word numbering.
    namespaces = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
        "dc": "http://purl.org/dc/elements/1.1/",
    }
    document_root = ET.fromstring(document_xml)
    image_properties = document_root.findall(".//wp:docPr", namespaces)
    assert any(
        item.get("descr") == "Overlaid free-energy pathways"
        and item.get("title") == "Free-energy pathways"
        for item in image_properties
    )
    assert document_root.find(".//w:trPr/w:tblHeader", namespaces) is not None
    styles_root = ET.fromstring(styles_xml)
    language_nodes = styles_root.findall(".//w:lang", namespaces)
    assert language_nodes
    assert all(
        node.get(qn("w:val")) == "en-US"
        and node.get(qn("w:eastAsia")) == "en-US"
        for node in language_nodes
    )
    core_root = ET.fromstring(core_xml)
    assert core_root.find("dc:language", namespaces).text == "en-US"
    assert doc.tables
    borders = doc.tables[0]._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders.find(qn("w:top")).get(qn("w:val")) == "single"
    assert borders.find(qn("w:bottom")).get(qn("w:val")) == "single"
    assert borders.find(qn("w:insideV")).get(qn("w:val")) == "nil"
    all_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "Adsorption Screening Report" in all_text
    assert "Table 1. Candidate evaluation" in all_text
    assert "Figure 1. Free-energy pathways" in all_text
    assert any(paragraph.style.name == "Heading 1" for paragraph in doc.paragraphs)
    assert any(paragraph.style.name == "Caption" for paragraph in doc.paragraphs)
    assert result["accessibility"]["docx"]["status"] == "conditional"
    assert result["accessibility"]["docx"]["source_figure_alt_complete"] is True


def test_pdf_is_a4_multipage_and_contains_shared_model_text(tmp_path):
    from pypdf import PdfReader

    result = paper_report.render_report_bundle(
        _model(tmp_path),
        tmp_path / "pdf-only",
        formats=("pdf",),
    )
    reader = PdfReader(result["files"]["pdf"])
    catalog = reader.trailer["/Root"]
    assert len(reader.pages) >= 2
    page = reader.pages[0]
    width = float(page.mediabox.width)
    height = float(page.mediabox.height)
    assert width == pytest.approx(595.28, abs=0.5)
    assert height == pytest.approx(841.89, abs=0.5)
    text = "\n".join(item.extract_text() or "" for item in reader.pages)
    assert "Adsorption Screening Report" in text
    assert "Candidate A" in text
    assert "Figure 1." in text
    assert "Page 2" in text
    assert catalog["/Lang"] == "en-US"
    assert "/StructTreeRoot" not in catalog
    assert reader.metadata.title == "Adsorption Screening Report"
    assert reader.metadata.author == "VASP Catalyst Studio"
    assert reader.metadata.subject == "A reproducible multi-catalyst comparison"
    assert "scientific report" in str(reader.metadata.get("/Keywords"))
    assert result["accessibility"]["pdf"]["status"] == "partial"
    assert result["accessibility"]["pdf"]["tagged"] is False
    assert result["accessibility"]["pdf"]["pdf_ua"] is False


@pytest.mark.parametrize(
    ("alt_value", "title"),
    [
        (None, "Free-energy pathways"),
        ("Figure 1", "Free-energy pathways"),
        ("图 1", "自由能路径"),
    ],
)
def test_missing_or_generic_figure_alt_is_fail_closed(tmp_path, alt_value, title):
    model = _model(tmp_path)
    model["figures"][0]["title"] = title
    if alt_value is None:
        model["figures"][0].pop("alt", None)
    else:
        model["figures"][0]["alt"] = alt_value

    result = paper_report.render_report_bundle(
        model, tmp_path / f"alt-{alt_value or 'missing'}", formats=("html",)
    )
    record = result["accessibility"]["html"]

    assert record["status"] == "partial"
    assert record["figures_total"] == 1
    assert record["figures_missing_meaningful_alt"] == 1
    assert record["source_figure_alt_complete"] is False
    assert "1" in record["reason_zh"]
    html_document = result["files"]["html"].read_text(encoding="utf-8")
    if alt_value is None:
        assert "alt=''" in html_document
        assert "alt='Free-energy pathways'" not in html_document
    manifest = json.loads(result["manifest"].read_text(encoding="utf-8"))
    assert manifest["accessibility"]["html"] == record


def test_model_hash_tracks_figure_content_not_source_location(tmp_path):
    first_model = _model(tmp_path)
    first = paper_report.render_report_bundle(
        first_model,
        tmp_path / "first",
        formats=("html",),
    )
    moved = tmp_path / "moved" / "renamed.png"
    moved.parent.mkdir()
    moved.write_bytes(Path(first_model["figures"][0]["path"]).read_bytes())
    second_model = _model(tmp_path)
    second_model["figures"][0]["path"] = moved
    second = paper_report.render_report_bundle(
        second_model,
        tmp_path / "second",
        formats=("html",),
    )
    assert first["model_sha256"] == second["model_sha256"]

    _png(moved, color=(220, 38, 38))
    changed = paper_report.render_report_bundle(
        second_model,
        tmp_path / "changed",
        formats=("html",),
    )
    assert changed["model_sha256"] != second["model_sha256"]


def test_model_hash_uses_contract_semantics_not_time_or_locator(tmp_path):
    from vcstudio.project.report_contracts import (
        ReportSnapshot,
        ReportSpec,
        ValidationCheck,
        ValidationResult,
    )

    def contracts(day, template_locator, project_locator, outline):
        spec = ReportSpec(
            requested_kind="final",
            locale="en-US",
            formats=("html",),
            scope={"kind": "project", "project_ids": ["project-1"], "job_ids": []},
            outline=tuple(outline),
            template_ref={"id": "paper", "locator": template_locator},
            created_at_utc=f"2026-08-{day:02d}T01:00:00+00:00",
        )
        snapshot = ReportSnapshot(
            spec_sha256=spec.semantic_sha256,
            input_fingerprint="input-fingerprint-001",
            created_at_utc=f"2026-08-{day:02d}T02:00:00+00:00",
            sources=({"source_id": "project:1", "locator": project_locator},),
        )
        content_model = _model(tmp_path)
        content_model.update(
            report_kind="final",
            scientific_qualification="adsorption_result_verified",
            template_ref=spec.template_ref,
        )
        validation = ValidationResult(
            spec_sha256=spec.semantic_sha256,
            snapshot_sha256=snapshot.semantic_sha256,
            validated_at_utc=f"2026-08-{day:02d}T03:00:00+00:00",
            validator={"id": "gate", "version": "1"},
            status="passed",
            effective_kind="final",
            final_allowed=True,
            scientific_qualification="adsorption_result_verified",
            report_model_sha256=paper_report.report_content_sha256(
                content_model, outline=outline),
            checks=(ValidationCheck(id="gate", status="pass"),),
        )
        return spec.to_dict(), snapshot.to_dict(), validation.to_dict()

    first_model = _model(tmp_path)
    first_model.pop("contract_refs", None)
    first_model.update(
        report_kind="final",
        scientific_qualification="adsorption_result_verified",
    )
    spec, snapshot, validation = contracts(
        9, "C:/templates/a", "C:/project/a",
        ["executive_summary", "adsorption_table", "methods", "limitations"])
    first_model.update(
        report_spec=spec,
        report_snapshot=snapshot,
        validation=validation,
    )
    first = paper_report.render_report_bundle(
        first_model, tmp_path / "contract-first", formats=("html",)
    )
    moved = copy.deepcopy(first_model)
    spec, snapshot, validation = contracts(
        10, "D:/templates/moved", "Z:/project/moved",
        ["executive_summary", "adsorption_table", "methods", "limitations"])
    moved.update(
        report_spec=spec,
        report_snapshot=snapshot,
        validation=validation,
    )
    second = paper_report.render_report_bundle(
        moved, tmp_path / "contract-moved", formats=("html",)
    )
    changed_contract = copy.deepcopy(moved)
    spec, snapshot, validation = contracts(
        10, "D:/templates/moved", "Z:/project/moved",
        ["limitations", "methods", "adsorption_table", "executive_summary"])
    changed_contract.update(
        report_spec=spec,
        report_snapshot=snapshot,
        validation=validation,
    )
    third = paper_report.render_report_bundle(
        changed_contract, tmp_path / "contract-changed", formats=("html",)
    )

    assert first["model_sha256"] == second["model_sha256"]
    assert {
        key: record["sha256"] for key, record in first["contracts"].items()
    } == {
        key: record["sha256"] for key, record in second["contracts"].items()
    }
    assert first["contracts"]["spec"]["file_sha256"] != (
        second["contracts"]["spec"]["file_sha256"]
    )
    assert third["model_sha256"] != second["model_sha256"]


def test_model_hash_keeps_arbitrary_claim_extension_locator_and_time_semantics(tmp_path):
    from vcstudio.project.report_contracts import (
        ClaimRecord,
        ReportSnapshot,
        ReportSpec,
        ValidationResult,
    )

    def model_with_claim(locator, created_at):
        model = _validated_model(tmp_path, formats=("html",))
        spec = ReportSpec.from_mapping(model["report_spec"])
        snapshot = ReportSnapshot.from_mapping(model["report_snapshot"], spec=spec)
        claim = ClaimRecord(
            id="claim.adsorption",
            text="The adsorption result passed the declared gate.",
            qualification="adsorption_result_verified",
            status="supported",
            extensions={
                "source_locator": locator,
                "created_at_utc": created_at,
                "semantic_note": "same claim",
            },
        )
        validation_data = copy.deepcopy(model["validation"])
        validation_data["claims"] = [claim.to_dict()]
        validation = ValidationResult.from_mapping(
            validation_data, spec=spec, snapshot=snapshot
        )
        model["validation"] = validation.to_dict()
        model["claims"] = [claim.to_dict()]
        return model, validation.semantic_sha256

    first_model, first_validation_hash = model_with_claim(
        "C:/evidence/claim.json", "2026-08-09T01:00:00+00:00"
    )
    moved_model, moved_validation_hash = model_with_claim(
        "Z:/archive/claim.json", "2026-08-10T02:00:00+00:00"
    )
    assert first_validation_hash != moved_validation_hash

    first = paper_report.render_report_bundle(
        first_model, tmp_path / "claim-first", formats=("html",)
    )
    moved = paper_report.render_report_bundle(
        moved_model, tmp_path / "claim-moved", formats=("html",)
    )
    assert first["model_sha256"] != moved["model_sha256"]


def test_full_contract_must_match_explicit_semantic_reference(tmp_path):
    from vcstudio.project.report_contracts import ReportSpec

    model = _validated_model(tmp_path, formats=("html",))
    spec = ReportSpec.from_mapping(model["report_spec"])
    model["contract_refs"] = {"spec": {
        "schema": spec.schema,
        "sha256": spec.semantic_sha256,
    }}
    result = paper_report.render_report_bundle(
        model, tmp_path / "contract-match", formats=("html",)
    )
    assert result["contracts"]["spec"]["sha256"] == spec.semantic_sha256

    mismatched = copy.deepcopy(model)
    mismatched["contract_refs"]["spec"]["sha256"] = "0" * 64
    with pytest.raises(ValueError, match="conflicts with the full contract"):
        paper_report.render_report_bundle(
            mismatched, tmp_path / "contract-mismatch", formats=("html",)
        )


def test_blocked_validation_cannot_be_labeled_final(tmp_path):
    from vcstudio.project.report_contracts import (
        ReportSnapshot,
        ReportSpec,
        ValidationCheck,
        ValidationResult,
    )

    spec = ReportSpec(
        requested_kind="final", locale="en-US", formats=("html",),
        scope={"kind": "project", "project_ids": ["p1"], "job_ids": []})
    snapshot = ReportSnapshot(
        spec_sha256=spec.semantic_sha256,
        input_fingerprint="blocked-input")
    validation = ValidationResult(
        spec_sha256=spec.semantic_sha256,
        snapshot_sha256=snapshot.semantic_sha256,
        validator={"id": "gate", "version": "1"},
        status="blocked",
        effective_kind="diagnostic",
        final_allowed=False,
        scientific_qualification="diagnostic",
        checks=(ValidationCheck(id="gate", status="fail"),),
    )
    model = _model(tmp_path)
    model.pop("contract_refs", None)
    model.update(
        report_kind="final",
        scientific_qualification="adsorption_result_verified",
        input_fingerprint="blocked-input",
        report_spec=spec.to_dict(),
        report_snapshot=snapshot.to_dict(),
        validation=validation.to_dict(),
    )

    with pytest.raises(ValueError, match="conflicts with validation.effective_kind"):
        paper_report.render_report_bundle(
            model, tmp_path / "blocked-as-final", formats=("html",)
        )


def test_final_without_validation_is_downgraded_fail_closed(tmp_path):
    model = _model(tmp_path)
    model.update(
        report_kind="final",
        scientific_qualification="adsorption_result_verified",
    )

    result = paper_report.render_report_bundle(
        model, tmp_path / "unvalidated-final", formats=("html",)
    )

    assert result["report_kind"] == "diagnostic"
    assert result["scientific_qualification"] == "diagnostic"
    manifest = json.loads(result["manifest"].read_text(encoding="utf-8"))
    assert manifest["report_kind"] == "diagnostic"
    assert "DIAGNOSTIC REPORT" in result["files"]["html"].read_text(encoding="utf-8")


def test_reference_only_validation_cannot_authorize_final(tmp_path):
    model = _model(tmp_path)
    model.update(
        report_kind="final",
        scientific_qualification="adsorption_result_verified",
        contract_refs={
            "spec": {"schema": "vcstudio.report-spec/v1", "sha256": "a" * 64},
            "snapshot": {
                "schema": "vcstudio.report-snapshot/v1",
                "sha256": "b" * 64,
                "input_fingerprint": model["input_fingerprint"],
            },
            "validation": {
                "schema": "vcstudio.report-validation/v1",
                "sha256": "c" * 64,
                "status": "passed",
                "final_allowed": True,
            },
        },
    )

    result = paper_report.render_report_bundle(
        model, tmp_path / "reference-only", formats=("html",)
    )

    assert result["report_kind"] == "diagnostic"
    assert result["scientific_qualification"] == "diagnostic"
    assert result["contract_status"] == "references_only"
    assert result["contract_files"] == {}


def test_unbound_diagnostic_cannot_self_declare_verified_qualification(tmp_path):
    model = _model(tmp_path)
    model["scientific_qualification"] = "human_scientific_reviewed"

    result = paper_report.render_report_bundle(
        model, tmp_path / "unbound-high-qualification", formats=("html",)
    )

    assert result["report_kind"] == "diagnostic"
    assert result["scientific_qualification"] == "diagnostic"


def test_diagnostic_report_preserves_verified_lower_qualification(tmp_path):
    from vcstudio.project.report_contracts import (
        ReportSnapshot,
        ReportSpec,
        ValidationCheck,
        ValidationResult,
    )

    spec = ReportSpec(
        requested_kind="final", locale="en-US", formats=("html",),
        scope={"kind": "project", "project_ids": ["p1"], "job_ids": []})
    snapshot = ReportSnapshot(
        spec_sha256=spec.semantic_sha256,
        input_fingerprint="lower-qualified-input")
    model = _model(tmp_path)
    model["scientific_qualification"] = "adsorption_result_verified"
    report_model_sha256 = paper_report.report_content_sha256(
        model, outline=spec.outline)
    validation = ValidationResult(
        spec_sha256=spec.semantic_sha256,
        snapshot_sha256=snapshot.semantic_sha256,
        validator={"id": "publication-gate", "version": "1"},
        status="blocked",
        effective_kind="diagnostic",
        final_allowed=False,
        scientific_qualification="adsorption_result_verified",
        report_model_sha256=report_model_sha256,
        checks=(
            ValidationCheck(id="adsorption-delivery-gate", status="pass"),
            ValidationCheck(id="publication-gate", status="fail"),
        ),
    )
    model.pop("contract_refs", None)
    model.update(
        report_kind="diagnostic",
        scientific_qualification="adsorption_result_verified",
        input_fingerprint="lower-qualified-input",
        report_spec=spec.to_dict(),
        report_snapshot=snapshot.to_dict(),
        validation=validation.to_dict(),
    )

    result = paper_report.render_report_bundle(
        model, tmp_path / "lower-qualified", formats=("html",)
    )
    assert result["report_kind"] == "diagnostic"
    assert result["scientific_qualification"] == "adsorption_result_verified"


def test_partial_full_contract_chain_is_rejected(tmp_path):
    model = _model(tmp_path)
    complete = _validated_model(tmp_path, formats=("html",))
    model["report_spec"] = complete["report_spec"]

    with pytest.raises(ValueError, match="must include ReportSpec, ReportSnapshot"):
        paper_report.render_report_bundle(
            model, tmp_path / "partial-contract", formats=("html",)
        )


@pytest.mark.parametrize(
    ("field", "value"),
    (("report_spec", []), ("report_snapshot", "not-a-mapping"), ("validation", 7)),
)
def test_full_contract_fields_reject_non_mapping_values(tmp_path, field, value):
    model = _model(tmp_path)
    model[field] = value

    with pytest.raises(TypeError, match=field):
        paper_report.render_report_bundle(
            model, tmp_path / f"wrong-{field}", formats=("html",)
        )


@pytest.mark.parametrize(
    ("mutate", "message"),
    (
        (lambda model: model.update(locale="zh-CN"), "locale conflicts"),
        (
            lambda model: model.update(input_fingerprint="different-input"),
            "input_fingerprint conflicts",
        ),
    ),
)
def test_render_context_must_match_bound_contracts(tmp_path, mutate, message):
    model = _validated_model(tmp_path, formats=("html",))
    mutate(model)

    with pytest.raises(ValueError, match=message):
        paper_report.render_report_bundle(
            model, tmp_path / "context-conflict", formats=("html",)
        )


def test_requested_formats_must_match_report_spec(tmp_path):
    model = _validated_model(tmp_path, formats=("html",))

    with pytest.raises(ValueError, match="formats conflict"):
        paper_report.render_report_bundle(
            model, tmp_path / "format-conflict", formats=("docx",)
        )


def test_bound_outline_controls_section_order_and_exclusion(tmp_path):
    model = _validated_model(
        tmp_path,
        formats=("html",),
        outline=(
            "limitations", "methods", "adsorption_table", "executive_summary",
        ),
    )

    result = paper_report.render_report_bundle(
        model, tmp_path / "outline", formats=("html",)
    )
    html = result["files"]["html"].read_text(encoding="utf-8")

    assert html.index("Limitations") < html.index("Executive Summary")
    assert "Key Findings" not in html
    assert "Adsorption-Energy Results" in html
    assert "Cross-Project Comparison" not in html


def test_bound_final_methods_only_outline_is_rejected_by_renderer(tmp_path):
    model = _validated_model(
        tmp_path,
        formats=("html",),
        outline=("methods",),
    )

    with pytest.raises(
        ValueError,
        match="final report outline.*scientific result/evidence section",
    ):
        paper_report.render_report_bundle(
            model,
            tmp_path / "methods-only-final",
            formats=("html",),
        )


@pytest.mark.parametrize(
    ("preset_id", "outline", "empty_field", "message"),
    [
        (
            "scientific-review",
            ("executive_summary", "adsorption_table", "methods", "limitations"),
            "adsorption_table",
            "result/evidence content",
        ),
        (
            "scientific-review",
            ("executive_summary", "adsorption_table", "methods", "limitations"),
            "limitations",
            "limitations content",
        ),
        (
            "quick-decision-brief",
            ("key_findings", "figures", "limitations"),
            "figures",
            "rendered figure/table",
        ),
    ],
)
def test_final_renderer_requires_real_visible_content_not_only_outline_names(
    tmp_path, preset_id, outline, empty_field, message,
):
    model = _validated_model(
        tmp_path,
        formats=("html",),
        outline=outline,
        preset_id=preset_id,
    )
    if empty_field == "figures":
        model[empty_field] = []
    else:
        model[empty_field] = None
    _rebind_report_content(model)

    with pytest.raises(ValueError, match=message):
        paper_report.render_report_bundle(
            model, tmp_path / f"empty-{empty_field}", formats=("html",)
        )


def test_bound_preview_redacts_local_paths_from_visible_metadata(tmp_path):
    model = _validated_model(tmp_path, formats=("html",))
    secret = r"C:\Users\alice\secret\project.yaml"
    model["metadata"]["Source"] = secret
    _rebind_report_content(model)

    preview = paper_report.render_report_html_preview(model)

    assert secret not in preview["html"]
    assert "local path redacted" in preview["html"]


def test_bound_outline_rejects_unknown_renderer_section(tmp_path):
    model = _validated_model(
        tmp_path,
        formats=("html",),
        outline=("limitations", "custom_section"),
    )

    with pytest.raises(ValueError, match="unsupported sections: custom_section"):
        paper_report.render_report_bundle(
            model, tmp_path / "unknown-outline", formats=("html",)
        )


@pytest.mark.parametrize(
    "mutate",
    (
        lambda model: model.update(
            executive_summary="Validated energy is +999 eV."),
        lambda model: model["adsorption_table"]["rows"][0].update(energy=999.0),
    ),
)
def test_bound_report_rejects_content_changed_after_validation(tmp_path, mutate):
    model = _validated_model(tmp_path, formats=("html",))
    mutate(model)

    with pytest.raises(ValueError, match="report content conflicts"):
        paper_report.render_report_bundle(
            model, tmp_path / "content-mismatch", formats=("html",)
        )


@pytest.mark.parametrize(
    ("field", "value", "message"),
    (
        ("preset_id", "different-preset", "preset_id conflicts"),
        ("claim_ceiling", "different-ceiling", "claim_ceiling conflicts"),
        (
            "claims",
            [{"id": "unbound-claim", "text": "not present in validation"}],
            "claims conflict",
        ),
        ("template_ref", {"id": "other-template"}, "template_ref conflicts"),
        ("policy_refs", [{"id": "other-policy"}], "policy_refs conflict"),
    ),
)
def test_top_level_science_fields_must_match_bound_contracts(
        tmp_path, field, value, message):
    model = _validated_model(tmp_path, formats=("html",))
    model[field] = value

    with pytest.raises(ValueError, match=message):
        paper_report.render_report_bundle(
            model, tmp_path / f"context-{field}", formats=("html",)
        )


def test_full_contract_sidecars_are_canonical_and_rehashable(tmp_path):
    from vcstudio.project.report_contracts import (
        ReportSnapshot,
        ReportSpec,
        ValidationResult,
    )

    model = _validated_model(tmp_path, formats=("html",))
    model["report_spec"]["formats"] = ["HTML", "html"]
    model["report_spec"]["created_at_utc"] = "2026-08-09T01:00:00Z"
    spec = ReportSpec.from_mapping(model["report_spec"])
    assert spec.semantic_sha256 == model["report_snapshot"]["spec_sha256"]

    model["report_snapshot"]["sources"] = [
        {"source_id": "project:z", "locator": "Z:/moved/project.yaml"},
        {"source_id": "project:a", "locator": "C:/work/project.yaml"},
    ]
    snapshot = ReportSnapshot.from_mapping(model["report_snapshot"], spec=spec)
    model["validation"]["snapshot_sha256"] = snapshot.semantic_sha256

    result = paper_report.render_report_bundle(
        model, tmp_path / "canonical-contracts", formats=("html",)
    )
    manifest = json.loads(result["manifest"].read_text(encoding="utf-8"))
    classes = {
        "spec": ReportSpec,
        "snapshot": ReportSnapshot,
        "validation": ValidationResult,
    }
    loaded = {}
    for key, cls in classes.items():
        data = json.loads(result["contract_files"][key].read_text(encoding="utf-8"))
        if key == "spec":
            loaded[key] = cls.from_mapping(data)
        elif key == "snapshot":
            loaded[key] = cls.from_mapping(data, spec=loaded["spec"])
        else:
            loaded[key] = cls.from_mapping(
                data, spec=loaded["spec"], snapshot=loaded["snapshot"]
            )
        assert loaded[key].semantic_sha256 == manifest["contracts"][key]["sha256"]

    spec_data = json.loads(result["contract_files"]["spec"].read_text(encoding="utf-8"))
    snapshot_data = json.loads(
        result["contract_files"]["snapshot"].read_text(encoding="utf-8")
    )
    assert spec_data["formats"] == ["html"]
    assert spec_data["created_at_utc"].endswith("+00:00")
    assert [item["source_id"] for item in snapshot_data["sources"]] == [
        "project:a", "project:z"
    ]


def test_render_accepts_read_only_mapping(tmp_path):
    model = MappingProxyType(_model(tmp_path))

    result = paper_report.render_report_bundle(
        model, tmp_path / "mapping", formats=("html",)
    )

    assert result["files"]["html"].is_file()


@pytest.mark.parametrize(
    "mutate",
    (
        lambda model, value: model.update(title=value),
        lambda model, value: model.update(extensions={"unstable": value}),
        lambda model, value: model.update(executive_summary=value),
    ),
)
def test_non_json_model_values_are_rejected_instead_of_hashing_repr(
        tmp_path, mutate):
    class UnstableRepresentation:
        pass

    model = _model(tmp_path)
    mutate(model, UnstableRepresentation())

    with pytest.raises(TypeError, match="unsupported|non-JSON"):
        paper_report.render_report_bundle(
            model, tmp_path / "non-json", formats=("html",)
        )


def test_model_hash_excludes_only_known_reference_administration(tmp_path):
    first_model = _model(tmp_path)
    first_model.update(
        revision={
            "id": "revision-1",
            "sequence": 1,
            "created_at_utc": "2026-08-09T01:00:00+00:00",
        },
        template_ref={"id": "paper", "version": "1", "locator": "C:/templates/a"},
        policy_refs=[{"id": "gate", "version": "1", "locator": "C:/policies/a"}],
        extensions={"semantic": {"mode": "screening"}, "source_locator": "C:/work/a"},
    )
    first_model["figures"][0]["extensions"] = {
        "dataset": "adsorption",
        "source_locator": "C:/figures/a.png",
        "created_at_utc": "2026-08-09T01:00:00+00:00",
    }
    first = paper_report.render_report_bundle(
        first_model, tmp_path / "admin-first", formats=("html",)
    )

    moved_model = copy.deepcopy(first_model)
    moved_model["revision"]["created_at_utc"] = "2026-08-10T02:00:00+00:00"
    moved_model["template_ref"]["locator"] = "D:/templates/moved"
    moved_model["policy_refs"][0]["locator"] = "Z:/policies/moved"
    moved = paper_report.render_report_bundle(
        moved_model, tmp_path / "admin-moved", formats=("html",)
    )

    assert first["model_sha256"] == moved["model_sha256"]

    changed_extensions = copy.deepcopy(moved_model)
    changed_extensions["extensions"]["source_locator"] = "D:/archive/moved"
    changed_extensions["figures"][0]["extensions"]["source_locator"] = (
        "Z:/figures/moved.png"
    )
    changed_extensions["figures"][0]["extensions"]["created_at_utc"] = (
        "2026-08-10T02:00:00+00:00"
    )
    changed = paper_report.render_report_bundle(
        changed_extensions, tmp_path / "extension-semantic-change", formats=("html",)
    )
    assert changed["model_sha256"] != moved["model_sha256"]

    changed_revision = copy.deepcopy(moved_model)
    changed_revision["revision"]["sequence"] = 2
    revision_result = paper_report.render_report_bundle(
        changed_revision, tmp_path / "revision-semantic-change", formats=("html",)
    )
    assert revision_result["model_sha256"] != moved["model_sha256"]


def test_missing_requested_dependency_has_clear_error(monkeypatch, tmp_path):
    real_import = paper_report.importlib.import_module

    def missing_docx(name, *args, **kwargs):
        if name == "docx":
            raise ImportError("simulated")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(paper_report.importlib, "import_module", missing_docx)
    with pytest.raises(paper_report.ReportDependencyError, match="python-docx"):
        paper_report.render_report_bundle(
            _model(tmp_path),
            tmp_path / "missing",
            formats=("docx",),
        )


def test_report_capabilities_exposes_each_optional_dependency(monkeypatch):
    real_import = paper_report.importlib.import_module

    def missing_docx(name, *args, **kwargs):
        if name == "docx":
            raise ImportError("simulated")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(paper_report.importlib, "import_module", missing_docx)
    result = paper_report.report_capabilities()

    assert result["schema"] == "vcstudio.paper-report.capabilities/v1"
    assert result["formats"]["html"] == {"available": True, "reason": ""}
    assert result["formats"]["docx"]["available"] is False
    assert "python-docx" in result["formats"]["docx"]["reason"]
    assert result["formats"]["pdf"]["available"] is True
    assert result["accessibility"]["html"]["status"] == "conditional"
    assert result["accessibility"]["docx"]["status"] == "conditional"
    assert result["accessibility"]["pdf"]["status"] == "partial"
    assert result["accessibility"]["pdf"]["visual"] is True
    assert result["accessibility"]["pdf"]["searchable"] is True
    assert result["accessibility"]["pdf"]["semantic_structure"] is False
    assert result["accessibility"]["pdf"]["tagged"] is False
    assert result["accessibility"]["pdf"]["pdf_ua"] is False
    assert "untagged" in result["accessibility"]["pdf"]["reason_en"]


def test_report_capabilities_rejects_corrupt_pdf_fonts(monkeypatch, tmp_path):
    font_root = tmp_path / "corrupt-fonts"
    font_root.mkdir()
    names = (
        "dejavu-sans-400.ttf",
        "dejavu-sans-700.ttf",
        "dejavu-sans-400-italic.ttf",
        "dejavu-sans-700-italic.ttf",
        "noto-sans-sc-400.ttf",
        "noto-sans-sc-700.ttf",
    )
    for name in names:
        (font_root / name).write_bytes(b"not-a-font")

    monkeypatch.setattr(paper_report, "_pdf_font_root", lambda: font_root)
    monkeypatch.setattr(
        paper_report,
        "_latin_pdf_font_files",
        lambda _package: {
            "PaperSans": font_root / "dejavu-sans-400.ttf",
            "PaperSansBold": font_root / "dejavu-sans-700.ttf",
            "PaperSansItalic": font_root / "dejavu-sans-400-italic.ttf",
            "PaperSansBoldItalic": font_root / "dejavu-sans-700-italic.ttf",
        },
    )

    result = paper_report.report_capabilities()

    assert result["formats"]["html"]["available"] is True
    assert result["formats"]["pdf"]["available"] is False
    assert result["formats"]["pdf"]["reason"]


def test_zh_locale_localizes_sections_captions_and_page_furniture(tmp_path):
    from docx import Document
    from pypdf import PdfReader

    model = _model(tmp_path)
    model.update(
        locale="zh-CN",
        title="锂硫电池吸附能筛选报告",
        subtitle="热力学筛选与后续计算建议",
        kicker="第一性原理计算报告",
        executive_summary="吸附能需要满足抓得住、转得动和脱得掉。",
    )
    result = paper_report.render_report_bundle(
        model,
        tmp_path / "zh-bundle",
        stem="中文报告",
    )

    html = result["files"]["html"].read_text(encoding="utf-8")
    for heading in (
        "执行摘要",
        "核心结论",
        "候选评价",
        "吸附能结果",
        "多项目比较",
        "图表",
        "计算方法",
        "局限性",
        "后续建议",
    ):
        assert heading in html
    assert "表 1  Candidate evaluation" in html
    assert "图 1  Free-energy pathways" in html
    assert '"第 " counter(page) " 页"' in html

    doc = Document(result["files"]["docx"])
    doc_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "1. 执行摘要" in doc_text
    assert "表 1  Candidate evaluation" in doc_text
    assert "图 1  Free-energy pathways" in doc_text
    with zipfile.ZipFile(result["files"]["docx"]) as archive:
        footer_xml = archive.read("word/footer1.xml").decode("utf-8")
    assert "第 " in footer_xml and " 页" in footer_xml and "PAGE" in footer_xml

    reader = PdfReader(result["files"]["pdf"])
    pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    compact_pdf_text = re.sub(r"\s+", " ", pdf_text)
    assert "锂硫电池吸附能筛选报告" in pdf_text
    assert "01 执行摘要" in compact_pdf_text
    assert "表 1 Candidate evaluation" in compact_pdf_text
    assert "图 1 Free-energy pathways" in compact_pdf_text
    assert "第 2 页" in compact_pdf_text


def test_table_columns_accept_plain_string_keys(tmp_path):
    model = _model(tmp_path)
    model["comparison_table"] = {
        "title": "Direct string columns",
        "columns": ["candidate", "score"],
        "rows": [
            {"candidate": "A", "score": 0.91},
            {"candidate": "B", "score": 0.86},
        ],
    }
    result = paper_report.render_report_bundle(
        model,
        tmp_path / "string-columns",
        formats=("html",),
    )
    html = result["files"]["html"].read_text(encoding="utf-8")
    assert "<th class=''>candidate</th>" in html
    assert "<th class='numeric'>score</th>" in html
    assert "<td class='numeric'>0.91</td>" in html


def test_pdf_font_root_supports_frozen_bundle_layout(monkeypatch, tmp_path):
    frozen = tmp_path / "frozen"
    # packaging/build_exe.py installs the complete Web asset tree here.
    fonts = frozen / "vcstudio_assets" / "fonts"
    fonts.mkdir(parents=True)
    for name in ("noto-sans-sc-400.ttf", "noto-sans-sc-700.ttf"):
        (fonts / name).write_bytes(b"frozen-font-fixture")
    monkeypatch.setattr(sys, "_MEIPASS", str(frozen), raising=False)
    assert paper_report._pdf_font_root() == fonts


def test_source_pdf_fonts_are_preconverted_true_type():
    root = paper_report._pdf_font_root()
    for name in (
        "noto-sans-sc-400.ttf",
        "noto-sans-sc-700.ttf",
        "dejavu-sans-400.ttf",
        "dejavu-sans-700.ttf",
        "dejavu-sans-400-italic.ttf",
        "dejavu-sans-700-italic.ttf",
    ):
        data = (root / name).read_bytes()
        assert data[:4] in {b"\x00\x01\x00\x00", b"OTTO"}
        assert len(data) > 500_000


def test_bundled_latin_fonts_cover_scientific_text_without_matplotlib(monkeypatch, tmp_path):
    import reportlab
    from reportlab.pdfbase.ttfonts import TTFont

    real_import = paper_report.importlib.import_module

    def no_matplotlib(name, *args, **kwargs):
        if name == "matplotlib":
            raise AssertionError("standalone PDF fonts must not import matplotlib")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(paper_report.importlib, "import_module", no_matplotlib)
    files = paper_report._latin_pdf_font_files(reportlab)
    regular = TTFont("StandaloneScientificCoverage", str(files["PaperSans"]))
    for character in "ΔηLi₂S₈→⁺⁻":
        assert ord(character) in regular.face.charToGlyph

    model = _model(tmp_path)
    model["title"] = "Li₂S₈ → Li₂S · ΔG · η"
    result = paper_report.render_report_bundle(
        model,
        tmp_path / "standalone-fonts",
        formats=("pdf",),
    )
    assert result["files"]["pdf"].stat().st_size > 1_000


def test_font_notice_records_licenses_and_current_hashes():
    root = paper_report._pdf_font_root()
    notice = (root / "NOTICE.md").read_text(encoding="utf-8")
    assert (root / "OFL-1.1.txt").is_file()
    assert (root / "LICENSE-DEJAVU.txt").is_file()
    for name in (
        "noto-sans-sc-400.ttf",
        "noto-sans-sc-700.ttf",
        "dejavu-sans-400.ttf",
        "dejavu-sans-700.ttf",
        "dejavu-sans-400-italic.ttf",
        "dejavu-sans-700-italic.ttf",
    ):
        digest = hashlib.sha256((root / name).read_bytes()).hexdigest()
        assert f"`{name}`" in notice
        assert f"`{digest}`" in notice


def test_same_stem_publish_guard_serializes_independent_processes(tmp_path):
    output = tmp_path / "cross-process-lock"
    output.mkdir()
    release = output / "release-a"
    worker = r"""
import sys
import time
from pathlib import Path
from vcstudio.project.paper_report import _report_publish_guard

output = Path(sys.argv[1])
name = sys.argv[2]
release = Path(sys.argv[3])
(output / f"{name}.started").write_text("started", encoding="utf-8")
with _report_publish_guard(output, "report"):
    (output / f"{name}.entered").write_text("entered", encoding="utf-8")
    if name == "a":
        deadline = time.monotonic() + 15
        while not release.exists():
            if time.monotonic() >= deadline:
                raise TimeoutError("release signal was not received")
            time.sleep(0.02)
"""
    root = Path(__file__).resolve().parents[1]
    env = os.environ.copy()
    env["PYTHONPATH"] = str(root) + os.pathsep + env.get("PYTHONPATH", "")

    def wait_for(path: Path, timeout=10):
        deadline = time.monotonic() + timeout
        while not path.exists():
            if time.monotonic() >= deadline:
                raise AssertionError(f"subprocess did not create {path.name}")
            time.sleep(0.02)

    command = [sys.executable, "-X", "utf8", "-c", worker, str(output)]
    proc_a = subprocess.Popen(
        [*command, "a", str(release)],
        cwd=root,
        env=env,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    proc_b = None
    try:
        wait_for(output / "a.entered")
        proc_b = subprocess.Popen(
            [*command, "b", str(release)],
            cwd=root,
            env=env,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
        wait_for(output / "b.started")
        time.sleep(0.25)
        assert not (output / "b.entered").exists()
        release.write_text("release", encoding="utf-8")
        stdout_a, stderr_a = proc_a.communicate(timeout=15)
        assert proc_a.returncode == 0, stdout_a + stderr_a
        assert proc_b is not None
        stdout_b, stderr_b = proc_b.communicate(timeout=15)
        assert proc_b.returncode == 0, stdout_b + stderr_b
        assert (output / "b.entered").is_file()
        assert (output / ".report.publish.lock").is_file()
    finally:
        release.write_text("release", encoding="utf-8")
        for process in (proc_b, proc_a):
            if process is None or process.poll() is not None:
                continue
            try:
                process.communicate(timeout=3)
            except subprocess.TimeoutExpired:
                process.terminate()
                try:
                    process.communicate(timeout=3)
                except subprocess.TimeoutExpired:
                    process.kill()
                    process.communicate(timeout=3)


def test_publish_failure_rolls_back_all_formats_and_manifest(monkeypatch, tmp_path):
    output = tmp_path / "transactional"
    old = paper_report.render_report_bundle(
        _model(tmp_path),
        output,
        formats=("html", "docx", "pdf"),
    )
    published = {**old["files"], "model": old["model_file"]}
    before = {kind: path.read_bytes() for kind, path in published.items()}

    replacement = _model(tmp_path)
    replacement["title"] = "A replacement that must not be published partially"
    real_replace = paper_report.os.replace

    def fail_while_docx_is_open(source, destination):
        source_path = Path(source)
        destination_path = Path(destination)
        if (
            source_path.suffix == ".docx"
            and destination_path.suffix == ".docx"
            and destination_path.parent == output
        ):
            raise PermissionError("simulated Word file lock")
        return real_replace(source, destination)

    monkeypatch.setattr(paper_report.os, "replace", fail_while_docx_is_open)
    with pytest.raises(PermissionError, match="Word file lock"):
        paper_report.render_report_bundle(
            replacement,
            output,
            formats=("html", "docx", "pdf"),
        )

    for kind, expected in before.items():
        assert published[kind].read_bytes() == expected


def test_manifest_commit_failure_rolls_back_contract_sidecars(monkeypatch, tmp_path):
    output = tmp_path / "contract-transaction"
    old = paper_report.render_report_bundle(
        _validated_model(tmp_path, formats=("html",)),
        output,
        formats=("html",),
    )
    published = {
        **old["files"], **old["contract_files"], "model": old["model_file"]}
    before = {kind: path.read_bytes() for kind, path in published.items()}

    replacement = _model(tmp_path)
    replacement["title"] = "Replacement report that must roll back with its contracts"
    real_replace = paper_report.os.replace

    def fail_at_manifest_commit(source, destination):
        source_path = Path(source)
        destination_path = Path(destination)
        if (
            destination_path.name == "report.manifest.json"
            and source_path.parent.name != "rollback"
        ):
            raise PermissionError("simulated manifest lock")
        return real_replace(source, destination)

    monkeypatch.setattr(paper_report.os, "replace", fail_at_manifest_commit)
    with pytest.raises(PermissionError, match="manifest lock"):
        paper_report.render_report_bundle(
            replacement,
            output,
            formats=("html",),
        )

    for kind, expected in before.items():
        assert published[kind].read_bytes() == expected


def test_incomplete_rollback_withholds_manifest_and_preserves_recovery(
        monkeypatch, tmp_path):
    output = tmp_path / "incomplete-rollback"
    old = paper_report.render_report_bundle(
        _model(tmp_path), output, formats=("html",)
    )
    old_html = old["files"]["html"].read_bytes()
    old_manifest = old["manifest"].read_bytes()

    replacement = _model(tmp_path)
    replacement["title"] = "NEW CONTENT THAT MUST NOT KEEP A COMPLETE MARKER"
    real_replace = paper_report.os.replace
    manifest_failed = {"value": False}

    def fail_manifest_then_html_restore(source, destination):
        source_path = Path(source)
        destination_path = Path(destination)
        if (
            destination_path.name == "report.manifest.json"
            and source_path.parent.name != "rollback"
        ):
            manifest_failed["value"] = True
            raise PermissionError("simulated manifest commit lock")
        if (
            manifest_failed["value"]
            and source_path.parent.name == "rollback"
            and destination_path.name == "report.html"
        ):
            raise PermissionError("simulated HTML rollback lock")
        return real_replace(source, destination)

    monkeypatch.setattr(
        paper_report.os, "replace", fail_manifest_then_html_restore)
    with pytest.raises(paper_report.ReportRecoveryError) as raised:
        paper_report.render_report_bundle(
            replacement, output, formats=("html",)
        )

    recovery_dir = raised.value.recovery_dir
    record_path = recovery_dir / "RECOVERY.json"
    assert recovery_dir.is_dir() and record_path.is_file()
    assert recovery_dir.parent.is_dir()  # the outer temporary generation survived
    assert not (output / "report.manifest.json").exists()
    assert (recovery_dir / "report.manifest.json").read_bytes() == old_manifest
    assert (recovery_dir / "report.html").read_bytes() == old_html
    assert b"NEW CONTENT" in (output / "report.html").read_bytes()

    record = json.loads(record_path.read_text(encoding="utf-8"))
    assert record["artifact_status"] == "recovery_required"
    assert record["canonical_manifest_withheld"] is True
    assert any("HTML rollback lock" in item for item in record["rollback_errors"])
    assert {Path(item["target"]).name for item in record["backups"]} >= {
        "report.html", "report.manifest.json"
    }


def test_unbound_replacement_removes_stale_contract_sidecars(tmp_path):
    output = tmp_path / "remove-stale-contracts"
    bound = paper_report.render_report_bundle(
        _validated_model(tmp_path, formats=("html",)),
        output,
        formats=("html",),
    )
    stale_paths = tuple(bound["contract_files"].values())
    assert all(path.is_file() for path in stale_paths)

    replacement = paper_report.render_report_bundle(
        _model(tmp_path),
        output,
        formats=("html",),
    )

    assert replacement["contract_status"] == "absent"
    assert replacement["contract_files"] == {}
    assert all(not path.exists() for path in stale_paths)
    manifest = json.loads(replacement["manifest"].read_text(encoding="utf-8"))
    assert manifest["contract_status"] == "absent"
    assert manifest["contracts"] == {}


def test_model_sidecar_is_replaced_and_recreated_when_missing(tmp_path):
    output = tmp_path / "model-sidecar-replacement"
    first = paper_report.render_report_bundle(
        _model(tmp_path), output, formats=("html",)
    )
    sidecar = first["model_file"]
    first_bytes = sidecar.read_bytes()
    first_hash = first["model_sha256"]

    replacement_model = _model(tmp_path)
    replacement_model["title"] = "Replacement normalized model"
    replacement = paper_report.render_report_bundle(
        replacement_model, output, formats=("html",)
    )
    assert replacement["model_file"] == sidecar
    assert sidecar.read_bytes() != first_bytes
    assert replacement["model_sha256"] != first_hash
    assert json.loads(sidecar.read_text(encoding="utf-8"))["title"] == (
        "Replacement normalized model"
    )

    sidecar.unlink()
    assert not sidecar.exists()
    recreated = paper_report.render_report_bundle(
        replacement_model, output, formats=("html",)
    )
    assert recreated["model_file"] == sidecar and sidecar.is_file()
    manifest = json.loads(recreated["manifest"].read_text(encoding="utf-8"))
    record = manifest["model_file"]
    assert record["path"] == sidecar.name
    assert record["sha256"] == hashlib.sha256(sidecar.read_bytes()).hexdigest()
    assert record["sha256"] == recreated["model_sha256"]
    assert record["size"] == sidecar.stat().st_size
    assert "model" not in recreated["files"]


def test_replacement_removes_unrequested_stale_formats(tmp_path):
    output = tmp_path / "remove-stale-formats"
    first = paper_report.render_report_bundle(
        _validated_model(tmp_path),
        output,
        formats=("html", "docx", "pdf"),
    )
    stale_docx = first["files"]["docx"]
    stale_pdf = first["files"]["pdf"]
    assert stale_docx.is_file() and stale_pdf.is_file()

    replacement = paper_report.render_report_bundle(
        _validated_model(tmp_path, formats=("html",)),
        output,
        formats=("html",),
    )

    assert set(replacement["files"]) == {"html", "manifest"}
    assert not stale_docx.exists()
    assert not stale_pdf.exists()
    manifest = json.loads(replacement["manifest"].read_text(encoding="utf-8"))
    assert manifest["formats"] == ["html"]
    assert set(manifest["files"]) == {"html"}


def test_manifest_failure_restores_formats_deleted_by_replacement(monkeypatch, tmp_path):
    output = tmp_path / "rollback-stale-formats"
    first = paper_report.render_report_bundle(
        _validated_model(tmp_path),
        output,
        formats=("html", "docx", "pdf"),
    )
    published = {
        **first["files"], **first["contract_files"], "model": first["model_file"]}
    before = {key: path.read_bytes() for key, path in published.items()}
    replacement = _validated_model(tmp_path, formats=("html",))
    replacement["title"] = "HTML-only replacement that must roll back"
    _rebind_report_content(replacement)
    real_replace = paper_report.os.replace

    def fail_at_manifest_commit(source, destination):
        if (
            Path(destination).name == "report.manifest.json"
            and Path(source).parent.name != "rollback"
        ):
            raise PermissionError("simulated manifest lock after format deletion")
        return real_replace(source, destination)

    monkeypatch.setattr(paper_report.os, "replace", fail_at_manifest_commit)
    with pytest.raises(PermissionError, match="after format deletion"):
        paper_report.render_report_bundle(
            replacement,
            output,
            formats=("html",),
        )

    for key, expected in before.items():
        assert published[key].read_bytes() == expected


@pytest.mark.parametrize("stem", ("../escape", r"C:\escape", "", ".."))
def test_unsafe_stem_is_rejected(tmp_path, stem):
    with pytest.raises(ValueError, match="safe filename"):
        paper_report.render_report_bundle(
            _model(tmp_path),
            tmp_path / "bundle",
            stem=stem,
            formats=("html",),
        )
