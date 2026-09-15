"""Focused structural tests for the Office-free DOCX/PDF report renderers."""

from __future__ import annotations

import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest

pytest.importorskip("docx")
pytest.importorskip("reportlab")

from docx import Document
from PIL import Image
from pypdf import PdfReader

from vcstudio.project import report_documents


_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _write_figure(path: Path) -> None:
    image = Image.new("RGB", (800, 400), (245, 248, 252))
    image.save(path, format="PNG")


def _model(tmp_path: Path, figure_path: str = "ladder.png") -> dict:
    return {
        "title": "Catalysis Report 催化计算报告",
        "subtitle": "Adsorption-energy and pathway comparison",
        "report_label": "SCIENTIFIC COMPUTATION REPORT",
        "running_title": "Catalysis Report",
        "footer_label": "Project Li-S | Auditable result",
        "base_dir": str(tmp_path),
        "metadata": {
            "Project": "Li-S screening",
            "Status": "Final",
        },
        "abstract": [
            "This report compares stable adsorption configurations.",
            "中文摘要可与英文内容共同排版。",
        ],
        "sections": [
            {
                "title": "Results",
                "blocks": [
                    {"type": "paragraph", "text": "All energies are reported in eV."},
                    {
                        "type": "table",
                        "caption": "Adsorption energies",
                        "columns": ["Site", "Delta E / eV", "State"],
                        "rows": [
                            {"Site": "top", "Delta E / eV": -1.23, "State": "DONE"},
                            ["bridge", -0.98, "DONE"],
                        ],
                        "column_widths": [2, 1, 1],
                        "note": "Lower values indicate stronger adsorption.",
                    },
                    {
                        "type": "figure",
                        "path": figure_path,
                        "caption": "Adsorption trend",
                        "alt_text": "Adsorption-energy pathway",
                    },
                ],
            },
        ],
    }


def test_generate_documents_are_valid_and_share_content(tmp_path):
    _write_figure(tmp_path / "ladder.png")
    destination = tmp_path / "reports"
    docx_path = destination / "report.docx"
    pdf_path = destination / "report.pdf"

    written = report_documents.generate_report_documents(
        _model(tmp_path), docx_path, pdf_path
    )

    assert written == (docx_path, pdf_path)
    assert zipfile.is_zipfile(docx_path)
    assert pdf_path.read_bytes().startswith(b"%PDF-")

    document = Document(docx_path)
    paragraphs = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Catalysis Report 催化计算报告" in paragraphs
    assert "Table 1. Adsorption energies" in paragraphs
    assert "Figure 1. Adsorption trend" in paragraphs
    assert len(document.tables) == 1
    assert len(document.inline_shapes) == 1

    with zipfile.ZipFile(docx_path) as archive:
        footer_xml = archive.read("word/footer1.xml").decode("utf-8")
        header_xml = archive.read("word/header1.xml").decode("utf-8")
    assert " PAGE " in footer_xml
    assert "Project Li-S" in footer_xml
    assert "Catalysis Report" in header_xml

    reader = PdfReader(str(pdf_path))
    assert reader.pages
    pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "Catalysis Report" in pdf_text
    assert "催化计算报告" in pdf_text
    assert "Adsorption energies" in pdf_text
    assert "Adsorption trend" in pdf_text
    assert "Page 1" in pdf_text


def test_docx_table_has_fixed_matching_ooxml_geometry(tmp_path):
    _write_figure(tmp_path / "ladder.png")
    output = report_documents.generate_docx(_model(tmp_path), tmp_path / "report.docx")

    with zipfile.ZipFile(output) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    namespace = {"w": _W}
    table = root.find(".//w:tbl", namespace)
    assert table is not None
    assert table.find("./w:tblPr/w:tblLayout", namespace).get(f"{{{_W}}}type") == "fixed"
    assert table.find("./w:tblPr/w:tblW", namespace).get(f"{{{_W}}}w") == "9360"
    assert table.find("./w:tblPr/w:tblInd", namespace).get(f"{{{_W}}}w") == "120"

    grid_widths = [
        int(column.get(f"{{{_W}}}w"))
        for column in table.findall("./w:tblGrid/w:gridCol", namespace)
    ]
    assert grid_widths == [4680, 2340, 2340]
    for row in table.findall("./w:tr", namespace):
        cell_widths = [
            int(cell.find("./w:tcPr/w:tcW", namespace).get(f"{{{_W}}}w"))
            for cell in row.findall("./w:tc", namespace)
        ]
        assert cell_widths == grid_widths


@pytest.mark.parametrize("figure_name", ["missing.png", "broken.png"])
def test_invalid_images_are_annotated_without_aborting(tmp_path, figure_name):
    if figure_name == "broken.png":
        (tmp_path / figure_name).write_bytes(b"not an image")
    model = _model(tmp_path, figure_name)
    docx_path, pdf_path = report_documents.generate_report_documents(
        model,
        tmp_path / f"{figure_name}.docx",
        tmp_path / f"{figure_name}.pdf",
    )

    document = Document(docx_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert f'Figure unavailable: "{figure_name}"' in text
    assert "Figure 1. Adsorption trend" in text
    assert len(document.inline_shapes) == 0

    pdf_text = "\n".join(
        page.extract_text() or "" for page in PdfReader(str(pdf_path)).pages
    )
    assert "Figure unavailable" in pdf_text
    assert figure_name in pdf_text
    assert "Adsorption trend" in pdf_text


def test_atomic_docx_write_preserves_existing_destination(tmp_path, monkeypatch):
    target = tmp_path / "existing.docx"
    target.write_bytes(b"previous report")
    sources = []

    def fail_replace(source, destination):
        sources.append(Path(source))
        assert Path(source).parent == target.parent
        assert Path(destination) == target
        raise OSError("simulated replace failure")

    monkeypatch.setattr(report_documents.os, "replace", fail_replace)
    with pytest.raises(OSError, match="simulated replace failure"):
        report_documents.generate_docx(_model(tmp_path, "missing.png"), target)

    assert target.read_bytes() == b"previous report"
    assert len(sources) == 1
    assert not sources[0].exists()
    assert not list(tmp_path.glob(f".{target.name}.*.tmp"))


def test_inconsistent_table_shape_is_rejected_before_writing(tmp_path):
    model = _model(tmp_path)
    model["sections"][0]["blocks"][1]["rows"] = [["top", -1.2]]
    output = tmp_path / "bad.docx"

    with pytest.raises(report_documents.ReportModelError, match="expected 3"):
        report_documents.generate_docx(model, output)
    assert not output.exists()


def test_pdf_list_markers_are_bullets_and_incrementing_numbers(tmp_path):
    model = {
        "title": "List report",
        "sections": [
            {
                "title": "Checks",
                "blocks": [
                    {"type": "bullets", "items": ["Bullet alpha", "Bullet beta"]},
                    {"type": "numbered", "items": ["Step alpha", "Step beta"]},
                ],
            }
        ],
    }
    output = report_documents.generate_pdf(model, tmp_path / "lists.pdf")
    lines = [
        line.strip()
        for page in PdfReader(str(output)).pages
        for line in (page.extract_text() or "").splitlines()
        if line.strip()
    ]

    bullet_alpha = lines.index("Bullet alpha")
    step_alpha = lines.index("Step alpha")
    step_beta = lines.index("Step beta")
    assert lines[bullet_alpha - 1] != "1"
    assert lines[step_alpha - 1] == "1"
    assert lines[step_beta - 1] == "2"
